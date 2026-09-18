import pandas as pd
from copy import copy 
from pathlib import Path
import shutil
from datetime import date, datetime
from openpyxl import load_workbook


# IMPORTA FILE DI SUPPORTO
#-------------------------------------------------------
from OneDrive import copia_in_onedrive 

# IMPORTAZIONE VARIABILI GLOBALI
#--------------------------------------------------------
from config import (COL_NOMI, LISTA_NOMI as LISTA_NOMI_GARA,
                            FILE_MASTER_GIORGIO as FILE_MASTER, COL_DATA_ATTIVAZIONE, COL_DATA_INSERIMENTO, 
                            ONEDRIVE_GIORGIO as PERCORSO_GIORGIO, elimina_file_da_cartella)

# LEGGERE AGENTI DA FILE LISTA NOMI
#--------------------------------------------------------
def leggi_agenti_gara(lista_nomi):
    df = pd.read_excel(lista_nomi, sheet_name="nomi", header=None)
    return df[COL_NOMI]

# COPIA FORMATTAZIONE CELLE
#--------------------------------------------------------
def copia_formattaz_cella(cella_input, cella_output):
    if cella_input.has_style:
        cella_output.font = copy(cella_input.font)
        cella_output.border = copy(cella_input.border)
        cella_output.fill = copy(cella_input.fill)
        cella_output.number_format = copy(cella_input.number_format)
        cella_output.protection  = copy(cella_input.protection)
        cella_output.alignment   = copy(cella_input.alignment)

# COPIA TEMPLATE (una volta sola)
#--------------------------------------------------------
def copia_template(file_master, cartella_output):
    cartella_output = Path(cartella_output)
    cartella_output.mkdir(parents=True, exist_ok=True)
    file_output = cartella_output / "FCST_GIORGIO.xlsx"
    shutil.copy(file_master, file_output)
    return file_output

# SCRIVI UN DATAFRAME IN UN FOGLIO ACCODANDO LE RIGHE
#--------------------------------------------------------
def scrivi_df_in_foglio(ws, df, riga_start):
    if df is None or df.empty:
        return riga_start
    n_cols = len(df.columns)
    stili_template = [ws.cell(row=3, column=idx_col) for idx_col in range(2, n_cols + 2)]
    for idx_riga, row_data in enumerate(df.values, start=riga_start):
        for idx_col, valore in enumerate(row_data, start=2):
            cella = ws.cell(row=idx_riga, column=idx_col)
            copia_formattaz_cella(stili_template[idx_col - 2], cella)
            if isinstance(valore, (datetime, pd.Timestamp)) and not pd.isnull(valore):
                cella.value = valore if isinstance(valore, datetime) else valore.to_pydatetime()
                cella.number_format = "DD/MM/YYYY"
            else:
                cella.value = valore
    return riga_start + len(df)

# POPOLA FOGLIO OPPORTUNITA
#--------------------------------------------------------
def popola_foglio_opp(wb, df_opp, righe_correnti):
    ws_opp = wb["OPPORTUNITA"]
    if df_opp is None or df_opp.empty:
        print("  ⚠️  Nessuna opportunità aperta trovata")
        return
    righe_correnti["OPPORTUNITA"] = scrivi_df_in_foglio(ws_opp, df_opp, righe_correnti["OPPORTUNITA"])

# POPOLA FOGLIO APPUNTAMENTI
#--------------------------------------------------------
def popola_foglio_app(wb, df_app, righe_correnti):
    ws_app = wb["APPUNTAMENTI"]
    if df_app is None or df_app.empty:
        print("  ⚠️  Nessun appuntamento trovato")
        return
    df_app["Data appuntamento"] = pd.to_datetime(df_app["Data appuntamento"], format="%d/%m/%Y", errors="coerce")
    righe_correnti["APPUNTAMENTI"] = scrivi_df_in_foglio(ws_app, df_app, righe_correnti["APPUNTAMENTI"])

# POPOLA FOGLIO FCST
#--------------------------------------------------------
def popola_foglio_fcst(wb, df_da_inserire, df_in_attivazione, df_attive, agenti):
    ws_fcst = wb["FCST"]
    ws_fcst.cell(row=2, column=3).value = date.today().strftime("%d/%m/%Y")
    for idx, nome_gara in enumerate(agenti):
        riga = 10 + idx
        ws_fcst.cell(row=riga, column=1).value = nome_gara
        ws_fcst.cell(row=riga, column=3).value = nome_gara

# SCRIVI DATI NEI FOGLI
#--------------------------------------------------------
def scrivi_dati_nei_fogli(wb, df_gara, df_opp, df_app, righe_correnti, agenti):

    inserimento_vuoto = (df_gara[COL_DATA_INSERIMENTO].isna()) | (df_gara[COL_DATA_INSERIMENTO].astype(str).str.strip() == "")
    attivazione_vuota = (df_gara[COL_DATA_ATTIVAZIONE].isna()) | (df_gara[COL_DATA_ATTIVAZIONE].astype(str).str.strip() == "")

    df_da_inserire    = df_gara[ inserimento_vuoto &  attivazione_vuota]
    df_in_attivazione = df_gara[~inserimento_vuoto &  attivazione_vuota]
    df_attive         = df_gara[~inserimento_vuoto & ~attivazione_vuota]

    for nome_foglio, df in [("DA INSERIRE", df_da_inserire), ("IN ATTIVAZIONE", df_in_attivazione), ("ATTIVE", df_attive)]:
        if df.empty:
            continue
        ws = wb[nome_foglio]
        righe_correnti[nome_foglio] = scrivi_df_in_foglio(ws, df, righe_correnti[nome_foglio])
    
    popola_foglio_opp(wb, df_opp, righe_correnti)
    popola_foglio_app(wb, df_app, righe_correnti)
    popola_foglio_fcst(wb, df_da_inserire, df_in_attivazione, df_attive, agenti)

# FUNZIONE MAIN
# I file sono già puliti — legge direttamente senza prepara_*
#--------------------------------------------------------
def main(df_gara, cartella_output, df_opp, df_app):
    percorso = PERCORSO_GIORGIO / "FCST_GIORGIO.xlsx!"
    elimina_file_da_cartella(percorso)

    righe_correnti = {
        "DA INSERIRE":    2,
        "IN ATTIVAZIONE": 2,
        "ATTIVE":         2,
        "OPPORTUNITA":    2,
        "FCST":           10,
        "APPUNTAMENTI":   2,
    }

    file_consolidato = copia_template(FILE_MASTER, cartella_output)
    wb = load_workbook(file_consolidato)

    agenti = leggi_agenti_gara(LISTA_NOMI_GARA)

    scrivi_dati_nei_fogli(wb, df_gara, df_opp, df_app, righe_correnti, agenti)
    wb.save(file_consolidato)
    wb.close()

    try:
        copia_in_onedrive(file_consolidato, PERCORSO_GIORGIO)
        print("\n  ✅ salvato su TEAMS\n")
    except Exception as e:
        print(f"Agente: GIORGIO non trovato, vado avanti {e}\n")

    print(f"\n📁 File consolidato salvato in: {file_consolidato}")
    print("\nElaborazione completata! 🎉")