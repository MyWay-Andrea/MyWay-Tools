import os
from io import BytesIO

import pandas as pd
from pathlib import Path
import sys
from colorama import Fore, init
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

BUSINESS_DIR = Path(__file__).resolve().parents[1]
SCRIPT_DIR = BUSINESS_DIR.parent
if str(BUSINESS_DIR) not in sys.path:
    sys.path.insert(0, str(BUSINESS_DIR))
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from common.crm_reportistica import API_KEY as REPORTISTICA_API_KEY, BASE_URL as REPORTISTICA_BASE_URL
from graph_sharepoint import GraphSharePointClient

init(autoreset=True)


'''
file configurazione per i dati dell'API e per la mapaptura dei dati
'''

API_KEY = REPORTISTICA_API_KEY
BASE_URL = REPORTISTICA_BASE_URL

# PERCORSI SHAREPOINT GRAPH
SHAREPOINT_HOSTNAME = os.getenv("SHAREPOINT_HOSTNAME", "").strip()
SHAREPOINT_SCAMBIO_SITE_PATH = os.getenv("SHAREPOINT_SCAMBIO_SITE_PATH", "").strip()
SHAREPOINT_SCAMBIO_LIBRARY_NAME = os.getenv("SHAREPOINT_SCAMBIO_LIBRARY_NAME", "").strip()
PATH_PARQUET = os.getenv("SHAREPOINT_SCAMBIO_PARQUET_FOLDER", "").strip(" /")
PATH_BASI_DATI = os.getenv("SHAREPOINT_BUSINESS_BASI_DATI_FOLDER", "").strip(" /")
PATH_STORICO_BASI_DATI = f"{PATH_BASI_DATI}/STORICO"
PATH_STORICO_PARQUET = f"{PATH_PARQUET}/STORICO"

_GRAPH_CLIENT = None
_GRAPH_DRIVE_ID = None


def _drive_sharepoint():
    global _GRAPH_CLIENT, _GRAPH_DRIVE_ID
    configurazione = {
        "SHAREPOINT_HOSTNAME": SHAREPOINT_HOSTNAME,
        "SHAREPOINT_SCAMBIO_SITE_PATH": SHAREPOINT_SCAMBIO_SITE_PATH,
        "SHAREPOINT_SCAMBIO_LIBRARY_NAME": SHAREPOINT_SCAMBIO_LIBRARY_NAME,
        "SHAREPOINT_SCAMBIO_PARQUET_FOLDER": PATH_PARQUET,
        "SHAREPOINT_BUSINESS_BASI_DATI_FOLDER": PATH_BASI_DATI,
    }
    mancanti = [nome for nome, valore in configurazione.items() if not valore]
    if mancanti:
        raise ValueError(
            "Configurazione SharePoint Opportunita incompleta: " + ", ".join(mancanti)
        )
    if _GRAPH_CLIENT is None:
        _GRAPH_CLIENT = GraphSharePointClient()
        sito = _GRAPH_CLIENT.trova_sito(
            SHAREPOINT_HOSTNAME, SHAREPOINT_SCAMBIO_SITE_PATH
        )
        raccolta = _GRAPH_CLIENT.trova_raccolta_documenti(
            sito["id"], SHAREPOINT_SCAMBIO_LIBRARY_NAME
        )
        _GRAPH_DRIVE_ID = str(raccolta["id"])
    return _GRAPH_CLIENT, _GRAPH_DRIVE_ID


STATO_MAP = {
    59096 : "Aperta",
    59101 : "Persa",
    59100 : "Acquisita",
}

TIP_SERVIZIO_MAP = {

    "1": "ALTRO",
    "2": "SIM VOCE",
    "3": "SIM DATI",
    "4": "DSL",
    "5": "ON TOP DSL",
    "6": "VRU INTERNO",
    "7": "LINK SA",
    "8": "LINK SU",
    "9": "TW SA",
    "10": "ON TOP LINK",
    "11": "M2M",
    "12": "SAAS",
    "13": "COLLABORATION",
    "14": "IAAS",
    "15": "SECURITY",
    "16": "COMPLEX DEAL",
    "17": "CUSTOM APP",
    "18": "EASY DEAL",
    "19": "EOLO",
    "20": "RATA TEL",
    "21": "SOL TEL",
    "22": "RATA TEL VOCE",
    "23": "SOL TEL VOCE",
    "24": "ON TOP EASY DEAL",
    "25": "RATA TEL DATI",
    "26": "SOL TEL DATI",
    "27": "NOLEGGIO OP.",
    "28": "ENEGERGIA"
}

CONVERSIONE_NOMI = {
    "Giorgio Mandelli": "G.MANDELLI",
    "Giovanni Marra": "G.MARRA",
    "Eduart Agalliu": "E.AGALLIU",
    "Alberto Milani": "A.MILANI",
    "Danilo Carugo": "D.CARUGO",
    "Doriano Gnani": "D.GNANI",
    "Enrico Pignatta": "E.PIGNATTA",
    "Mauro Sala": "M.SALA",
    "Riccardo Scimone": "R.SCIMONE",
    "Fabio Arlotta": "F.ARLOTTA",
    "CRM MWR": "CRM My Way",
}

COLONNE_OPPORTUNITA = [
    "Codice Opportunità",
    "Potenziale cliente (Opportunità)",
    "Partita IVA (Opportunità)",
    "Proprietario (Opportunità)",
    "Stato (Opportunità)",
    "Data creazione",
    "Data chiusura prev. (Opportunità)",
    "Data modifica",
    "Tipologia Servizio",
    "Prodotto esistente",
    "Quantità",
    "Prezzo",
    "Importo totale",
    "Probabilità (Opportunità)",
]


def trova_file(path_cartella: str, parametro: str) -> dict:
    client, drive_id = _drive_sharepoint()
    file_trovati = client.elenca_file_cartella(drive_id, path_cartella)
    parola = "storico opportunita" if parametro == "Storico" else "opportunita"
    candidati = [
        file for file in file_trovati
        if parola in str(file.get("name", "")).casefold()
    ]
    if not candidati:
        raise FileNotFoundError(
            f"Nessun file {parametro!r} trovato in SharePoint: {path_cartella}"
        )
    candidati.sort(
        key=lambda file: str(file.get("lastModifiedDateTime", "")), reverse=True
    )
    scelto = candidati[0]
    return {**scelto, "content": client.scarica_file(drive_id, scelto["id"])}


def _formatta_excel(contenuto: BytesIO) -> bytes:
    col_prezzi = set()
    col_date = set()
    col_prob = set()

    contenuto.seek(0)
    wb = load_workbook(contenuto)
    ws = wb.active

    fill_blu = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    fill_giallo = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    fill_rosso = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    font_bianco = Font(color="FFFFFF", bold=True)
    font_nero = Font(color="000000", bold=True)
    font_verde_testo = Font(color="2FBF47", bold=True)
    font_arancione_testo = Font(color="EA6B14", bold=True)
    font_rosso_testo = Font(color="FF0000", bold=True)
    font_azzurro_testo = Font(color="00B0F0", bold=True)

    for cella in ws[1]:
        nome_colonna = str(cella.value).strip().lower() if cella.value else ""

        if "proprietario" in nome_colonna:
            cella.fill = fill_giallo
            cella.font = font_nero
        elif "probabilit" in nome_colonna:
            cella.fill = fill_rosso
            cella.font = font_bianco
        else:
            cella.fill = fill_blu
            cella.font = font_bianco

        if "prezzo" in nome_colonna or "importo" in nome_colonna:
            col_prezzi.add(cella.column)
        elif "data" in nome_colonna:
            col_date.add(cella.column)
        elif "probabilit" in nome_colonna:
            col_prob.add(cella.column)

    for riga in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cella in riga:
            if cella.column in col_date:
                cella.number_format = "DD/MM/YYYY"
            elif cella.column in col_prezzi:
                cella.number_format = "#,##0.00 €"
            elif cella.column in col_prob:
                cella.number_format = r"0\%"
            elif str(cella.value).strip().lower() == "aperta":
                cella.font = font_verde_testo
            elif str(cella.value).strip().lower() == "acquisita":
                cella.font = font_azzurro_testo
            elif str(cella.value).strip().lower() == "persa":
                cella.font = font_rosso_testo

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def carica_su_teams(df: pd.DataFrame, *paths: str) -> None:
    try:
        if not isinstance(df, pd.DataFrame):
            raise TypeError(f"Atteso pandas.DataFrame, ricevuto {type(df).__name__}")

        for path in paths:
            storico = path in (PATH_STORICO_BASI_DATI, PATH_STORICO_PARQUET)
            nome_base = "Storico Opportunita" if storico else "opportunita"
            buffer = BytesIO()
            if path in (PATH_PARQUET, PATH_STORICO_PARQUET):
                nome_file = f"{nome_base}.parquet"
                df.to_parquet(buffer, index=False)
                contenuto = buffer.getvalue()
                content_type = "application/octet-stream"
            elif path in (PATH_BASI_DATI, PATH_STORICO_BASI_DATI):
                nome_file = f"{nome_base}.xlsx"
                df.to_excel(buffer, index=False, engine="openpyxl")
                contenuto = _formatta_excel(buffer)
                content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            else:
                raise ValueError(
                    f"Destinazione SharePoint Opportunita non riconosciuta: {path}"
                )

            client, drive_id = _drive_sharepoint()
            caricato = client.carica_bytes(
                drive_id, path, nome_file, contenuto, content_type=content_type
            )
            print(Fore.GREEN + f"SharePoint: {caricato.get('webUrl', nome_file)}")

        print(Fore.GREEN + "Caricamento su SharePoint avvenuto con successo\n")
    except Exception as errore:
        print(f"[{Fore.RED}ERRORE{Fore.RESET}] Caricamento file su SharePoint fallito --> {errore}")
        raise


def elimina_file(file_path: Path) -> None:
    if file_path.is_file():
        file_path.unlink()
