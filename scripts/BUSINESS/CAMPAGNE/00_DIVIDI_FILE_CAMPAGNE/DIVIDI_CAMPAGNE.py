"""Divide i file campagna per venditore e li pubblica tramite Microsoft Graph."""

from __future__ import annotations

import argparse
import shutil
import tempfile
from datetime import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from config import (
    PATH_OUTPUT_TEST,
    copia_stile_cella,
    pubblica_campagna,
    pubblica_campagne_divise,
    scarica_file_campagna,
    trova_venditore,
)


def trova_colonna_agente(df):
    for colonna in df.columns:
        nome = str(colonna).casefold()
        if "agente" in nome or "assegnazione" in nome or "assegnat" in nome:
            return colonna
    return None


def normalizza_agente(valore):
    return " ".join(str(valore).strip().split(" - ", 1)[0].replace(",", " ").split())


def leggi_agenti(files_sorgente):
    agenti = []
    for file in files_sorgente:
        df = pd.read_excel(file, dtype=str)
        colonna = trova_colonna_agente(df)
        if colonna is None:
            print(f"  ATTENZIONE: colonna Agente/Assegnazione assente in {file.name}")
            continue
        for valore in df[colonna].dropna().unique():
            agente = normalizza_agente(valore)
            if agente and agente.casefold() not in {a.casefold() for a in agenti}:
                agenti.append(agente)
    return agenti


def filtra_per_agente(file_excel, agente):
    df = pd.read_excel(file_excel, dtype=str)
    colonna = trova_colonna_agente(df)
    if colonna is None:
        return None, None
    maschera = df[colonna].apply(
        lambda valore: not pd.isna(valore)
        and normalizza_agente(valore).casefold() == agente.casefold()
    )
    indici = df[maschera].index.tolist()
    if not indici:
        return None, None
    filtrato = df[maschera].copy()
    filtrato.insert(0, "Gestito (Si/No)", "No")
    filtrato.insert(1, "Concluso (Si/No)", "No")
    filtrato.insert(2, "Tipologia", file_excel.stem)
    return filtrato, indici


def crea_campagna_agente(agente, files_sorgente, cartella_output):
    file_output = cartella_output / f"{dt.today():%Y-%m}_CAMPAGNA_{agente}.xlsx"
    fogli = {}
    for file in files_sorgente:
        df, indici = filtra_per_agente(file, agente)
        if df is not None:
            fogli[file.stem[:31]] = {"df": df, "indici": indici, "file_origine": file}
    if not fogli:
        return None, 0

    with pd.ExcelWriter(file_output, engine="openpyxl") as writer:
        for nome_foglio, dati in fogli.items():
            dati["df"].to_excel(writer, sheet_name=nome_foglio, index=False)

    wb_output = load_workbook(file_output)
    totale_righe = 0
    for nome_foglio, dati in fogli.items():
        ws_output = wb_output[nome_foglio]
        wb_origine = load_workbook(dati["file_origine"])
        ws_origine = wb_origine.active
        totale_righe += len(dati["df"])
        for colonna in range(1, ws_origine.max_column + 1):
            copia_stile_cella(ws_origine.cell(1, colonna), ws_output.cell(1, colonna + 3))
        for riga_output, indice_origine in enumerate(dati["indici"], start=2):
            for colonna in range(1, ws_origine.max_column + 1):
                copia_stile_cella(
                    ws_origine.cell(indice_origine + 2, colonna),
                    ws_output.cell(riga_output, colonna + 3),
                )
        for colonna in range(1, ws_origine.max_column + 1):
            lettera_origine = ws_origine.cell(1, colonna).column_letter
            lettera_destinazione = ws_output.cell(1, colonna + 3).column_letter
            ws_output.column_dimensions[lettera_destinazione].width = (
                ws_origine.column_dimensions[lettera_origine].width
            )
        wb_origine.close()
    wb_output.save(file_output)
    wb_output.close()
    return file_output, totale_righe


def _converti_xls(files_sorgente):
    """Converte eventuali .xls scaricati, mantenendo invariato il flusso storico."""
    risultati = []
    xls = [file for file in files_sorgente if file.suffix.casefold() == ".xls"]
    if not xls:
        return list(files_sorgente)
    try:
        import win32com.client
    except ImportError as errore:
        raise RuntimeError("Trovato un file .xls ma pywin32 non e installato") from errore
    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        for file in files_sorgente:
            if file.suffix.casefold() != ".xls":
                risultati.append(file)
                continue
            destinazione = file.with_suffix(".xlsx")
            workbook = excel.Workbooks.Open(str(file.resolve()))
            try:
                workbook.SaveAs(str(destinazione.resolve()), FileFormat=51)
            finally:
                workbook.Close(False)
            risultati.append(destinazione)
    finally:
        excel.Quit()
    return risultati


def esegui(files_sorgente, cartella_output, *, test, origine, pubblica_centrale=True):
    cartella_output.mkdir(parents=True, exist_ok=True)
    agenti = leggi_agenti(files_sorgente)
    if not agenti:
        raise ValueError("Nessun agente trovato nei file sorgente")

    print(f"Origine: {origine}")
    print(f"File sorgente: {len(files_sorgente)}")
    print(f"Agenti trovati: {len(agenti)}")
    print(f"Modalita: {'TEST (nessuna modifica remota)' if test else 'REALE'}\n")

    creati = 0
    non_mappati = []
    risultati = []
    for agente in agenti:
        file_output, righe = crea_campagna_agente(agente, files_sorgente, cartella_output)
        if file_output is None:
            continue
        venditore = trova_venditore(agente)
        risultati.append((agente, venditore, file_output))
        mappatura = venditore["nome_crm"] if venditore else "non mappato"
        print(f"OK {file_output.name}: {righe} righe, {mappatura}")
        creati += 1

    files_output = [risultato[2] for risultato in risultati]
    if pubblica_centrale:
        pubblica_campagne_divise(files_output, origine, test=test)
    else:
        print("\n[TEST LOCALE] Upload in CAMPAGNE_DIVISE non eseguito")

    print("\nDistribuzione nelle cartelle personali:")
    for agente, venditore, file_output in risultati:
        if venditore is None:
            non_mappati.append(agente)
            print(f"  SKIP {agente}: presente in CAMPAGNE_DIVISE ma non mappato nel JSON")
            continue
        pubblica_campagna(venditore, file_output, test=test)

    print(f"\nCreati: {creati}/{len(agenti)} file")
    print(f"Output locale: {cartella_output}")
    if non_mappati:
        print("Non distribuiti nelle cartelle personali: " + ", ".join(non_mappati))


def leggi_argomenti():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--test",
        action="store_true",
        help="crea i file localmente ma non cancella e non carica nulla su SharePoint",
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        help="cartella locale alternativa per provare lo script senza scaricare da SharePoint",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        help="cartella output locale (predefinita: OUTPUT_TEST in test)",
    )
    return parser.parse_args()


def main():
    args = leggi_argomenti()
    test = args.test or args.input_dir is not None
    if args.input_dir:
        input_dir = args.input_dir.resolve()
        files = sorted(
            file for file in input_dir.iterdir()
            if file.is_file()
            and file.suffix.casefold() in {".xls", ".xlsx"}
            and not file.name.startswith("~$")
        )
        if not files:
            raise FileNotFoundError(f"Nessun file Excel in {input_dir}")
        output = (args.output_dir or PATH_OUTPUT_TEST).resolve()
        esegui(
            _converti_xls(files),
            output,
            test=True,
            origine=str(input_dir),
            pubblica_centrale=False,
        )
        return

    with tempfile.TemporaryDirectory(prefix="campagne_") as directory:
        temporanea = Path(directory)
        origine, files = scarica_file_campagna(temporanea / "input")
        files = _converti_xls(files)
        output_temporaneo = temporanea / "output"
        output = (args.output_dir or PATH_OUTPUT_TEST).resolve() if test else output_temporaneo
        esegui(files, output, test=test, origine=origine)
        if not test and args.output_dir:
            args.output_dir.mkdir(parents=True, exist_ok=True)
            for file in output_temporaneo.glob("*.xlsx"):
                shutil.copy2(file, args.output_dir / file.name)


if __name__ == "__main__":
    main()
