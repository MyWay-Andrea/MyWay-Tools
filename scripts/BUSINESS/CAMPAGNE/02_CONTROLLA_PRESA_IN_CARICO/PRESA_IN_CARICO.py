import re
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from datetime import datetime
from colorama import Fore, init
init(autoreset=True)
from config import (
    applica_logica_assegnata,
    copia_stile_cella,
    trova_raw_file,
    _cancella_file,
    ONEDRIVE_VENDITORI,
    CARTELLA_RAW_FILE,
    ONEDRIVE_CRM_SCAMBIO_DOCUMENTI,
    TIPOLOGIE_INCLUSE,
)

def _trova_colonna_agente(df: pd.DataFrame) -> str | None:
    for col in df.columns:
        if re.search(r"agente", col.strip().lower()):
            return col
    return None

def normalizza_agente(valore: str) -> str:
    normalizzato = valore.strip().split(' - ')[0].strip()
    normalizzato = normalizzato.replace(',', '').strip()
    return normalizzato.lower()

def applica_formattazione(wb_dest, wb_orig_path, fogli_presenti):
    wb_orig = load_workbook(wb_orig_path)
    for nome_foglio in fogli_presenti:
        if nome_foglio not in wb_orig.sheetnames or nome_foglio not in wb_dest.sheetnames:
            continue
        ws_o = wb_orig[nome_foglio]
        ws_d = wb_dest[nome_foglio]
        for col in range(1, ws_d.max_column + 1):
            copia_stile_cella(ws_o.cell(1, col), ws_d.cell(1, col))
        if ws_o.max_row >= 2:
            for riga in range(2, ws_d.max_row + 1):
                for col in range(1, ws_d.max_column + 1):
                    copia_stile_cella(ws_o.cell(2, col), ws_d.cell(riga, col))
        for col in range(1, min(ws_o.max_column, ws_d.max_column) + 1):
            lettera = get_column_letter(col)
            if ws_o.column_dimensions[lettera].width:
                ws_d.column_dimensions[lettera].width = ws_o.column_dimensions[lettera].width
    wb_orig.close()

def salva_excel(file_output, fogli: dict, file_stile: Path):
    with pd.ExcelWriter(file_output, engine="openpyxl") as writer:
        for nome_foglio, df in fogli.items():
            df.to_excel(writer, sheet_name=nome_foglio, index=False)
    wb = load_workbook(file_output)
    applica_formattazione(wb, file_stile, list(fogli.keys()))
    wb.save(file_output)

def main():
    cartella_campagna, _ = trova_raw_file(CARTELLA_RAW_FILE)
    cartella_compilate   = cartella_campagna / "CAMPAGNE COMPILATE"
    cartella_pic         = cartella_campagna / "PRESE IN CARICO"

    if not cartella_compilate.exists():
        print(Fore.RED + "❌ Cartella 'CAMPAGNE COMPILATE' non trovata!")
        input("\nPremi INVIO per uscire...")
        return

    cartella_pic.mkdir(exist_ok=True)
    prefisso = datetime.today().strftime("%Y_%m")

    # ── STEP 1: leggi, trasforma, salva in PRESE IN CARICO e su OneDrive CRM ─
    print("=" * 50)
    print("STEP 1 – PRESE IN CARICO")
    print("=" * 50)

    # dati[tipologia][foglio] = df  |  stile[tipologia] = file_path
    dati:  dict[str, dict[str, pd.DataFrame]] = {}
    stile: dict[str, Path] = {}

    for file in sorted(cartella_compilate.glob("*.xlsx")):
        tipologia = file.stem.replace("COMPILATA_", "")
        print(Fore.LIGHTMAGENTA_EX + f"  → {file.name}  [tipologia: {tipologia}]")

        fogli = {}
        with pd.ExcelFile(file) as xl:
            for foglio in xl.sheet_names:
                df = pd.read_excel(xl, sheet_name=foglio, dtype=str)
                if "Partita_IVA" in df.columns:
                    df["Partita_IVA"] = df["Partita_IVA"].str.zfill(11)
                col_agente = _trova_colonna_agente(df)
                df = applica_logica_assegnata(df, col_agente=col_agente or "AGENTE")
                fogli[foglio] = df

        salva_excel(cartella_pic / file.name, fogli, file)
        print(Fore.GREEN + f"    ✓ Salvato in PRESE IN CARICO")

        tipologia_inclusa = any(
            re.search(t.strip().lower(), tipologia.strip().lower())
            for t in TIPOLOGIE_INCLUSE
        )
        if tipologia_inclusa:
            if ONEDRIVE_CRM_SCAMBIO_DOCUMENTI.exists():
                nome_crm = f"RAUL_{tipologia}.xlsx"
                salva_excel(ONEDRIVE_CRM_SCAMBIO_DOCUMENTI / nome_crm, fogli, file)
                print(Fore.CYAN + f"    ✓ Caricato su CRM: {nome_crm}")
            else:
                print(Fore.YELLOW + f"    ⚠ Cartella CRM non trovata: {ONEDRIVE_CRM_SCAMBIO_DOCUMENTI}")

        dati[tipologia]  = fogli
        stile[tipologia] = file

    if not dati:
        print(Fore.RED + "\n❌ Nessun dato trovato.")
        input("\nPremi INVIO per uscire...")
        return

    # ── STEP 2: dividi per venditore e invia su OneDrive ────────────────────
    print("\n" + "=" * 50)
    print("STEP 2 – DISTRIBUZIONE AI VENDITORI")
    print("=" * 50 + "\n")

    venditori_normalizzati = {
        normalizza_agente(venditore): (venditore, cartelle if isinstance(cartelle, list) else [cartelle])
        for venditore, cartelle in ONEDRIVE_VENDITORI.items()
    }

    for nome_norm, (venditore, cartelle) in venditori_normalizzati.items():

        # chiave = tipologia (es. "Reward", "evo voce") → 1 foglio per tipologia nel file finale
        fogli_venditore: dict[str, pd.DataFrame] = {}
        file_stile_ref:  Path | None = None

        for tipologia, fogli in dati.items():
            for nome_foglio, df in fogli.items():
                col_agente = _trova_colonna_agente(df)
                if col_agente is None:
                    continue
                df_filtrato = df[df[col_agente].astype(str).map(normalizza_agente) == nome_norm]
                if df_filtrato.empty:
                    continue

                # Usa tipologia come nome foglio, non nome_foglio
                if tipologia in fogli_venditore:
                    fogli_venditore[tipologia] = pd.concat(
                        [fogli_venditore[tipologia], df_filtrato], ignore_index=True
                    )
                else:
                    fogli_venditore[tipologia] = df_filtrato.copy()
                    if file_stile_ref is None:
                        file_stile_ref = stile[tipologia]

        if not fogli_venditore:
            print(Fore.YELLOW + f"  ⚠ {venditore}: nessuna riga, file non creato.")
            continue

        nome_file = f"{prefisso}_CAMPAGNA_{venditore}.xlsx"
        for cartella in cartelle:
            if not cartella.exists():
                print(Fore.YELLOW + f"  ⚠ Cartella non trovata: {cartella}")
                continue
            _cancella_file(cartella, rf"CAMPAGNA_{re.escape(venditore)}")
            salva_excel(cartella / nome_file, fogli_venditore, file_stile_ref)
            totale = sum(len(d) for d in fogli_venditore.values())
            print(Fore.GREEN + f"  ✓ {nome_file}  ({len(fogli_venditore)} fogli, {totale} righe)")
            print(Fore.CYAN  + f"     → {cartella / nome_file}")

    print("\n" + "=" * 50)
    print(Fore.GREEN + "✓ Completato!")
    print("=" * 50)
    

if __name__ == "__main__":
    main()