import pandas as pd
import os
from io import BytesIO
from pathlib import Path
import sys
from datetime import datetime
from colorama import Fore, Style, init
init(autoreset=True)
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

BUSINESS_DIR = Path(__file__).resolve().parents[1]
SCRIPT_DIR = BUSINESS_DIR.parent
if str(BUSINESS_DIR) not in sys.path:
    sys.path.insert(0, str(BUSINESS_DIR))
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from common.crm_reportistica import API_KEY as REPORTISTICA_API_KEY, BASE_URL as REPORTISTICA_BASE_URL
from graph_sharepoint import GraphSharePointClient

'''
file configurazione per i dati dell'API e per la mapaptura dei dati
'''

API_KEY = REPORTISTICA_API_KEY
BASE_URL = REPORTISTICA_BASE_URL


#PERCORSI SHAREPOINT GRAPH
#-------------------------------------------------------------------------------------
SHAREPOINT_HOSTNAME = os.getenv("SHAREPOINT_HOSTNAME", "").strip()
SHAREPOINT_SCAMBIO_SITE_PATH = os.getenv("SHAREPOINT_SCAMBIO_SITE_PATH", "").strip()
SHAREPOINT_SCAMBIO_LIBRARY_NAME = os.getenv("SHAREPOINT_SCAMBIO_LIBRARY_NAME", "").strip()
PATH_PARQUET = os.getenv("SHAREPOINT_SCAMBIO_PARQUET_FOLDER", "").strip(" /")
PATH_BASI_DATI = os.getenv("SHAREPOINT_BUSINESS_BASI_DATI_FOLDER", "").strip(" /")
PATH_STORICO_BASI_DATI = f"{PATH_BASI_DATI}/STORICO"
PATH_STORICO_PARQUET = f"{PATH_PARQUET}/STORICO"
FILE_STORICO_GARA_PARQUET = "Storico Gara.parquet"
PATH_BACKUP_STORICO_GARA = f"{PATH_STORICO_PARQUET}/BACKUP"
PATH_BACKUP_STORICO_GARA_XLSX = f"{PATH_STORICO_BASI_DATI}/BACKUP"

_GRAPH_CLIENT = None
_GRAPH_DRIVE_ID = None


def _drive_sharepoint():
    global _GRAPH_CLIENT, _GRAPH_DRIVE_ID
    mancanti = [
        nome for nome, valore in {
            "SHAREPOINT_HOSTNAME": SHAREPOINT_HOSTNAME,
            "SHAREPOINT_SCAMBIO_SITE_PATH": SHAREPOINT_SCAMBIO_SITE_PATH,
            "SHAREPOINT_SCAMBIO_LIBRARY_NAME": SHAREPOINT_SCAMBIO_LIBRARY_NAME,
            "SHAREPOINT_SCAMBIO_PARQUET_FOLDER": PATH_PARQUET,
            "SHAREPOINT_BUSINESS_BASI_DATI_FOLDER": PATH_BASI_DATI,
        }.items() if not valore
    ]
    if mancanti:
        raise ValueError("Configurazione SharePoint Gara incompleta: " + ", ".join(mancanti))
    if _GRAPH_CLIENT is None:
        _GRAPH_CLIENT = GraphSharePointClient()
        sito = _GRAPH_CLIENT.trova_sito(SHAREPOINT_HOSTNAME, SHAREPOINT_SCAMBIO_SITE_PATH)
        raccolta = _GRAPH_CLIENT.trova_raccolta_documenti(
            sito["id"], SHAREPOINT_SCAMBIO_LIBRARY_NAME
        )
        _GRAPH_DRIVE_ID = str(raccolta["id"])
    return _GRAPH_CLIENT, _GRAPH_DRIVE_ID


def elenca_file_sharepoint(cartella, crea=False):
    client, drive_id = _drive_sharepoint()
    return client.elenca_file_cartella(drive_id, cartella, crea=crea)


def scarica_file_sharepoint(cartella, nome_file, obbligatorio=True):
    file_trovati = elenca_file_sharepoint(cartella, crea=not obbligatorio)
    corrispondenze = [
        file for file in file_trovati
        if str(file.get("name", "")).casefold() == nome_file.casefold()
    ]
    if not corrispondenze:
        if obbligatorio:
            raise FileNotFoundError(f"{nome_file} non trovato in SharePoint: {cartella}")
        return None
    if len(corrispondenze) > 1:
        raise ValueError(f"Più file chiamati {nome_file!r} in {cartella}")
    client, drive_id = _drive_sharepoint()
    file_remoto = corrispondenze[0]
    return {**file_remoto, "content": client.scarica_file(drive_id, file_remoto["id"])}


def carica_file_sharepoint(cartella, nome_file, contenuto, content_type):
    client, drive_id = _drive_sharepoint()
    return client.carica_bytes(
        drive_id, cartella, nome_file, contenuto, content_type=content_type
    )

#COLONNE FILE 
#-------------------------------------------------------------------------------------
COLONNE_CORRETTE_GARA   = ['Cliente (Ordine)', 'Partita IVA (Ordine)', 'ID Riga CRM', 'ID Pratica',
       'Stato Ordine', 'Prodotto esistente', 'Tipologia Servizio (Prodotto esistente)',
       'Quantità', 'Prezzo unitario', 'Importo totale', 'Data creazione',
       'Data Inserimento OmniSales', 'Data Passaggio ACA', 'Data Attivazione',
       'Proprietario (Ordine)', 'ID ordine (Ordine)', 'Motivo stato (Ordine)',
       'Valuta', 'scala sconti']

COLONNE_OUTPUT_GARA = [col for col in COLONNE_CORRETTE_GARA if col != "ID Riga CRM"]

COLONNE_CORRETTE_STORICO = ['Cliente (Ordine)', 'Partita IVA (Ordine)', 'ID Pratica',
       'Prodotto esistente', 'Tipologia Servizio (Prodotto esistente)',
       'Quantità', 'Prezzo unitario', 'Importo totale', 'Data creazione',
       'Data Inserimento OmniSales', 'Data Passaggio ACA', 'Data Attivazione',
       'Proprietario (Ordine)', 'ID ordine (Ordine)', 'Motivo stato (Ordine)',
       'Valuta', 'scala sconti']

#FUNZIONI
#-------------------------------------------------------------------------------------
#trova file da elaborare
def trova_file(path_cartella, parametro):
    file_trovati = elenca_file_sharepoint(path_cartella, crea=True)
    parola = "inflow" if parametro == "Gara" else "storico gara"
    candidati = [
        file for file in file_trovati
        if parola in str(file.get("name", "")).casefold()
    ]
    if not candidati:
        print("["+Fore.RED+"ERRORE"+Fore.RESET+f"] Nessun file che rispetti il parametro: {parametro}")
        sys.exit(1)
    candidati.sort(
        key=lambda file: str(file.get("lastModifiedDateTime", "")), reverse=True
    )
    client, drive_id = _drive_sharepoint()
    scelto = candidati[0]
    return {**scelto, "content": client.scarica_file(drive_id, scelto["id"])}

#formatta file excel
def _formatta_excel(path):
    '''
    Funzione che serve per dare un pò di colore al file finale, e per dare il giusto formato 
    alle date ai prezzi.
    Si usano i set al posto delle liste perchè sono più veloci ed evitano di dover ciclare tutto il file 
    per trovare le colonne su cui applicare la formattazione.
    '''
    col_prezzi  = set()
    col_date    = set()
    wb = load_workbook(path)
    ws = wb.active

    # Colori
    fill_blu            = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    fill_giallo         = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    fill_rosso          = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    font_bianco         = Font(color="FFFFFF", bold=True)
    font_nero           = Font(color="000000", bold=True)
    font_verde_testo    = Font(color="2FBF47", bold = True)
    font_arancione_testo= Font(color="EA6B14", bold = True)
    font_rosso_testo    = Font(color="FF0000", bold = True)
    font_azzurro_testo  = Font(color="00B0F0", bold = True)
    
    for cell in ws[1]:  # prima riga = intestazioni
        col_name = str(cell.value).strip().lower() if cell.value else ""

        if "proprietario" in col_name:
            cell.fill = fill_giallo
            cell.font = font_nero
        elif "scala sconti" in col_name or "sconto" in col_name:
            cell.fill = fill_rosso
            cell.font = font_bianco
        else:
            cell.fill = fill_blu
            cell.font = font_bianco
        #crea lista di colonne data e prezzo/importo
        if "prezzo" in col_name or "importo" in col_name:
            col_prezzi.add(cell.column) 
        
        elif "data" in col_name:
            col_date.add(cell.column)

    #Formatta Date, Prezzi e colori dell stato ordine
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.column in col_date:
                cell.number_format = "DD/MM/YYYY"
            elif cell.column in col_prezzi:
                cell.number_format = "#,##0.00 €"
            elif str(cell.value).strip().lower() == "attivato":
                cell.font = font_verde_testo
            elif str(cell.value).strip().lower() == "da inserire":   
                cell.font = font_azzurro_testo
            elif str(cell.value).strip().lower() == "in attivazione":
                cell.font = font_arancione_testo
            elif str(cell.value).strip().lower() == "annullato":
                cell.font = font_rosso_testo
    if isinstance(path, BytesIO):
        output = BytesIO()
        wb.save(output)
        path.seek(0)
        path.truncate()
        path.write(output.getvalue())
        path.seek(0)
    else:
        wb.save(path)

# Carica i file nelle cartelle SharePoint ".parquet" e "BASI_DATI_xlsx".
def carica_su_teams(df, *args) -> None:
    '''
    Carica tramite Graph sia i file della gara corrente sia quelli dello storico,
    riconoscendo la destinazione SharePoint richiesta.
    '''
    try:
        if not isinstance(df, pd.DataFrame):
            raise TypeError(
                f"Atteso pandas.DataFrame, ricevuto {type(df).__name__}"
            )   
        
        for path in args:
            if path == PATH_PARQUET:
                nome_file, formato = "Gara (Inflow).parquet", "parquet"
            elif path == PATH_BASI_DATI:
                nome_file, formato = "Gara (Inflow).xlsx", "excel"
            elif path == PATH_STORICO_PARQUET:
                nome_file, formato = "Storico Gara.parquet", "parquet"
            elif path == PATH_STORICO_BASI_DATI:
                nome_file, formato = "Storico Gara.xlsx", "excel"
            else:
                raise ValueError(f"Destinazione SharePoint Gara non riconosciuta: {path}")

            buffer = BytesIO()
            if formato == "parquet":
                df.to_parquet(buffer, index=False)
                content_type = "application/octet-stream"
            else:
                df.to_excel(buffer, index=False, engine="openpyxl")
                buffer.seek(0)
                _formatta_excel(buffer)
                content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            caricato = carica_file_sharepoint(
                path, nome_file, buffer.getvalue(), content_type
            )
            print(Fore.GREEN + f"SharePoint: {caricato.get('webUrl', nome_file)}")

        print(Fore.GREEN + "Caricamento su SharePoint avvenuto con successo\n")
    except Exception as e:
        print("["+Fore.RED+"ERRORE"+Fore.RESET+f"] Caricamento file su SharePoint fallito --> {e}")
        raise

#CONVERSIONE API

STAGE_MAP = {
    3: "Parzialmente evaso",
    4: "annullato",
    6: "evaso",
    1: " NON INDICATO "
}

LISTINO_MAP = {
    "1": "Listino",
    "2": "Listino Standard",
    "3": "Inattivo",
    "4": "Listino Campagne",
    "5": "Listino Preventiva",
    "6": "Listino Maverick",
    "7": "Listino IT Farm"
}

TIPOLOGIA_SERVIZIO_MAP = {
    "1" : "ALTRO",
    "2" : "SIM VOCE",
    "3" : "SIM DATI",
    "4" : "DSL",
    "5" : "ON TOP DSL",
    "6" : "VRU INTERNO",
    "7" : "LINK SA",
    "8" : "LINK SU",
    "9" : "TW SA",
    "10" : "ON TOP LINK",
    "11" : "M2M",
    "12" : "SAAS",
    "13" : "Collaboration",
    "14" : "IAAS",
    "15" : "Security",
    "16" : "COMPLEX DEAL",
    "17" : "Custom App",
    "18" : "EASYDEAL",
    "19" : "EOLO",
    "20" : "RATA TEL",
    "21" : "SOL. TEL.",
    "22" : "RATA TEL VOCE",
    "23" : "SOL TEL VOCE",
    "24" : "ON TOP EASY DEAL",
    "25" : "RATA TEL DATI",
    "26" : "SOL TEL DATI",
    "27" : "NOLEGGIO OP.",
    "28" : "ENERGIA",
    None : "NON SEGNATA"

}

TIPO_RELAZIONE_MAP = {
    "01" : "Potenziale",
    "02" : "CB",
    "03" : "DB Finanziario",
    "04" : "Cliente",
    None : "NON SEGNATA"
}


PERIODO_GARE = {
    1: "Gen - Mar",
    2: "Gen - Mar",
    3: "Gen - Mar",

    4: "Apr - Giu",
    5: "Apr - Giu",
    6: "Apr - Giu",

    7: "Lug - Set",
    8: "Lug - Set",
    9: "Lug - Set",

    10: "Ott - Dic",
    11: "Ott - Dic",
    12: "Ott - Dic",
}

def trova_periodo_gara():
    oggi = datetime.today()
    anno = oggi.year
    gara_corrente = PERIODO_GARE[oggi.month]

    if gara_corrente == "Gen - Mar":
        inizio_gara = datetime(anno, 1, 1).date()
        fine_gara = datetime(anno, 3, 31).date()

    elif gara_corrente == "Apr - Giu":
        inizio_gara = datetime(anno, 4, 1).date()
        fine_gara = datetime(anno, 6, 30).date()

    elif gara_corrente == "Lug - Set":
        inizio_gara = datetime(anno, 7, 1).date()
        fine_gara = datetime(anno, 9, 30).date()

    elif gara_corrente == "Ott - Dic":
        inizio_gara = datetime(anno, 10, 1).date()
        fine_gara = datetime(anno, 12, 31).date()

    return inizio_gara, fine_gara, gara_corrente
