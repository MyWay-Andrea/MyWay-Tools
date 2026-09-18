import pandas as pd
import numpy as np
from openpyxl import load_workbook
import shutil

from config import (
    COLONNE, SERVIZI_DA_SCONTARE, COLONNA_AGENTE_GARA,
    MASTER_GARA as MASTER, SHEET_GARA,
    DATE_COLS_NAME as DATE_COLS,
    copia_formattaz_cella, copia_template,
    trova_raw_file, PATH_RAW_FILE, DEST,
    PERCORSO_TEAMS, copia_su_teams, svuota_cartella
)
from colorama import Fore, Style, init
init(autoreset=True)


COLONNE_CORRETTE = [
    'Cliente (Ordine)', 'Partita IVA (Ordine)', 'ID Pratica',
    'Prodotto esistente', 'Tipologia Servizio (Prodotto esistente)',
    'Quantità', 'Prezzo unitario', 'Importo totale', 'Data creazione',
    'Data Inserimento OmniSales', 'Data Passaggio ACA', 'Data Attivazione',
    'Proprietario (Ordine)', 'ID ordine (Ordine)', 'Motivo stato (Ordine)',
    'Valuta', 'Codice Aggregato', 'Famiglia', 'Squadra Energia',
    'scala sconti'
]


def forza_formato_dati(df):
    df.iloc[:, 5:8] = df.iloc[:, 5:8].apply(pd.to_numeric, errors='coerce')
    for col in DATE_COLS:
        df[col] = pd.to_datetime(df[col], errors='coerce').dt.strftime('%d/%m/%Y')
    df = df.rename(columns={df.columns[12]: COLONNA_AGENTE_GARA})
    df[COLONNA_AGENTE_GARA] = df[COLONNA_AGENTE_GARA].astype(str).str.strip()
    return df


def calcola_prezzo(df):
    try:
        df["Importo Calcolato"] = df.iloc[:, 7].copy()
        for idx, servizio in df.iloc[:, 4].items():
            if servizio in SERVIZI_DA_SCONTARE:
                df.loc[idx, "Importo Calcolato"] = (
                    df.loc[idx, df.columns[7]]
                    - (df.loc[idx, df.columns[7]] * (df.loc[idx, df.columns[16]] / 100))
                )
        return df
    except Exception as e:
        print(f"  ⚠️ Errore nel calcolo dell'importo calcolato: {e}")
        return df


def aggiungi_colonne_vuote(subset):
    colonne = subset.columns.tolist()
    nome_colonna_16 = colonne[16]
    scala_sconti = subset.pop(nome_colonna_16)
    subset[16] = np.nan
    subset[17] = np.nan
    subset[18] = np.nan
    subset[19] = scala_sconti
    return subset


def forza_formato_excel(ws):
    for row in range(2, ws.max_row + 1):
        for col in [9, 10, 11, 12]:
            ws.cell(row=row, column=col).number_format = 'DD/MM/YYYY'
        for col in [7, 8]:
            ws.cell(row=row, column=col).number_format = '#,##0.00 €'


def popola_template(file_gara, df, sheet):
    wb = load_workbook(file_gara)
    ws = wb[sheet]
    template_stili = {
        idx_col: ws.cell(row=2, column=idx_col)
        for idx_col in range(1, len(df.columns) + 2)
    }
    for idx_riga, row_data in enumerate(df.values, start=2):
        for idx_col, valore in enumerate(row_data, start=1):
            cella = ws.cell(row=idx_riga, column=idx_col)
            copia_formattaz_cella(template_stili[idx_col], cella)
            cella.value = valore
    forza_formato_excel(ws)
    wb.save(file_gara)
    wb.close()


def main(file_gara_raw):
    file_gara = None
    try:
        df = pd.read_excel(
            file_gara_raw,
            sheet_name=0,
            parse_dates=DATE_COLS,
            dtype={'Ordini di vendita P.IVA': str, 'Items ID Pratica': str},
            date_format='%d/%m/%Y'
        )

        df = df.dropna(how='all')
        df = calcola_prezzo(df)         # 1. calcola importo (usa colonna 16)
        df = forza_formato_dati(df)     # 2. normalizza tipi e nomi
        subset = df.iloc[:, COLONNE]    # 3. seleziona colonne (dopo il calcolo)
        subset = aggiungi_colonne_vuote(subset)
        subset.columns = COLONNE_CORRETTE

        file_gara = copia_template(MASTER, file=file_gara_raw)
        popola_template(file_gara, subset, SHEET_GARA)

        #copia su teams "file_gara_completo"
        copia_su_teams(file_gara, PERCORSO_TEAMS)

        #copia il file in formato '.parquet'
        copia_su_teams(subset)

        #elimina file dalla cartella una volta caricato sul teams 
        svuota_cartella(PATH_RAW_FILE, nome_file = "Gara")

        print(Fore.GREEN + "\n ✔ FORMATTAZIONE GARA ESEGUITA CON " + Style.BRIGHT + "SUCCESSO")

    except Exception as e:
        print(f"Si è verificato un errore --> {e}")
        if file_gara is not None:
            backup = DEST / "GARA_PULITA.xlsx"
            shutil.copy2(file_gara, backup)
            print(f"  💾 Backup salvato in: {backup}")


if __name__ == "__main__":
    file_gara = trova_raw_file(PATH_RAW_FILE)
    main(file_gara)