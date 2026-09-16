from openpyxl import load_workbook
import xlwings as xw
from copy import copy
from colorama import Fore, init
init(autoreset = True)

def copia_stile_cella(cella_origine, cella_destinazione):
    if cella_origine.has_style:
        cella_destinazione.font = copy(cella_origine.font)
        cella_destinazione.border = copy(cella_origine.border)
        cella_destinazione.fill = copy(cella_origine.fill)
        cella_destinazione.number_format = copy(cella_origine.number_format)
        cella_destinazione.protection = copy(cella_origine.protection)
        cella_destinazione.alignment = copy(cella_origine.alignment)


def popola_folgi_df(ws, df):
    n_colonne = len(df.columns)
    
    template_stili = {idx_col: ws.cell(row=2, column=idx_col) 
                      for idx_col in range(1, n_colonne + 1)}

    for idx_riga, row_data in enumerate(df.values, start = 2):
        for idx_col, valore in enumerate(row_data, start = 1):
            cella_destinazione = ws.cell(row = idx_riga, column = idx_col)
            copia_stile_cella(template_stili[idx_col], cella_destinazione)
            cella_destinazione.value = valore

def aggiorna_pivot(path_report):
    app = xw.App(visible=False)
    app.display_alerts = False
    try:
        wb = xw.Book(path_report)
        wb.api.RefreshAll()

        #ridimensiona le colonne per non vedere ####### al posto dei valori
        ws_pivot = wb.sheets["PIVOT"]
        ws_pivot.api.Columns.AutoFit()

        wb.save()
        wb.close()
    finally:
        app.quit()

def forza_formato(ws):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
        for cella in row:
            cella.number_format = "DD/MM/YYYY"

def inserisci_dati_report(path_report, REPORT_INFO, agente = None):
    wb = load_workbook(path_report)
    
    for nome_foglio, dati in REPORT_INFO.items():
        ws = wb[nome_foglio]
        popola_folgi_df(ws, dati)
        forza_formato(ws)

    if agente == "G.MANDELLI":
        print(Fore.YELLOW + f"\n ⌛ Inizo elaborazione {agente}...\n")
        ws = wb["PIVOT"]
        try:
            ws.sheet_state = "visible"
            wb.save(path_report)
            wb.close()
            aggiorna_pivot(path_report)
            print(Fore.GREEN + "   Foglio pivot scoperto con successo")
            
        except Exception as e:
            print(Fore.RED + f" ❌ ERRORE nella compilazione del file: {path_report.name}. --> {type(e).__name__}: {e}")
    else:
        wb.save(path_report)
        wb.close()
    
    return path_report

def main(path_report, df_app, df_corr, agente = None):

    # 1. diz con foglio, df
    REPORT_INFO = {
        "APPUNTAMENTI" : df_app,
        "SETTIMANA CORRENTE" : df_corr
    }  

    # 2. scrivi dati nel fogli del report 
    try:
        if agente == None:
            report = inserisci_dati_report(path_report, REPORT_INFO)
        else:
            report = inserisci_dati_report(path_report, REPORT_INFO, agente)

        print(Fore.GREEN + f"\n ✅ {path_report.name} è stato scritto correttamente")
    except Exception as e:
        print(Fore.RED + f" ❌ ATTENZIONE il report {path_report.name} ha causato un errore --> {type(e).__name__}: {e}\ncontinuo...")

    return report
    