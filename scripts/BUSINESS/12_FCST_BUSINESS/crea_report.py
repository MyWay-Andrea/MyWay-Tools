from elabora_fcst import main as elabora_fcst

from config import(
    trova_file,
    PATH_ONEDRIVE_FILE,
    PATH_OUTPUT_FCST,
    MESI_IT,
    trova_periodo_gara

)
import pandas as pd
import re
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from colorama import Fore, init
init(autoreset=True)

def trova_colonna_stato(df: pd.DataFrame) -> str:
    for col in df.columns:
        if "stato" in col and "ordine" in col:
            return col
    
    for col in df.columns:
        if "stato" in col:
            return col
    
    raise ValueError("Colonna stato ordine non trovata nel df gara")

def trova_colonna(df: pd.DataFrame, *parole_chiave: str) -> str:
    for col in df.columns:
        nome_colonna = str(col).lower()
        # all controlla che tutte ele condizioni siano verificate
        if all(parola in nome_colonna for parola in parole_chiave):
            return col

    return ""

def normalizza_serie_testo(serie: pd.Series) -> pd.Series:
    return serie.astype(str).str.strip().str.lower()

def filtra_per_venditore(df: pd.DataFrame, col_venditore: str, venditore: str) -> pd.DataFrame:
    if col_venditore not in df.columns:
        return df.iloc[0:0].copy()

    return df[
        df[col_venditore]
        .astype(str)
        .str.strip()
        .str.upper()
        .eq(venditore.upper())
    ].copy()

def filtra_contiene(df: pd.DataFrame, colonna: str, testo: str) -> pd.DataFrame:
    if colonna not in df.columns:
        return df.iloc[0:0].copy()

    return df[normalizza_serie_testo(df[colonna]).eq(testo.lower())].copy()

def filtra_per_stato(df: pd.DataFrame, col_stato: str, stato: str) -> pd.DataFrame:
    return df[
        df[col_stato]
        .astype(str)
        .str.strip()
        .str.lower()
        .eq(stato.lower())
    ].copy()

def pulisci_nome_file(nome: str) -> str:
    nome_pulito = re.sub(r'[<>:"/\\|?*]', "_", str(nome).strip())
    nome_pulito = re.sub(r"\s+", " ", nome_pulito)
    return nome_pulito or "SENZA_NOME"

def somma_per_mese(df: pd.DataFrame, col_data: str, col_valore: str, mesi: list[int]) -> list[float]:
    valori = []

    if col_data not in df.columns or col_valore not in df.columns:
        return [0 for _ in mesi]

    df_calc = df.copy()
    df_calc[col_data] = pd.to_datetime(df_calc[col_data], errors="coerce")
    df_calc[col_valore] = pd.to_numeric(df_calc[col_valore], errors="coerce").fillna(0)

    for mese in mesi:
        valori.append(df_calc.loc[df_calc[col_data].dt.month == mese, col_valore].sum())

    return valori

def conta_per_mese(df: pd.DataFrame, col_data: str, mesi: list[int]) -> list[int]:
    if col_data not in df.columns:
        return [0 for _ in mesi]

    date = pd.to_datetime(df[col_data], errors="coerce")
    return [int((date.dt.month == mese).sum()) for mese in mesi]

def crea_pivot_fcst(df_gara: pd.DataFrame, df_oppo: pd.DataFrame, df_appu: pd.DataFrame) -> pd.DataFrame:
    _, mesi, gara_corrente = trova_periodo_gara()
    col_venditore_gara = trova_colonna(df_gara, "proprietario", "ordine")
    col_venditore_oppo = trova_colonna(df_oppo, "proprietario", "opportun")
    col_venditore_appu = trova_colonna(df_appu, "in carico")
    col_stato_gara = trova_colonna_stato(df_gara)
    col_stato_oppo = trova_colonna(df_oppo, "stato", "opportun")
    col_data_attivazione = trova_colonna(df_gara, "data", "attivazione")
    col_data_opp = trova_colonna(df_oppo, "data", "chiusura")
    col_data_app = trova_colonna(df_appu, "data")
    col_importo_gara = trova_colonna(df_gara, "importo", "totale")
    col_importo_oppo = trova_colonna(df_oppo, "importo", "totale")

    venditori = venditori_presenti(
        (df_gara, col_venditore_gara),
        (df_oppo, col_venditore_oppo),
        (df_appu, col_venditore_appu),
    )

    righe = []
    for venditore in venditori:
        ordini_venditore = filtra_per_venditore(df_gara, col_venditore_gara, venditore)
        ordini_attivi = filtra_per_stato(ordini_venditore, col_stato_gara, "Attivato")
        opp_venditore = filtra_per_venditore(df_oppo, col_venditore_oppo, venditore)
        opp_aperte = filtra_contiene(opp_venditore, col_stato_oppo, "Aperta")
        app_venditore = filtra_per_venditore(df_appu, col_venditore_appu, venditore)

        riga = {"Venditore": venditore, "Periodo": gara_corrente}

        for mese, valore in zip(mesi, somma_per_mese(ordini_attivi, col_data_attivazione, col_importo_gara, mesi)):
            riga[f"Vendite {MESI_IT[mese]}"] = valore

        for mese, valore in zip(mesi, somma_per_mese(opp_aperte, col_data_opp, col_importo_oppo, mesi)):
            riga[f"Opp {MESI_IT[mese]}"] = valore

        for mese, valore in zip(mesi, conta_per_mese(app_venditore, col_data_app, mesi)):
            riga[f"App {MESI_IT[mese]}"] = valore

        righe.append(riga)

    return pd.DataFrame(righe)

def crea_fogli_report(df_gara: pd.DataFrame, df_oppo: pd.DataFrame, df_appu: pd.DataFrame) -> dict:
    col_stato = trova_colonna_stato(df_gara)

    return {
        "FCST": crea_pivot_fcst(df_gara, df_oppo, df_appu),
        "ordini attivi": filtra_per_stato(df_gara, col_stato, "Attivato"),
        "in attivazione": filtra_per_stato(df_gara, col_stato, "In Attivazione"),
        "da inserire": filtra_per_stato(df_gara, col_stato, "Da Inserire"),
        "annullati": filtra_per_stato(df_gara, col_stato, "Annullato"),
        "Opportunita": df_oppo,
        "Appuntamenti": df_appu,
    }

def formatta_excel(path_file) -> None:
    wb = load_workbook(path_file)

    fill_blu = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    fill_giallo = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    fill_rosso = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    font_bianco = Font(color="FFFFFF", bold=True)
    font_nero = Font(color="000000", bold=True)
    font_verde = Font(color="2FBF47", bold=True)
    font_arancione = Font(color="EA6B14", bold=True)
    font_rosso = Font(color="FF0000", bold=True)
    font_azzurro = Font(color="00B0F0", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        col_prezzi = set()
        col_date = set()
        col_prob = set()

        for cell in ws[1]:
            col_name = str(cell.value).strip().lower() if cell.value else ""
            cell.alignment = Alignment(horizontal="center")

            if "proprietario" in col_name or "venditore" in col_name or "in carico" in col_name:
                cell.fill = fill_giallo
                cell.font = font_nero
            elif "scala sconti" in col_name or "sconto" in col_name:
                cell.fill = fill_rosso
                cell.font = font_bianco
            else:
                cell.fill = fill_blu
                cell.font = font_bianco

            if "prezzo" in col_name or "importo" in col_name or "vendite" in col_name or col_name.startswith("opp "):
                col_prezzi.add(cell.column)
            elif "data" in col_name:
                col_date.add(cell.column)
            elif "probabilit" in col_name:
                col_prob.add(cell.column)

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                valore = str(cell.value).strip().lower()
                if cell.column in col_date:
                    cell.number_format = "DD/MM/YYYY"
                elif cell.column in col_prezzi:
                    cell.number_format = '#,##0.00 €'
                elif cell.column in col_prob:
                    cell.number_format = r"0\%"
                elif valore in ["attivato", "aperta", "attività fatta"]:
                    cell.font = font_verde
                elif valore in ["da inserire", "acquisita", "attività da fare"]:
                    cell.font = font_azzurro
                elif valore == "in attivazione":
                    cell.font = font_arancione
                elif valore in ["annullato", "persa", "annullata"]:
                    cell.font = font_rosso

        for col_cells in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 12), 45)

    wb.save(path_file)

def scrivi_report_excel(path_file, fogli: dict) -> None:
    with pd.ExcelWriter(path_file, engine="openpyxl") as writer:
        for nome_foglio, df in fogli.items():
            df.to_excel(writer, sheet_name=nome_foglio, index=False)
    formatta_excel(path_file)

def venditori_presenti(*dataframes: tuple[pd.DataFrame, str]) -> list[str]:
    venditori = set()

    for df, colonna in dataframes:
        if colonna not in df.columns:
            continue

        valori = (
            df[colonna]
            .dropna()
            .astype(str)
            .str.strip()
        )
        venditori.update(valore for valore in valori if valore and valore.lower() != "nan")

    return sorted(venditori, key=str.upper)

def main(df_gara, df_oppo, df_appu):
    print("\n === Creazione file FCST ===")

    PATH_OUTPUT_FCST.mkdir(parents=True, exist_ok=True)

    col_venditore_gara = trova_colonna(df_gara, "proprietario", "ordine")
    col_venditore_oppo = trova_colonna(df_oppo, "proprietario", "opportun")
    col_venditore_appu = trova_colonna(df_appu, "in carico")

    venditori = venditori_presenti(
        (df_gara, col_venditore_gara),
        (df_oppo, col_venditore_oppo),
        (df_appu, col_venditore_appu),
    )

    file_creati = []

    for venditore in venditori:
        df_gara_venditore = filtra_per_venditore(df_gara, col_venditore_gara, venditore)
        df_oppo_venditore = filtra_per_venditore(df_oppo, col_venditore_oppo, venditore)
        df_appu_venditore = filtra_per_venditore(df_appu, col_venditore_appu, venditore)

        path_file = PATH_OUTPUT_FCST / f"FCST_{pulisci_nome_file(venditore)}.xlsx"
        scrivi_report_excel(
            path_file,
            crea_fogli_report(df_gara_venditore, df_oppo_venditore, df_appu_venditore)
        )
        file_creati.append(path_file)
        print(Fore.GREEN + "✓ " + Fore.RESET + f"{path_file.name}")

    path_governance = PATH_OUTPUT_FCST / "FCST_GOVERNANCE.xlsx"
    scrivi_report_excel(
        path_governance,
        crea_fogli_report(df_gara, df_oppo, df_appu)
    )
    file_creati.append(path_governance)

    print(Fore.GREEN + "✓ " + Fore.RESET + f"{path_governance.name}")
    print(f"\nReport creati: {len(file_creati)}")
    print(f"Cartella output: {PATH_OUTPUT_FCST}")

    return file_creati


if __name__ == "__main__":

    # 1. importa file 
    print("\n === Ricerca file per file FCST ===")
    file_gara = trova_file(PATH_ONEDRIVE_FILE, "gara")
    file_opportunita = trova_file(PATH_ONEDRIVE_FILE, "opportunita")
    file_appuntamenti = trova_file(PATH_ONEDRIVE_FILE, "appuntamenti")

    
    # 2. chiama funzione 
    df_gara, df_oppo, df_appu = elabora_fcst(file_gara, file_opportunita, file_appuntamenti)

    # 3. crea report FCST
    main(df_gara, df_oppo, df_appu)
