from datetime import datetime
from io import BytesIO

import pandas as pd
from colorama import Fore, init
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from config import (
    PATH_INPUT_PARQUET,
    PATH_OUTPUT_FCST,
    MESI_IT,
    carica_report_venditore,
    trova_file,
    trova_periodo_gara,
)
from elabora_fcst import main as elabora_fcst_main
from crea_report import (
    conta_per_mese,
    crea_fogli_report,
    filtra_contiene,
    filtra_per_stato,
    filtra_per_venditore,
    pulisci_nome_file,
    scrivi_report_excel,
    somma_per_mese,
    trova_colonna,
    trova_colonna_stato,
    venditori_presenti,
)

init(autoreset=True)


def abbrevia_mese(mese: int) -> str:
    return MESI_IT[mese][:3].upper()


def valore_somma(valori: list[float]) -> float:
    return float(sum(valori))


def somma_colonna(df: pd.DataFrame, colonna: str) -> float:
    if colonna not in df.columns:
        return 0.0

    return float(pd.to_numeric(df[colonna], errors="coerce").fillna(0).sum())


def lettera_colonna(indice: int) -> str:
    lettere = ""
    while indice:
        indice, resto = divmod(indice - 1, 26)
        lettere = chr(65 + resto) + lettere
    return lettere


def aggiorna_bordo(cell, left=None, right=None, top=None, bottom=None) -> None:
    cell.border = Border(
        left=left or cell.border.left,
        right=right or cell.border.right,
        top=top or cell.border.top,
        bottom=bottom or cell.border.bottom,
    )


def filtra_per_venditori_presenti(df: pd.DataFrame, col_venditore: str, venditori: list[str]) -> pd.DataFrame:
    if col_venditore not in df.columns:
        return df.iloc[0:0].copy()

    venditori_normalizzati = {str(venditore).strip().upper() for venditore in venditori}
    return df[
        df[col_venditore]
        .astype(str)
        .str.strip()
        .str.upper()
        .isin(venditori_normalizzati)
    ].copy()


def venditore_crm(venditore: str) -> bool:
    return str(venditore).strip().upper().startswith("CRM")


def ordina_venditori_crm_ultimo(venditori: list[str]) -> list[str]:
    return sorted(venditori, key=lambda venditore: (venditore_crm(venditore), str(venditore).upper()))


def crea_pivot_fcst_aggregato(
    df_gara: pd.DataFrame,
    df_oppo: pd.DataFrame,
    df_appu: pd.DataFrame,
    includi_totale: bool = True,
) -> pd.DataFrame:
    _, mesi, periodo = trova_periodo_gara()
    label_mesi = [abbrevia_mese(mese) for mese in mesi]

    col_venditore_gara = trova_colonna(df_gara, "proprietario", "ordine")
    col_venditore_oppo = trova_colonna(df_oppo, "proprietario", "opportun")
    col_venditore_appu = trova_colonna(df_appu, "in carico")
    col_stato_gara = trova_colonna_stato(df_gara)
    col_stato_oppo = trova_colonna(df_oppo, "stato", "opportun")
    col_data_gara = trova_colonna(df_gara, "data", "attivazione")
    col_data_oppo = trova_colonna(df_oppo, "data", "chiusura")
    col_data_appu = trova_colonna(df_appu, "data")
    col_importo_gara = trova_colonna(df_gara, "importo", "totale")
    col_importo_oppo = trova_colonna(df_oppo, "importo", "totale")

    venditori = ordina_venditori_crm_ultimo(
        venditori_presenti(
            (df_gara, col_venditore_gara),
            (df_oppo, col_venditore_oppo),
            (df_appu, col_venditore_appu),
        )
    )

    righe = []

    if includi_totale:
        righe.append(
            crea_riga_fcst(
                "TOTALE",
                periodo,
                mesi,
                label_mesi,
                df_gara,
                df_oppo,
                df_appu,
                col_stato_gara,
                col_stato_oppo,
                col_data_gara,
                col_data_oppo,
                col_data_appu,
                col_importo_gara,
                col_importo_oppo,
            )
        )

    for venditore in venditori:
        righe.append(
            crea_riga_fcst(
                venditore,
                periodo,
                mesi,
                label_mesi,
                filtra_per_venditore(df_gara, col_venditore_gara, venditore),
                filtra_per_venditore(df_oppo, col_venditore_oppo, venditore),
                filtra_per_venditore(df_appu, col_venditore_appu, venditore),
                col_stato_gara,
                col_stato_oppo,
                col_data_gara,
                col_data_oppo,
                col_data_appu,
                col_importo_gara,
                col_importo_oppo,
            )
        )

    return pd.DataFrame(righe)


def crea_riga_fcst(
    venditore: str,
    periodo: str,
    mesi: list[int],
    label_mesi: list[str],
    df_gara: pd.DataFrame,
    df_oppo: pd.DataFrame,
    df_appu: pd.DataFrame,
    col_stato_gara: str,
    col_stato_oppo: str,
    col_data_gara: str,
    col_data_oppo: str,
    col_data_appu: str,
    col_importo_gara: str,
    col_importo_oppo: str,
) -> dict:
    ordini_attivi = filtra_per_stato(df_gara, col_stato_gara, "Attivato")
    ordini_da_inserire = filtra_per_stato(df_gara, col_stato_gara, "Da Inserire")
    ordini_in_attivazione = filtra_per_stato(df_gara, col_stato_gara, "In Attivazione")
    opportunita_aperte = filtra_contiene(df_oppo, col_stato_oppo, "Aperta")

    valori_attivi = somma_per_mese(ordini_attivi, col_data_gara, col_importo_gara, mesi)
    da_inserire = somma_colonna(ordini_da_inserire, col_importo_gara)
    in_attivazione = somma_colonna(ordini_in_attivazione, col_importo_gara)
    valore_opportunita = valore_somma(
        somma_per_mese(opportunita_aperte, col_data_oppo, col_importo_oppo, mesi)
    )
    appuntamenti = conta_per_mese(df_appu, col_data_appu, mesi)

    riga = {
        "Venditore": venditore,
        "Periodo": periodo,
    }

    for label_mese, valore in zip(label_mesi, valori_attivi):
        riga[f"{label_mese} Attivo"] = valore

    riga[f"{'+'.join(label_mesi)} ATTIVO"] = valore_somma(valori_attivi)
    riga["o/w da inserire"] = da_inserire
    riga["o/w in attivazione"] = in_attivazione
    riga["TOT"] = da_inserire + in_attivazione
    riga["TOT OPPS CRM"] = valore_opportunita

    for label_mese, valore in zip(label_mesi, appuntamenti):
        riga[label_mese] = valore

    riga["TRIMESTRE"] = int(sum(appuntamenti))
    return riga


def formatta_fcst(path_file) -> None:
    wb = load_workbook(path_file)
    ws = wb["FCST"]

    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))

    oggi = datetime.today()
    _, mesi, _ = trova_periodo_gara()
    mesi_label = [abbrevia_mese(mese) for mese in mesi]
    mese_corrente = MESI_IT[oggi.month].upper()

    header_row_1 = 5
    header_row_2 = 6
    first_data_row = 7
    first_col = 2
    last_col = 14

    gia_formattato = str(ws["B5"].value).strip().upper() == "GARA"
    source_row = first_data_row if gia_formattato else 1
    source_col = first_col if gia_formattato else 1
    source_last_col = source_col + (12 if gia_formattato else 13)

    rows = []
    for row_idx in range(source_row, ws.max_row + 1):
        values = [ws.cell(row=row_idx, column=col_idx).value for col_idx in range(source_col, source_last_col + 1)]
        if any(value is not None for value in values):
            rows.append(values if gia_formattato else values[:1] + values[2:])

    data_rows = rows if gia_formattato else rows[1:]
    last_row = first_data_row + len(data_rows) - 1

    max_row_to_clear = max(ws.max_row, last_row + 2, 10)
    for row in ws.iter_rows(min_row=1, max_row=max_row_to_clear, min_col=1, max_col=20):
        for cell in row:
            cell.value = None
            cell.border = Border()
            cell.fill = PatternFill(fill_type=None)
            cell.font = Font(name="Calibri", color="000000")
            cell.alignment = Alignment()

    fill_verde = PatternFill(start_color="2E7031", end_color="2E7031", fill_type="solid")
    fill_bianco = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_rosso = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    fill_azzurro = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
    lato_thick = Side(style="thick", color="000000")
    lato_thin = Side(style="thin", color="000000")
    lato_rosso = Side(style="thick", color="FF0000")

    def border(left=lato_thin, right=lato_thin, top=lato_thin, bottom=lato_thin):
        return Border(left=left, right=right, top=top, bottom=bottom)

    def style_range(min_row, max_row, min_col, max_col, fill=None, font=None, border_style=None):
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if fill is not None:
                    cell.fill = fill
                if font is not None:
                    cell.font = font
                if border_style is not None:
                    cell.border = border_style
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("B2:D3")
    ws["B2"] = oggi.strftime("%d/%m/%Y")
    ws["B2"].font = Font(name="Calibri", size=26, bold=True, color="000000")
    style_range(2, 3, 2, 4, border_style=border(lato_thick, lato_thick, lato_thick, lato_thick))

    ws["B5"] = "GARA"
    ws["B6"] = mese_corrente
    ws["C5"] = "ACTUAL"
    ws["G5"] = f"{mese_corrente}\n(NON ATTIVO)"
    ws["J5"] = "OPPORTUNITA APERTE"
    ws["K5"] = "APPUNTAMENTI A SISTEMA"

    ws.merge_cells("C5:F5")
    ws.merge_cells("G5:I5")
    ws.merge_cells("K5:N5")

    headers = [
        f"{mesi_label[0]} Attivo",
        f"{mesi_label[1]} Attivo",
        f"{mesi_label[2]} Attivo",
        f"{'+'.join(mesi_label)} ATTIVO",
        "o/w da inserire",
        "o/w in attivazione",
        "TOT",
        "TOT OPPS CRM",
        mesi_label[0],
        mesi_label[1],
        mesi_label[2],
        "TRIMESTRE",
    ]
    for col_idx, value in enumerate(headers, start=3):
        ws.cell(row=header_row_2, column=col_idx).value = value

    style_range(header_row_1, header_row_2, first_col, last_col, fill=fill_bianco, font=Font(name="Aharoni", bold=True), border_style=border(lato_thick, lato_thick, lato_thick, lato_thick))

    ws["B5"].font = Font(name="Aharoni", bold=True, size=20)
    ws["B6"].font = Font(name="Aharoni", bold=True, size=20)
    ws["C5"].font = Font(name="Aharoni", bold=True, size=20, color="FFFFFF")
    ws["C5"].fill = fill_verde
    for col_idx in range(3, 7):
        ws.cell(row=header_row_2, column=col_idx).fill = fill_bianco
        ws.cell(row=header_row_2, column=col_idx).font = Font(name="Aharoni", bold=True, size=16, color="2E7031")
    ws["F6"].font = Font(name="Aharoni", bold=True, size=12, color="2E7031")

    ws["G5"].font = Font(name="Aharoni", bold=True, size=14, color="000000")
    ws["G6"].fill = fill_rosso
    ws["G6"].font = Font(name="Aharoni", bold=True, color="FFFFFF")
    ws["H6"].fill = fill_azzurro
    ws["H6"].font = Font(name="Aharoni", bold=True, color="FFFFFF")
    ws["I6"].fill = fill_bianco
    ws["I6"].font = Font(name="Aharoni", bold=True, size=20, color="00B0F0")

    ws["J5"].fill = fill_bianco
    ws["J5"].font = Font(name="Aharoni", bold=True, size=14, color="000000")
    ws["J6"].fill = fill_bianco
    ws["J6"].font = Font(name="Aharoni", bold=True, size=18, color="000000")

    ws["K5"].font = Font(name="Aharoni", bold=True, size=20, color="FFFFFF")
    ws["K5"].fill = fill_verde
    for col_idx in range(11, 15):
        ws.cell(row=header_row_2, column=col_idx).fill = fill_bianco
        ws.cell(row=header_row_2, column=col_idx).font = Font(name="Aharoni", bold=True, size=16, color="2E7031")

    for row_idx, row_values in enumerate(data_rows, start=first_data_row):
        for col_idx, value in enumerate(row_values[:13], start=first_col):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.fill = fill_bianco
            cell.font = Font(name="Calibri", color="000000")
            cell.border = border()
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=first_data_row, max_row=max(last_row, first_data_row), min_col=3, max_col=10):
        for cell in row:
            cell.number_format = '#,##0.00 \u20ac'
    for row in ws.iter_rows(min_row=first_data_row, max_row=max(last_row, first_data_row), min_col=11, max_col=14):
        for cell in row:
            cell.number_format = "0"

    has_total = last_row >= first_data_row and str(ws.cell(first_data_row, first_col).value).strip().upper() == "TOTALE"
    if has_total:
        for cell in ws[first_data_row][first_col - 1:last_col]:
            cell.font = Font(name="Calibri", bold=True)

        ws["B7"].border = border(lato_thick, lato_thick, lato_thick, lato_thick)
        for col_idx in range(3, 7):
            ws.cell(first_data_row, col_idx).border = border(
                lato_thick if col_idx == 3 else lato_thin,
                lato_thick if col_idx == 6 else lato_thin,
                lato_thick,
                lato_thick,
            )
        for col_idx in range(7, 9):
            ws.cell(first_data_row, col_idx).border = border(lato_thin, lato_thin, lato_thin, lato_thick)
        ws["I7"].border = border(lato_thick, lato_rosso, lato_thick, lato_thick)
        ws["K7"].border = border(lato_rosso, lato_thin, lato_thick, lato_thick)
        ws["L7"].border = border(lato_thin, lato_thin, lato_thick, lato_thick)
        ws["M7"].border = border(lato_thin, lato_thin, lato_thick, lato_thick)
        ws["N7"].border = border(lato_thick, lato_thick, lato_thick, lato_thick)

    for row_idx in range(first_data_row + 1, last_row + 1):
        ws.cell(row_idx, 2).border = border(lato_thick, lato_thick, lato_thin, lato_thin)
        ws.cell(row_idx, 2).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx in range(3, 15):
            ws.cell(row_idx, col_idx).border = border()

    for row_idx in range(header_row_2, last_row + 1):
        ws.cell(row_idx, 6).font = Font(name="Aharoni" if row_idx == header_row_2 else "Calibri", color="2E7031" if row_idx == header_row_2 else "000000", bold=True, size=12 if row_idx == header_row_2 else None)
        ws.cell(row_idx, 6).border = border(lato_thick, lato_thick, lato_thick if row_idx in [header_row_2, first_data_row] else lato_thin, lato_thick if row_idx in [first_data_row, last_row] else lato_thin)

        ws.cell(row_idx, 9).font = Font(name="Aharoni" if row_idx == header_row_2 else "Calibri", color="00B0F0", bold=True, size=20 if row_idx == header_row_2 else None)
        ws.cell(row_idx, 9).border = border(lato_thick, lato_rosso, lato_thick if row_idx in [header_row_2, first_data_row] else lato_thin, lato_thick if row_idx in [first_data_row, last_row] else lato_thin)

    for row_idx in range(first_data_row + 1, last_row + 1):
        for col_idx in range(7, 10):
            ws.cell(row_idx, col_idx).font = Font(name="Calibri", color="00B0F0", bold=(col_idx == 9))
        ws.cell(row_idx, 10).border = border(lato_thin, lato_rosso, lato_thin, lato_rosso if row_idx == last_row else lato_thin)
        ws.cell(row_idx, 14).border = border(lato_thick, lato_thin, lato_thin, lato_thin)

    for row_idx in range(header_row_1, last_row + 1):
        aggiorna_bordo(ws.cell(row_idx, first_col), left=lato_thick)
        aggiorna_bordo(ws.cell(row_idx, last_col), right=lato_thick)
    for col_idx in range(first_col, last_col + 1):
        aggiorna_bordo(ws.cell(header_row_1, col_idx), top=lato_thick)
        if last_row >= first_data_row:
            aggiorna_bordo(ws.cell(last_row, col_idx), bottom=lato_thick)

    if last_row >= first_data_row:
        for row_idx in range(header_row_1, last_row + 1):
            ws.cell(row_idx, 10).border = border(
                lato_rosso,
                lato_rosso,
                lato_thick if row_idx in [header_row_1, first_data_row] else lato_thin,
                lato_thick if row_idx in [header_row_1, first_data_row] else (lato_rosso if row_idx == last_row else lato_thin),
            )
        ws["J5"].border = border(lato_rosso, lato_rosso, lato_rosso, lato_thick)
        ws["J6"].border = border(lato_rosso, lato_rosso, lato_thin, lato_thin)
        aggiorna_bordo(ws["G5"], right=lato_rosso)
        aggiorna_bordo(ws["I5"], right=lato_rosso)
        aggiorna_bordo(ws["K5"], left=lato_rosso)
        aggiorna_bordo(ws["I6"], right=lato_rosso)
        aggiorna_bordo(ws["K6"], left=lato_rosso)

    ws.freeze_panes = None
    ws.auto_filter.ref = None
    ws.row_dimensions[header_row_1].height = 41
    ws.row_dimensions[header_row_2].height = 54

    for sheet in wb.worksheets:
        sheet.sheet_properties.tabColor = "2E7031" if sheet.title == "FCST" else "FFD966"

    for col_idx, col_cells in enumerate(ws.iter_cols(min_col=1, max_col=last_col), start=1):
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        ws.column_dimensions[lettera_colonna(col_idx)].width = min(max(max_len + 2, 12), 45)

    ws.column_dimensions["A"].width = 2.29
    ws.column_dimensions["B"].width = 28.14
    ws.column_dimensions["C"].width = 12.29
    ws.column_dimensions["D"].width = 12.29
    ws.column_dimensions["E"].width = 12.29
    ws.column_dimensions["F"].width = 17.86
    ws.column_dimensions["G"].width = 9.29
    ws.column_dimensions["H"].width = 15
    ws.column_dimensions["I"].width = 10.29
    ws.column_dimensions["K"].width = 12.43
    ws.column_dimensions["L"].width = 12.43
    ws.column_dimensions["M"].width = 12.43
    ws.column_dimensions["N"].width = 18

    wb.save(path_file)


def formatta_foglio_fcst(path_file) -> None:
    formatta_fcst(path_file)


def main(df_gara: pd.DataFrame, df_oppo: pd.DataFrame, df_appu: pd.DataFrame) -> list:
    print("\n === Creazione report aggregati FCST ===")

    col_venditore_gara = trova_colonna(df_gara, "proprietario", "ordine")
    col_venditore_oppo = trova_colonna(df_oppo, "proprietario", "opportun")
    col_venditore_appu = trova_colonna(df_appu, "in carico")

    venditori = ordina_venditori_crm_ultimo(
        venditori_presenti(
            (df_gara, col_venditore_gara),
            (df_oppo, col_venditore_oppo),
            (df_appu, col_venditore_appu),
        )
    )

    venditori_con_path = ordina_venditori_crm_ultimo(venditori)
    file_creati = []

    for venditore in venditori_con_path:
        # CRM contribuisce al totale di Mandelli ma non ha un report personale.
        if venditore_crm(venditore):
            continue

        if venditore.strip().upper() == "G.MANDELLI":
            df_gara_report = filtra_per_venditori_presenti(df_gara, col_venditore_gara, venditori_con_path)
            df_oppo_report = filtra_per_venditori_presenti(df_oppo, col_venditore_oppo, venditori_con_path)
            df_appu_report = filtra_per_venditori_presenti(df_appu, col_venditore_appu, venditori_con_path)
            includi_totale = True
        else:
            df_gara_report = filtra_per_venditore(df_gara, col_venditore_gara, venditore)
            df_oppo_report = filtra_per_venditore(df_oppo, col_venditore_oppo, venditore)
            df_appu_report = filtra_per_venditore(df_appu, col_venditore_appu, venditore)
            includi_totale = False

        fogli = crea_fogli_report(df_gara_report, df_oppo_report, df_appu_report)
        fogli["FCST"] = crea_pivot_fcst_aggregato(
            df_gara_report,
            df_oppo_report,
            df_appu_report,
            includi_totale=includi_totale,
        )

        nome_file = f"FCST_{pulisci_nome_file(venditore)}.xlsx"
        buffer = BytesIO()
        scrivi_report_excel(buffer, fogli)
        formatta_foglio_fcst(buffer)
        contenuto = buffer.getvalue()
        try:
            caricato = carica_report_venditore(venditore, nome_file, contenuto)
        except (KeyError, ValueError) as errore:
            print(f"[{Fore.RED}ATTENZIONE{Fore.RESET}] salto report {venditore}: {errore}")
            continue

        file_creati.append(caricato)
        print(Fore.GREEN + "OK " + Fore.RESET + nome_file)

    print(f"\nFile creati: {len(file_creati)}")
    print(f"Destinazione: {PATH_OUTPUT_FCST}")

    return file_creati


if __name__ == "__main__":
    file_gara = trova_file(PATH_INPUT_PARQUET, "gara")
    file_opportunita = trova_file(PATH_INPUT_PARQUET, "opportunita")
    file_appuntamenti = trova_file(PATH_INPUT_PARQUET, "appuntamenti")
    df_gara, df_oppo, df_appu = elabora_fcst_main(file_gara, file_opportunita, file_appuntamenti)
    main(df_gara, df_oppo, df_appu)
