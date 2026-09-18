from pathlib import Path
from colorama import Fore, init, Style
init(autoreset=True)
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import warnings 

warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed"
)


PATH_HOME = Path.home()

COLONNE = ["NEGOZIO_ORIGINE", "DATA", "VENDITORE", "NOME", "COGNOME", "CODICE FISCALE", "SERVIZIO", "OFFERTA"]

ONEDRIVE_PATH = {
    "CARPI"                 : PATH_HOME / "My Way S.r.l" / "TEAM - CARPI - TEAM - CARPI" / "TRACCIAMENTO",
    "EUROSIA"               :PATH_HOME / "My Way S.r.l" / "TEAM - EUROSIA - TEAM - EUROSIA" / "TRACCIAMENTO",
    "FIDENZA"               :PATH_HOME / "My Way S.r.l" / "TEAM - FIDENZA - TEAM - FIDENZA" / "TRACCIAMENTO",
    "GALASSIA"              :PATH_HOME / "My Way S.r.l" / "TEAM - GALASSIA - TEAM - GALASSIA" / "TRACCIAMENTO",
    "PIACENZA MEDIAWORLD"   :PATH_HOME / "My Way S.r.l" / "TEAM - PIACENZA MED - TEAM - PIACENZA MED" / "TRACCIAMENTO",
    "SASSUOLO"              :PATH_HOME / "My Way S.r.l" / "TEAM - SASSUOLO - TEAM - SASSUOLO" / "TRACCIAMENTO",
    "TORRI"                 :PATH_HOME / "My Way S.r.l" / "TEAM - TORRI - TEAM - TORRI" / "TRACCIAMENTO",
}

PATH_TEAM_CONSUMER = PATH_HOME / "My Way S.r.l" / "TEAM CONSUMER - TEAM CONSUMER"

def trova_file_da_teams(dict_path: dict, parametro: str) -> list:

    files = []
    if isinstance(dict_path, dict):
        for negozio, path in dict_path.items():
            print(f"\nRicerca file in cartella: {negozio}...")
            if path.is_dir():
                files_cartella = path.glob(parametro)
                for file in files_cartella:
                    print("File trovato: " + Fore.LIGHTBLUE_EX + f"{file.name}")
                    files.append(file)

    return files


def crea_lista_df(files, colonne):
    lista_df = []
    for file in files:
        try:
            #estrae nome
            nome_negozio = file.name.split("_")[0].upper()
            print(f"\nLettura in corso: " + Fore.CYAN+ f"{nome_negozio}...")
            df = pd.read_excel(file, sheet_name = "INSERIMENTO ATTIVATO", dtype=str)
            #rimuove righe vuote (dove non c'è la data)
            df = df.dropna(subset=['DATA'])
            #aggiunge il negozio
            df['NEGOZIO_ORIGINE'] = nome_negozio
            #rimuove le colonne inutili
            subset = df[colonne].copy()
            print(subset)
            lista_df.append(subset)

        except Exception as e:
            print(Fore.RED + f"Errore su {file.name}: {e}")

    return lista_df

def aggrega_df(lista_df):
    if lista_df:
        try:
            df_aggregato = pd.concat(lista_df, ignore_index = True)
            #pulizia/ordine dati
            df_aggregato['DATA'] = pd.to_datetime(df_aggregato['DATA'])
            df_aggregato["CODICE FISCALE"] = df_aggregato["CODICE FISCALE"].str.upper().str.strip()
            df_aggregato = df_aggregato.sort_values(by = "DATA", ascending = False)

            print(f"\n{Fore.GREEN}Successo!{Style.RESET_ALL} Aggregati {len(lista_df)} file.")
            print(f"Totale righe: {len(df_aggregato)}")

        except Exception as e:
            print(Fore.RED + f"errore: {e}")

        print(f"df_aggregato:\n\n{df_aggregato}")
    return df_aggregato

def salva_file_excel(nome_file: str, df_aggregato, path):
    filename = Path(path / nome_file)
    try:
        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            df_aggregato.to_excel(writer, index=False, sheet_name='Trattative firmate')
            
            # Accediamo al foglio di lavoro appena creato
            workbook = writer.book
            worksheet = writer.sheets['Trattative firmate']
            
            # --- PERSONALIZZAZIONE INTESTAZIONI ---
            # Definiamo lo stile: Blu scuro, testo Bianco, grassetto
            header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            for cell in worksheet[1]:  # La riga 1 contiene le intestazioni
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                
            # --- FORMATO DATA E LARGHEZZA COLONNE ---
            for col_idx, column_name in enumerate(df_aggregato.columns, 1):
                letter = get_column_letter(col_idx)
                
                # Applichiamo il formato data se la colonna si chiama "DATA"
                if "DATA" in column_name.upper():
                    for cell in worksheet[letter]:
                        if cell.row > 1: # Saltiamo l'intestazione
                            cell.number_format = 'DD/MM/YYYY'
                
                # Auto-regolazione della larghezza delle colonne
                max_length = max(df_aggregato[column_name].fillna("").astype(str).map(lambda x: len(x.strip())).max(),len(column_name)) + 2
                worksheet.column_dimensions[letter].width = max_length

        print(Fore.GREEN +"\nFile Gara salvato in: " + Fore.YELLOW + f"{filename}")
    except Exception as e:
        print(Fore.RED + f"Errore durante il salvataggio: {e}")