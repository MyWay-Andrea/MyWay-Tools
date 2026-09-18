from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from config import ORDINE_NEGOZI


BLU = "1F4E78"
AZZURRO = "D9EAF7"
VERDE = "00B050"
ROSSO = "FF0000"
GIALLO = "FFD966"
BIANCO = "FFFFFF"
GRIGIO = "E7E6E6"
NERO = "000000"
GIALLO_ENERGIA = "EBE600"
GIALLO_GADGET = "FFC000"
AZZURRO_DIGI = "5B9BD5"

SOTTILE = Side(style="thin", color=NERO)
SPESSO = Side(style="thick", color=NERO)
SOTTILE_CHIARO = Side(style="thin", color="BFBFBF")
NESSUNO = Side(style=None)

LARGHEZZE_CRUSCOTTO = {
    "A": 3.140625,
    "B": 20.0,
    "C": 3.140625,
    "D": 23.5703125,
    "E": 22.7109375,
    "F": 15.85546875,
    "G": 16.28515625,
    "H": 17.140625,
    "I": 3.140625,
    "J": 11.7109375,
    "K": 3.140625,
    "L": 11.7109375,
    "M": 3.140625,
    "N": 14.0,
    "O": 19.5703125,
    "P": 25.140625,
}

ALTEZZE_CRUSCOTTO = {
    1: 15.75,
    2: 15.75,
    3: 15.0,
    5: 15.75,
    6: 15.75,
    7: 21.75,
    8: 16.5,
    9: 25.5,
    10: 16.5,
    11: 15.75,
    12: 15.75,
    13: 15.75,
    14: 15.75,
    15: 15.75,
    16: 15.75,
    17: 16.5,
}


def _valore_excel(valore):
    if pd.isna(valore):
        return None
    if isinstance(valore, pd.Timestamp):
        return valore.to_pydatetime()
    return valore


def _scrivi_dataframe(ws, df: pd.DataFrame, nome_tabella: str, start_row: int = 1) -> None:
    df_excel = df.copy().astype(object).where(pd.notna(df), None)
    for indice_riga, riga in enumerate(
        dataframe_to_rows(df_excel, index=False, header=True),
        start=start_row,
    ):
        for indice_colonna, valore in enumerate(riga, start=1):
            ws.cell(indice_riga, indice_colonna, _valore_excel(valore))

    ultima_riga = start_row + max(len(df), 1)
    ultima_colonna = max(len(df.columns), 1)
    if df.empty:
        for colonna in range(1, ultima_colonna + 1):
            ws.cell(start_row + 1, colonna, None)

    ws.freeze_panes = f"A{start_row + 1}"
    ws.sheet_view.showGridLines = False

    for cella in ws[start_row]:
        cella.font = Font(color=BIANCO, bold=True)
        cella.fill = PatternFill("solid", fgColor=BLU)
        cella.alignment = Alignment(horizontal="center", vertical="center")

    for indice, colonna in enumerate(df.columns, start=1):
        valori = [str(colonna)] + [str(v) for v in df[colonna].dropna().head(500)]
        larghezza = min(max(max(map(len, valori), default=10) + 2, 12), 35)
        ws.column_dimensions[get_column_letter(indice)].width = larghezza
        if "DATA" in str(colonna).upper():
            for cella in ws[get_column_letter(indice)][start_row:]:
                cella.number_format = "dd/mm/yyyy"


def _prepara_andamento_giornaliero(
    ws,
    df: pd.DataFrame,
    larghezze: dict[str, float] | None = None,
    colonna_data: str | None = "A",
) -> None:
    """Ricrea lo stile del foglio omonimo nel master CRUSCOTTO2.xlsx."""
    df_excel = df.copy().astype(object).where(pd.notna(df), None)
    for indice_riga, riga in enumerate(
        dataframe_to_rows(df_excel, index=False, header=True),
        start=1,
    ):
        for indice_colonna, valore in enumerate(riga, start=1):
            ws.cell(indice_riga, indice_colonna, _valore_excel(valore))

    ultima_riga = max(len(df) + 1, 2)
    ultima_colonna = max(len(df.columns), 1)
    if df.empty:
        for colonna in range(1, ultima_colonna + 1):
            ws.cell(2, colonna, None)

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    for cella in ws[1]:
        cella.font = Font(name="Calibri", size=11, color=BIANCO, bold=True)
        cella.fill = PatternFill("solid", fgColor="0070C0")
        cella.alignment = Alignment(horizontal="center", vertical="center")
        cella.border = Border(
            left=SOTTILE,
            right=SOTTILE,
            top=SOTTILE,
        )

    for indice_riga in range(2, ultima_riga + 1):
        for indice_colonna in range(1, ultima_colonna + 1):
            cella = ws.cell(indice_riga, indice_colonna)
            cella.font = Font(name="Calibri", size=11)
            cella.fill = PatternFill(fill_type=None)
            cella.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cella.border = Border(
                left=SPESSO if indice_colonna == 1 else SOTTILE_CHIARO,
                right=SPESSO
                if indice_colonna == ultima_colonna
                else SOTTILE_CHIARO,
                top=SPESSO
                if indice_riga == 2
                else SOTTILE_CHIARO,
                bottom=SPESSO
                if indice_riga == ultima_riga
                else SOTTILE_CHIARO,
            )

    # Misure rilevate dal master CRUSCOTTO2, foglio ANDAMENTO GIORNALIERO.
    larghezze = larghezze or {
        "A": 21.0,
        "B": 16.140625,
        "C": 28.140625,
        "D": 20.42578125,
        "E": 15.140625,
        "F": 16.42578125,
    }
    for colonna, larghezza in larghezze.items():
        ws.column_dimensions[colonna].width = larghezza

    ws.row_dimensions[1].height = 24.0
    for indice_riga in range(2, ultima_riga + 1):
        ws.row_dimensions[indice_riga].height = 15.75

    if colonna_data:
        for cella in ws[colonna_data][1:ultima_riga]:
            cella.number_format = "mm-dd-yy"


def _bordo_cruscotto(
    left: str | None = "thin",
    right: str | None = "thin",
    top: str | None = "thin",
    bottom: str | None = "thin",
) -> Border:
    lati = {"thin": SOTTILE, "thick": SPESSO, None: NESSUNO}
    return Border(
        left=lati[left],
        right=lati[right],
        top=lati[top],
        bottom=lati[bottom],
    )


def _stile_blocco_cruscotto(
    ws,
    intervallo: str,
    colore: str,
    colore_testo: str = NERO,
) -> None:
    for riga in ws[intervallo]:
        for cella in riga:
            cella.fill = PatternFill("solid", fgColor=colore)
            cella.font = Font(
                name="Aharoni",
                size=11,
                bold=True,
                color=colore_testo,
            )
            cella.alignment = Alignment(horizontal="center", vertical="center")


def _prepara_dettaglio_business(
    ws,
    dettaglio: pd.DataFrame,
) -> None:
    """Crea una tabella per negozio, tre tabelle per ogni riga."""

    ws.title = "DETTAGLIO BUSINESS"
    ws.sheet_view.showGridLines = False

    larghezze = {
        "A": 6.85546875,
        "B": 29.5703125,
        "C": 13.0,
        "D": 13.0,
        "E": 13.0,
        "F": 13.0,
        "G": 13.0,
        "H": 20.0,
        "I": 13.28515625,
        "J": 13.0,
        "K": 13.0,
        "L": 13.0,
        "M": 13.0,
        "N": 19.0,
        "O": 12.28515625,
        "P": 13.0,
        "Q": 13.0,
        "R": 13.0,
    }
    for colonna, larghezza in larghezze.items():
        ws.column_dimensions[colonna].width = larghezza

    if dettaglio.empty:
        return

    negozi = dettaglio["NEGOZIO"].drop_duplicates().tolist()
    piste = dettaglio["FAMIGLIA_PRODOTTO"].drop_duplicates().tolist()
    colonne_iniziali = (2, 8, 14)  # B, H, N
    altezza_tabella = len(piste) + 2
    passo_verticale = altezza_tabella + 2
    intestazioni = (
        None,
        "QTÀ INS",
        "PUNTI INS",
        "QTÀ _ATT",
        "PUNTI_ATT",
    )
    campi = (
        "FAMIGLIA_PRODOTTO",
        "QUANTITA_INSERITA",
        "PUNTI_INSERITI",
        "QUANTITA_ATTIVATA",
        "PUNTI_ATTIVATI",
    )

    for indice_negozio, negozio in enumerate(negozi):
        gruppo_riga = indice_negozio // 3
        posizione = indice_negozio % 3
        riga_iniziale = 3 + gruppo_riga * passo_verticale
        colonna_iniziale = colonne_iniziali[posizione]
        dati_negozio = (
            dettaglio[dettaglio["NEGOZIO"] == negozio]
            .set_index("FAMIGLIA_PRODOTTO")
            .reindex(piste)
        )

        ws.row_dimensions[riga_iniziale].height = 27.0
        ws.row_dimensions[riga_iniziale + 1].height = (
            20.25 if gruppo_riga == 0 else 19.5
        )
        for offset_pista in range(len(piste)):
            ws.row_dimensions[riga_iniziale + 2 + offset_pista].height = (
                16.5
                if offset_pista in (0, len(piste) - 1)
                else 15.75
            )
        for offset_vuoto in (altezza_tabella, altezza_tabella + 1):
            ws.row_dimensions[riga_iniziale + offset_vuoto].height = 15.75

        for offset_colonna, intestazione in enumerate(intestazioni):
            cella = ws.cell(
                riga_iniziale,
                colonna_iniziale + offset_colonna,
            )
            cella.value = negozio if offset_colonna == 0 else intestazione
            cella.fill = PatternFill(
                "solid",
                fgColor=(
                    "92D050"
                    if offset_colonna == 0
                    else ROSSO
                    if offset_colonna in (1, 2)
                    else VERDE
                ),
            )
            cella.font = Font(
                name="Aharoni",
                size=20 if offset_colonna == 0 else 11,
                bold=offset_colonna != 0,
                color=NERO if offset_colonna == 0 else BIANCO,
            )
            cella.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cella.border = Border(
                left=SPESSO,
                right=SPESSO,
                top=SPESSO,
                bottom=SOTTILE,
            )

        riga_totale = riga_iniziale + 1
        totali = [
            "TOT.",
            dati_negozio["QUANTITA_INSERITA"].sum(),
            dati_negozio["PUNTI_INSERITI"].sum(),
            dati_negozio["QUANTITA_ATTIVATA"].sum(),
            dati_negozio["PUNTI_ATTIVATI"].sum(),
        ]
        for offset_colonna, valore in enumerate(totali):
            cella = ws.cell(
                riga_totale,
                colonna_iniziale + offset_colonna,
            )
            cella.value = _valore_excel(valore)
            cella.font = Font(
                name="Aharoni" if offset_colonna == 0 else "Calibri",
                size=14,
                bold=offset_colonna != 0,
            )
            cella.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cella.number_format = "0"
            cella.border = Border(
                left=SPESSO,
                right=SPESSO,
                top=SOTTILE,
                bottom=SPESSO,
            )

        for offset_pista, pista in enumerate(piste):
            riga = riga_iniziale + 2 + offset_pista
            valori = [
                pista,
                dati_negozio.at[pista, "QUANTITA_INSERITA"],
                dati_negozio.at[pista, "PUNTI_INSERITI"],
                dati_negozio.at[pista, "QUANTITA_ATTIVATA"],
                dati_negozio.at[pista, "PUNTI_ATTIVATI"],
            ]
            prima_pista = offset_pista == 0
            ultima_pista = offset_pista == len(piste) - 1
            for offset_colonna, valore in enumerate(valori):
                cella = ws.cell(
                    riga,
                    colonna_iniziale + offset_colonna,
                )
                cella.value = _valore_excel(valore)
                cella.font = Font(
                    name="Aharoni" if offset_colonna == 0 else "Calibri",
                    size=11 if offset_colonna == 0 else 12,
                    bold=offset_colonna == 0,
                )
                cella.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )
                cella.number_format = "0"

                if offset_colonna == 0:
                    sinistro, destro = SPESSO, NESSUNO
                elif offset_colonna == 1:
                    sinistro, destro = SPESSO, SOTTILE_CHIARO
                elif offset_colonna == 4:
                    sinistro, destro = SOTTILE_CHIARO, SPESSO
                else:
                    sinistro = destro = SOTTILE_CHIARO
                cella.border = Border(
                    left=sinistro,
                    right=destro,
                    top=SPESSO if prima_pista else SOTTILE_CHIARO,
                    bottom=SPESSO if ultima_pista else SOTTILE_CHIARO,
                )

    ultima_riga = (
        3
        + ((len(negozi) - 1) // 3) * passo_verticale
        + altezza_tabella
        - 1
    )
    ws.print_area = f"B3:R{ultima_riga}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2


def _bordo_report_verticale(
    colonna: int,
    ultima_colonna: int,
    *,
    riepilogo: bool,
    ultima_riga_gruppo: bool = False,
) -> Border:
    return Border(
        left=SPESSO
        if colonna in (1, 3)
        else SOTTILE_CHIARO,
        right=SPESSO
        if colonna in (2, ultima_colonna)
        else SOTTILE_CHIARO,
        top=SOTTILE if riepilogo else SOTTILE_CHIARO,
        bottom=SOTTILE
        if ultima_riga_gruppo
        else SOTTILE_CHIARO,
    )


def _stile_intestazione_report(
    ws,
    ultima_colonna: int,
    colori: dict[int, str],
    colonne_testo_nero: set[int] | None = None,
) -> None:
    colonne_testo_nero = colonne_testo_nero or set()
    ws.row_dimensions[1].height = 36.75
    for colonna in range(1, ultima_colonna + 1):
        cella = ws.cell(1, colonna)
        cella.fill = PatternFill("solid", fgColor=colori[colonna])
        cella.font = Font(
            name="Aharoni",
            size=14,
            bold=True,
            color=NERO if colonna in colonne_testo_nero else BIANCO,
        )
        cella.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cella.border = Border(
            left=SPESSO if colonna == 3 else SOTTILE,
            right=SOTTILE,
            top=SOTTILE,
            bottom=SPESSO,
        )


def _prepara_dettaglio_business_verticale(
    ws,
    dettaglio: pd.DataFrame,
) -> None:
    """Ricrea il nuovo DETTAGLIO BUSINESS verticale."""
    ws.title = "DETTAGLIO BUSINESS"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    for colonna, larghezza in {
        "A": 13.0,
        "B": 17.5703125,
        "C": 22.0,
        "D": 20.0,
        "E": 20.0,
        "F": 20.0,
    }.items():
        ws.column_dimensions[colonna].width = larghezza

    ws.merge_cells("A1:B1")
    ws["A1"] = "NEGOZIO / PISTA"
    for colonna, valore in enumerate(
        (
            "QTA INSERITO",
            "PUNTI INSERITO",
            "QTA ATTIVATO",
            "PUNTI ATTIVATO",
        ),
        start=3,
    ):
        ws.cell(1, colonna, valore)
    _stile_intestazione_report(
        ws,
        6,
        {
            1: "0070C0",
            2: "0070C0",
            3: ROSSO,
            4: ROSSO,
            5: VERDE,
            6: VERDE,
        },
    )

    if dettaglio.empty:
        return

    negozi = dettaglio["NEGOZIO"].drop_duplicates().tolist()
    piste = dettaglio["FAMIGLIA_PRODOTTO"].drop_duplicates().tolist()
    riga = 2
    for negozio in negozi:
        dati_negozio = (
            dettaglio[dettaglio["NEGOZIO"] == negozio]
            .set_index("FAMIGLIA_PRODOTTO")
            .reindex(piste)
            .fillna(0)
        )
        ws.merge_cells(
            start_row=riga,
            start_column=1,
            end_row=riga,
            end_column=2,
        )
        ws.cell(riga, 1, negozio)
        totali = (
            dati_negozio["QUANTITA_INSERITA"].sum(),
            dati_negozio["PUNTI_INSERITI"].sum(),
            dati_negozio["QUANTITA_ATTIVATA"].sum(),
            dati_negozio["PUNTI_ATTIVATI"].sum(),
        )
        for colonna, valore in enumerate(totali, start=3):
            ws.cell(riga, colonna, _valore_excel(valore))

        for colonna in range(1, 7):
            cella = ws.cell(riga, colonna)
            cella.fill = PatternFill("solid", fgColor="FFFFB7")
            cella.font = Font(
                name="Aharoni" if colonna <= 2 else "Calibri",
                size=14 if colonna <= 2 else 12,
                bold=True,
            )
            cella.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cella.number_format = "0"
            cella.border = _bordo_report_verticale(
                colonna,
                6,
                riepilogo=True,
            )
        ws.row_dimensions[riga].height = 19.5
        riga += 1

        for indice_pista, pista in enumerate(piste):
            valori = (
                None,
                pista,
                dati_negozio.at[pista, "QUANTITA_INSERITA"],
                dati_negozio.at[pista, "PUNTI_INSERITI"],
                dati_negozio.at[pista, "QUANTITA_ATTIVATA"],
                dati_negozio.at[pista, "PUNTI_ATTIVATI"],
            )
            ultima_pista = indice_pista == len(piste) - 1
            for colonna, valore in enumerate(valori, start=1):
                cella = ws.cell(riga, colonna, _valore_excel(valore))
                cella.font = Font(
                    name="Aharoni" if colonna <= 2 else "Calibri",
                    size=11 if colonna <= 2 else 12,
                    bold=colonna == 2,
                )
                cella.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )
                cella.number_format = "0"
                cella.border = _bordo_report_verticale(
                    colonna,
                    6,
                    riepilogo=False,
                    ultima_riga_gruppo=ultima_pista,
                )
            ws.row_dimensions[riga].height = 15.75
            riga += 1

    ultima_riga = riga - 1
    ws.print_area = f"A1:F{ultima_riga}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2


def _prepara_dettaglio_venditori(
    ws,
    venditori: pd.DataFrame,
) -> None:
    """Ricrea il nuovo DETTAGLIO VENDITORI raggruppato per negozio."""
    ws.title = "DETTAGLIO VENDITORI"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    for colonna, larghezza in {
        "A": 18.0,
        "B": 17.5703125,
        "C": 29.28515625,
        "D": 28.28515625,
        "E": 21.28515625,
        "F": 15.42578125,
        "G": 24.28515625,
        "H": 15.42578125,
    }.items():
        ws.column_dimensions[colonna].width = larghezza

    ws.merge_cells("A1:B1")
    ws["A1"] = "NEGOZIO / VENDITORE"
    for colonna, valore in enumerate(
        (
            "BUSINESS INSERITO",
            "BUSINESS ATTIVATO",
            "TRAT ENERGIA",
            "DIGI",
            "GADGET",
            "CROSS SELLING",
        ),
        start=3,
    ):
        ws.cell(1, colonna, valore)
    _stile_intestazione_report(
        ws,
        8,
        {
            1: "0070C0",
            2: "0070C0",
            3: ROSSO,
            4: VERDE,
            5: "FFFF00",
            6: AZZURRO_DIGI,
            7: "FFC000",
            8: "FFC000",
        },
        colonne_testo_nero={5},
    )

    if venditori.empty:
        return

    colonne_valori = (
        "BUSINESS_INSERITO",
        "BUSINESS_ATTIVATO",
        "ENERGIA",
        "DIGI",
        "GADGET",
        "CROSS_SELLING",
    )
    riga = 2
    for negozio in venditori["NEGOZIO"].drop_duplicates():
        gruppo = venditori[venditori["NEGOZIO"] == negozio]
        ws.merge_cells(
            start_row=riga,
            start_column=1,
            end_row=riga,
            end_column=2,
        )
        ws.cell(riga, 1, negozio)
        for colonna, campo in enumerate(colonne_valori, start=3):
            ws.cell(riga, colonna, _valore_excel(gruppo[campo].sum()))

        for colonna in range(1, 9):
            cella = ws.cell(riga, colonna)
            cella.fill = PatternFill("solid", fgColor="FFFFB7")
            cella.font = Font(
                name="Aharoni" if colonna <= 2 else "Calibri",
                size=14 if colonna <= 2 else 12,
                bold=True,
            )
            cella.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cella.number_format = "0"
            cella.border = _bordo_report_verticale(
                colonna,
                8,
                riepilogo=True,
            )
        ws.row_dimensions[riga].height = 15.75
        riga += 1

        for indice_venditore, (_, venditore) in enumerate(
            gruppo.iterrows()
        ):
            ultima_riga_gruppo = indice_venditore == len(gruppo) - 1
            valori = (
                None,
                venditore["VENDITORE"],
                *[venditore[campo] for campo in colonne_valori],
            )
            for colonna, valore in enumerate(valori, start=1):
                cella = ws.cell(riga, colonna, _valore_excel(valore))
                cella.font = Font(name="Calibri", size=11)
                cella.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )
                cella.number_format = "0" if colonna >= 3 else "General"
                cella.border = _bordo_report_verticale(
                    colonna,
                    8,
                    riepilogo=False,
                    ultima_riga_gruppo=ultima_riga_gruppo,
                )
            ws.row_dimensions[riga].height = 15.75
            riga += 1

    ultima_riga = riga - 1
    ws.print_area = f"A1:H{ultima_riga}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2


def _prepara_dashboard(
    ws,
    dashboard: pd.DataFrame,
    riferimento: pd.Timestamp,
) -> None:
    """Genera CRUSCOTTO con layout e stile del master CRUSCOTTO2.xlsx."""

    ws.title = "CRUSCOTTO"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100

    for colonna, larghezza in LARGHEZZE_CRUSCOTTO.items():
        ws.column_dimensions[colonna].width = larghezza
    for riga, altezza in ALTEZZE_CRUSCOTTO.items():
        ws.row_dimensions[riga].height = altezza
    ultima_riga_negozi = 9 + len(ORDINE_NEGOZI)
    for riga in range(10, ultima_riga_negozi + 1):
        ws.row_dimensions[riga].height = 15.75

    intervalli_uniti = (
        "C1:D1",
        "C2:D2",
        "B3:P4",
    )
    for intervallo in intervalli_uniti:
        ws.merge_cells(intervallo)

    ws["B1"] = "Periodo:"
    ws["C1"] = riferimento.to_pydatetime()
    ws["C1"].number_format = "mmm-yy"
    ws["B2"] = "Aggiornato il :"
    ws["C2"] = datetime.now()
    ws["C2"].number_format = "m/d/yy h:mm"
    for cella in (ws["B1"], ws["B2"]):
        cella.font = Font(name="Calibri", size=12, bold=True)
        cella.alignment = Alignment(horizontal="right")
    for cella in (ws["C1"], ws["C2"]):
        cella.font = Font(name="Calibri", size=12, bold=True)
        cella.alignment = Alignment(horizontal="left")

    ws["B3"] = "CRUSCOTTO AGGREGATIVO"
    ws["B3"].font = Font(name="Aharoni", size=24, bold=True)
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B3"].fill = PatternFill("solid", fgColor=BIANCO)

    kpi = (
        (
            "D6:D6",
            "D6",
            "BUSINESS_INSERITO",
            BLU,
            BIANCO,
            "D7:D7",
            "D7",
            "BUSINESS_INSERITO",
        ),
        (
            "E6:E6",
            "E6",
            "BUSINESS ATTIVATO",
            VERDE,
            BIANCO,
            "E7:E7",
            "E7",
            "BUSINESS_ATTIVATO",
        ),
        (
            "F6:F6",
            "F6",
            "TARGET",
            BLU,
            "FFFF00",
            "F7:F7",
            "F7",
            "TARGET",
        ),
        (
            "G6:G6",
            "G6",
            "PROIEZIONE",
            BLU,
            BIANCO,
            "G7:G7",
            "G7",
            "PROIEZIONE",
        ),
        (
            "H6:H6",
            "H6",
            "VAR_TARGET",
            BLU,
            BIANCO,
            "H7:H7",
            "H7",
            "VAR_TARGET_TOTALE",
        ),
        (
            "J6:J6",
            "J6",
            "ENERGIA",
            GIALLO_ENERGIA,
            NERO,
            "J7:J7",
            "J7",
            "ENERGIA",
        ),
        (
            "L6:L6",
            "L6",
            "DIGI",
            AZZURRO_DIGI,
            BIANCO,
            "L7:L7",
            "L7",
            "DIGI",
        ),
        (
            "N6:N6",
            "N6",
            "GADGET",
            GIALLO_GADGET,
            NERO,
            "N7:N7",
            "N7",
            "GADGET",
        ),
        (
            "O6:O6",
            "O6",
            "CROSS SELLING",
            GIALLO_GADGET,
            NERO,
            "O7:O7",
            "O7",
            "CROSS_SELLING",
        ),
        (
            "P6:P6",
            "P6",
            "% CROSS SELLING TOTALE",
            GIALLO_GADGET,
            NERO,
            "P7:P7",
            "P7",
            "PERC_CROSS_SELLING_TOTALE",
        ),
    )
    for (
        header_range,
        header_cell,
        titolo,
        colore,
        colore_testo,
        value_range,
        value_cell,
        colonna_dati,
    ) in kpi:
        _stile_blocco_cruscotto(
            ws,
            header_range,
            colore,
            colore_testo,
        )
        ws[header_cell] = titolo
        for riga in ws[value_range]:
            for cella in riga:
                cella.font = Font(name="Calibri", size=16, bold=True)
                cella.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )
        if colonna_dati == "PERC_CROSS_SELLING_TOTALE":
            totale_gadget = pd.to_numeric(
                dashboard["GADGET"],
                errors="coerce",
            ).fillna(0).sum()
            totale_cross_selling = pd.to_numeric(
                dashboard["CROSS_SELLING"],
                errors="coerce",
            ).fillna(0).sum()
            ws[value_cell] = (
                float(totale_cross_selling / totale_gadget)
                if totale_gadget > 0
                else 0
            )
            ws[value_cell].number_format = "0.00%"
        elif colonna_dati == "VAR_TARGET_TOTALE":
            totale_target = pd.to_numeric(
                dashboard["TARGET"],
                errors="coerce",
            ).fillna(0).sum()
            totale_proiezione = pd.to_numeric(
                dashboard["PROIEZIONE"],
                errors="coerce",
            ).fillna(0).sum()
            var_target = (
                float((totale_proiezione - totale_target) / totale_target)
                if totale_target > 0
                else 0
            )
            ws[value_cell] = var_target
            ws[value_cell].number_format = "0.0%"
            ws[value_cell].font = Font(
                name="Calibri",
                size=16,
                bold=True,
                color=VERDE if var_target >= 0 else ROSSO,
            )
        else:
            ws[value_cell] = float(
                pd.to_numeric(
                    dashboard[colonna_dati],
                    errors="coerce",
                ).fillna(0).sum()
            )
            if colonna_dati in ("TARGET", "PROIEZIONE"):
                ws[value_cell].number_format = "#,##0"

        ultima_header = ws[header_range.split(":")[-1]].column
        ultima_value = ws[value_range.split(":")[-1]].column
        prima_header = ws[header_cell].column
        prima_value = ws[value_cell].column
        for cella in ws[header_range][0]:
            cella.border = _bordo_cruscotto(
                "thick" if cella.column == prima_header else None,
                "thick" if cella.column == ultima_header else None,
                "thick",
                "thin",
            )
        for cella in ws[value_range][0]:
            cella.border = _bordo_cruscotto(
                "thick" if cella.column == prima_value else None,
                "thick" if cella.column == ultima_value else None,
                "thin",
                "thick",
            )

    gruppi_intestazioni = (
        (["B"], ["NEGOZIO"]),
        (
            ["D", "E", "F", "G", "H"],
            [
                "BUSINESS_INSERITO",
                "BUSINESS_ATTIVATO",
                "TARGET",
                "PROIEZIONE",
                "VAR_TARGET",
            ],
        ),
        (["J"], ["ENERGIA"]),
        (["L"], ["DIGI"]),
        (
            ["N", "O", "P"],
            ["GADGET", "CROSS_SELLING", "PERC_CROSS_SELLING"],
        ),
    )
    for colonne, intestazioni in gruppi_intestazioni:
        for indice, (colonna, intestazione) in enumerate(
            zip(colonne, intestazioni)
        ):
            cella = ws[f"{colonna}9"]
            cella.value = intestazione
            cella.fill = PatternFill(
                "solid",
                fgColor=(
                    VERDE
                    if intestazione == "BUSINESS_ATTIVATO"
                    else AZZURRO_DIGI if intestazione == "DIGI" else BLU
                ),
            )
            cella.font = Font(
                name="Aharoni",
                size=11,
                bold=True,
                color="FFFF00" if intestazione == "TARGET" else BIANCO,
            )
            cella.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            if intestazione == "BUSINESS_ATTIVATO":
                sinistro, destro = "thick", "thick"
            elif intestazione == "TARGET":
                sinistro, destro = None, "thin"
            else:
                sinistro = "thick" if indice == 0 else "thin"
                destro = (
                    "thick" if indice == len(colonne) - 1 else "thin"
                )
            cella.border = _bordo_cruscotto(
                sinistro,
                destro,
                "thick",
                "thick",
            )

    record_per_negozio = {
        record["NEGOZIO"]: record
        for record in dashboard.to_dict("records")
    }
    ordine_negozi = ORDINE_NEGOZI
    colonne_dati = {
        "B": "NEGOZIO",
        "D": "BUSINESS_INSERITO",
        "E": "BUSINESS_ATTIVATO",
        "F": "TARGET",
        "G": "PROIEZIONE",
        "H": "VAR_TARGET",
        "J": "ENERGIA",
        "L": "DIGI",
        "N": "GADGET",
        "O": "CROSS_SELLING",
        "P": "PERC_CROSS_SELLING",
    }
    gruppi_colonne = (("B", "B"), ("D", "H"), ("J", "J"), ("L", "L"), ("N", "P"))

    for offset, negozio in enumerate(ordine_negozi):
        riga = 10 + offset
        record = record_per_negozio.get(negozio, {"NEGOZIO": negozio})
        for colonna, campo in colonne_dati.items():
            valore = record.get(campo)
            if campo in ("TARGET", "VAR_TARGET", "PERC_CROSS_SELLING") and pd.isna(valore):
                valore = "-"
            elif pd.isna(valore):
                valore = 0

            cella = ws[f"{colonna}{riga}"]
            cella.value = valore
            cella.font = Font(
                name="Calibri",
                size=12,
                bold=campo
                in (
                    "NEGOZIO",
                    "TARGET",
                    "VAR_TARGET",
                    "PERC_CROSS_SELLING",
                ),
            )
            cella.alignment = Alignment(
                horizontal=None if campo == "NEGOZIO" else "center",
                vertical="center",
            )
            if campo == "VAR_TARGET":
                cella.number_format = "0.0%"
            elif campo == "PERC_CROSS_SELLING":
                cella.number_format = "0.00%"
            elif campo == "NEGOZIO":
                cella.number_format = "General"
            else:
                cella.number_format = "#,##0"

            if campo == "VAR_TARGET" and valore != "-":
                cella.font = Font(
                    name="Calibri",
                    size=12,
                    bold=True,
                    color=VERDE if float(valore) >= 0 else ROSSO,
                )

        for colonna_iniziale, colonna_finale in gruppi_colonne:
            for codice in range(
                ord(colonna_iniziale),
                ord(colonna_finale) + 1,
            ):
                colonna = chr(codice)
                if colonna not in colonne_dati:
                    continue
                ws[f"{colonna}{riga}"].border = _bordo_cruscotto(
                    "thick" if colonna == colonna_iniziale else "thin",
                    "thick" if colonna == colonna_finale else "thin",
                    None if riga == 10 else "thin",
                    "thick" if riga == ultima_riga_negozi else "thin",
                )

    ws.print_area = f"B1:P{ultima_riga_negozi}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2


def _controlli(sorgenti: dict[str, object], dati: dict[str, pd.DataFrame | pd.Timestamp]) -> pd.DataFrame:
    righe = []
    associazioni = (
        ("Business", sorgenti["business"], len(dati["business"])),
        ("Pedonalita", sorgenti["pedonalita"], None),
    )
    for tipo, file_remoto, numero_righe in associazioni:
        ultima_modifica = pd.to_datetime(
            file_remoto.get("lastModifiedDateTime"),
            errors="coerce",
            utc=True,
        )
        if not pd.isna(ultima_modifica):
            ultima_modifica = ultima_modifica.tz_localize(None)
        righe.append({
            "SORGENTE": tipo,
            "NEGOZIO": "",
            "FILE": file_remoto.get("webUrl", file_remoto.get("name", "")),
            "ULTIMA_MODIFICA": ultima_modifica,
            "RIGHE_LETTE": numero_righe,
            "ESITO": "OK",
        })
    for negozio, file_remoto in sorgenti["tracciamenti"].items():
        ultima_modifica = pd.to_datetime(
            file_remoto.get("lastModifiedDateTime"),
            errors="coerce",
            utc=True,
        )
        if not pd.isna(ultima_modifica):
            ultima_modifica = ultima_modifica.tz_localize(None)
        righe.append({
                "SORGENTE": "Tracciamento",
                "NEGOZIO": negozio,
                "FILE": file_remoto.get("webUrl", file_remoto.get("name", "")),
                "ULTIMA_MODIFICA": ultima_modifica,
                "RIGHE_LETTE": sum(
                    int((dati[chiave]["NEGOZIO"] == negozio).sum())
                    for chiave in ("energia", "digi", "gadget")
                ),
                "ESITO": "OK",
        })
    return pd.DataFrame(righe)


def _verifica_business_attivato(
    dashboard: pd.DataFrame,
    business: pd.DataFrame,
) -> None:
    """Impedisce di generare un cruscotto con attivati non riconciliati."""

    richieste_dashboard = {"NEGOZIO", "BUSINESS_ATTIVATO"}
    richieste_business = {"NEGOZIO", "PUNTI_ATTIVATI"}
    if not richieste_dashboard.issubset(dashboard.columns):
        mancanti = sorted(richieste_dashboard.difference(dashboard.columns))
        raise ValueError(
            "Colonne mancanti nel dashboard: " + ", ".join(mancanti)
        )
    if not richieste_business.issubset(business.columns):
        mancanti = sorted(richieste_business.difference(business.columns))
        raise ValueError(
            "Colonne mancanti nei dati business: " + ", ".join(mancanti)
        )

    attesi = (
        business.assign(
            PUNTI_ATTIVATI=pd.to_numeric(
                business["PUNTI_ATTIVATI"],
                errors="coerce",
            ).fillna(0)
        )
        .groupby("NEGOZIO", dropna=False)["PUNTI_ATTIVATI"]
        .sum()
    )
    ottenuti = (
        dashboard.assign(
            BUSINESS_ATTIVATO=pd.to_numeric(
                dashboard["BUSINESS_ATTIVATO"],
                errors="coerce",
            ).fillna(0)
        )
        .groupby("NEGOZIO", dropna=False)["BUSINESS_ATTIVATO"]
        .sum()
    )
    confronto = pd.concat(
        [attesi.rename("ATTESO"), ottenuti.rename("OTTENUTO")],
        axis=1,
    ).fillna(0)
    differenze = confronto[
        ~confronto["ATTESO"].round(6).eq(
            confronto["OTTENUTO"].round(6)
        )
    ]
    if not differenze.empty:
        dettagli = "; ".join(
            f"{negozio}: atteso {riga.ATTESO:g}, ottenuto {riga.OTTENUTO:g}"
            for negozio, riga in differenze.iterrows()
        )
        raise ValueError(
            "BUSINESS ATTIVATO non riconciliato con "
            f"DATI_BUSINESS ({dettagli})"
        )


def genera_cruscotto(
    dati: dict[str, pd.DataFrame | pd.Timestamp],
    sorgenti: dict[str, object],
    percorso_output: Path,
) -> Path:
    _verifica_business_attivato(
        dati["dashboard"],
        dati["business_att"],
    )

    wb = Workbook()
    _prepara_dashboard(wb.active, dati["dashboard"], dati["riferimento"])

    ws_dettaglio_business = wb.create_sheet("DETTAGLIO BUSINESS")
    _prepara_dettaglio_business_verticale(
        ws_dettaglio_business,
        dati["dettaglio_business"],
    )

    ws_venditori = wb.create_sheet("DETTAGLIO VENDITORI")
    _prepara_dettaglio_venditori(
        ws_venditori,
        dati["venditori"],
    )

    ws_andamento = wb.create_sheet("ANDAMENTO GIORNALIERO")
    _prepara_andamento_giornaliero(ws_andamento, dati["giornaliero"])

    ws_controlli = wb.create_sheet("CONTROLLI")
    _scrivi_dataframe(
        ws_controlli,
        _controlli(sorgenti, dati),
        "TabellaControlli",
    )

    fogli_nascosti = (
        ("DATI_BUSINESS", dati["business"], "TabellaDatiBusiness"),
        ("DATI_ENERGIA", dati["energia"], "TabellaDatiEnergia"),
        ("DATI_DIGI", dati["digi"], "TabellaDatiDigi"),
        ("DATI_GADGET", dati["gadget"], "TabellaDatiGadget"),
    )
    for nome, frame, tabella in fogli_nascosti:
        ws = wb.create_sheet(nome)
        _scrivi_dataframe(ws, frame, tabella)
        ws.sheet_state = "hidden"

    percorso_output = Path(percorso_output)
    percorso_output.parent.mkdir(parents=True, exist_ok=True)
    temporaneo = percorso_output.with_name(f".{percorso_output.stem}.tmp.xlsx")
    wb.save(temporaneo)
    try:
        os.replace(temporaneo, percorso_output)
    except PermissionError as errore:
        temporaneo.unlink(missing_ok=True)
        raise PermissionError(
            f"Impossibile aggiornare {percorso_output}. Chiudi il file in Excel e riprova."
        ) from errore
    return percorso_output
