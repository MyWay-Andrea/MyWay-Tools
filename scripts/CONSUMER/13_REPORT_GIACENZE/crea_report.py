from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BLU = "FF0070C0"
AZZURRO = "FFCCECFF"
ROSSO = "FFFF0000"
ROSSO_SFOCATO = "FFFFC1C1"
GIALLO = "FFFFFF00"
GRIGIO_TOTALE = "FFF2F2F2"
GRIGIO_BORDO = "FFBFBFBF"
GRIGIO_TRATTINO = "FF7F7F7F"
NERO = "FF000000"
BIANCO = "FFFFFFFF"


def _lato(stile: str | None, colore: str = NERO) -> Side:
    return Side(style=stile, color=colore) if stile else Side()


def _bordi_cella(
    colonna: int,
    prima_colonna: int,
    colonna_descrizione: int,
    primo_negozio: int,
    ultimo_negozio: int,
    colonna_totale: int,
    colonna_valore: int,
    *,
    bordo_alto: str | None,
    bordo_basso: str | None,
) -> Border:
    if colonna == prima_colonna:
        sinistro = _lato("thick")
    elif colonna in {primo_negozio, colonna_totale}:
        sinistro = _lato("medium")
    elif colonna == colonna_valore:
        sinistro = _lato("dashed")
    else:
        sinistro = _lato("thin", GRIGIO_BORDO)

    if colonna == colonna_valore:
        destro = _lato("thick")
    elif colonna in {colonna_descrizione, ultimo_negozio}:
        destro = _lato("medium")
    elif colonna == colonna_totale:
        destro = _lato("dashed")
    else:
        destro = _lato("thin", GRIGIO_BORDO)

    return Border(
        left=sinistro,
        right=destro,
        top=_lato(bordo_alto, GRIGIO_BORDO if bordo_alto == "thin" else NERO),
        bottom=_lato(
            bordo_basso,
            GRIGIO_BORDO if bordo_basso == "thin" else NERO,
        ),
    )


def crea_report_giacenze(dati: pd.DataFrame, percorso_output: Path) -> Path:
    """Genera il report Excel con una sola riga per prodotto."""

    richieste = {"cod.art", "descrizione", "valore"}
    mancanti = richieste - set(dati.columns)
    if mancanti:
        raise ValueError(
            "Colonne mancanti nel DataFrame: " + ", ".join(sorted(mancanti))
        )
    if dati["cod.art"].duplicated().any():
        raise ValueError("Il report richiede una sola riga per prodotto")

    negozi = [colonna for colonna in dati.columns if colonna not in richieste]
    if not negozi:
        raise ValueError("Nessuna colonna negozio disponibile")

    # La tabella deve mostrare per primi gli articoli con il valore maggiore.
    # I valori mancanti, rappresentati da "-", vengono mantenuti in fondo.
    dati = (
        dati.assign(
            _valore_ordinamento=pd.to_numeric(dati["valore"], errors="coerce")
        )
        .sort_values(
            by="_valore_ordinamento",
            ascending=False,
            na_position="last",
            kind="stable",
        )
        .drop(columns="_valore_ordinamento")
        .reset_index(drop=True)
    )

    percorso_output = Path(percorso_output)
    percorso_output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    foglio = workbook.active
    foglio.title = "GIACENZE SMARTPHONE"
    foglio.sheet_view.showGridLines = False

    riga_intestazioni, riga_totali, prima_riga = 7, 8, 9
    ultima_riga = prima_riga + len(dati) - 1
    colonna_codice, colonna_descrizione, primo_negozio = 2, 3, 4
    ultimo_negozio = primo_negozio + len(negozi) - 1
    colonna_totale, colonna_valore = ultimo_negozio + 1, ultimo_negozio + 2

    foglio.column_dimensions["A"].width = 0.86
    foglio.column_dimensions[get_column_letter(colonna_codice)].width = 10.14
    foglio.column_dimensions[get_column_letter(colonna_descrizione)].width = 75
    for colonna in range(primo_negozio, ultimo_negozio + 1):
        foglio.column_dimensions[get_column_letter(colonna)].width = 9.14
    foglio.column_dimensions[get_column_letter(colonna_totale)].width = 8
    foglio.column_dimensions[get_column_letter(colonna_valore)].width = 12.86
    foglio.row_dimensions[1].height = 4.5
    for riga in range(2, 7):
        foglio.row_dimensions[riga].height = 15.75
    foglio.row_dimensions[riga_intestazioni].height = 16.5
    foglio.row_dimensions[riga_totali].height = 19.5
    for riga in range(prima_riga, ultima_riga + 1):
        foglio.row_dimensions[riga].height = 15.75

    foglio.merge_cells(
        start_row=2,
        start_column=colonna_codice,
        end_row=4,
        end_column=colonna_valore,
    )
    titolo = foglio.cell(2, colonna_codice, "GIACENZA SMARTPHONE")
    titolo.font = Font(name="Aharoni", size=20, color=BLU)
    titolo.alignment = Alignment(horizontal="center", vertical="center")
    for riga in range(2, 5):
        for colonna in range(colonna_codice, colonna_valore + 1):
            foglio.cell(riga, colonna).border = Border(
                left=_lato("medium") if colonna == colonna_codice else Side(),
                right=_lato("medium") if colonna == colonna_valore else Side(),
                top=_lato("medium") if riga == 2 else Side(),
                bottom=_lato("medium") if riga == 4 else Side(),
            )

    intestazioni = ["COD. ART", "DESCRIZIONE", *negozi, "TOT", "VALORE"]
    for offset, testo in enumerate(intestazioni):
        colonna = colonna_codice + offset
        cella = foglio.cell(riga_intestazioni, colonna, testo)
        if colonna <= colonna_descrizione:
            riempimento, colore_font = BLU, BIANCO
        elif colonna <= ultimo_negozio:
            riempimento, colore_font = AZZURRO, BLU
        else:
            riempimento, colore_font = ROSSO, BIANCO
        cella.fill = PatternFill("solid", fgColor=riempimento)
        cella.font = Font(
            name="Calibri", size=12, bold=True, color=colore_font
        )
        cella.alignment = Alignment(
            horizontal="left" if colonna == colonna_codice else "center",
            vertical="center",
        )
        cella.border = _bordi_cella(
            colonna,
            colonna_codice,
            colonna_descrizione,
            primo_negozio,
            ultimo_negozio,
            colonna_totale,
            colonna_valore,
            bordo_alto="thick",
            bordo_basso="dashed",
        )

    foglio.merge_cells(
        start_row=riga_totali,
        start_column=colonna_codice,
        end_row=riga_totali,
        end_column=colonna_descrizione,
    )
    foglio.cell(riga_totali, colonna_codice, "TOTALE")
    for colonna in range(colonna_codice, colonna_valore + 1):
        cella = foglio.cell(riga_totali, colonna)
        cella.fill = PatternFill(
            "solid",
            fgColor=GIALLO if colonna >= colonna_totale else GRIGIO_TOTALE,
        )
        cella.font = Font(
            name="Calibri",
            size=14 if colonna == colonna_valore else 12,
            bold=True,
            color=NERO,
        )
        cella.alignment = Alignment(horizontal="center", vertical="center")
        cella.border = _bordi_cella(
            colonna,
            colonna_codice,
            colonna_descrizione,
            primo_negozio,
            ultimo_negozio,
            colonna_totale,
            colonna_valore,
            bordo_alto="dashed",
            bordo_basso="thick",
        )

    for colonna in range(primo_negozio, ultimo_negozio + 1):
        lettera = get_column_letter(colonna)
        foglio.cell(
            riga_totali,
            colonna,
            f"=SUM({lettera}{prima_riga}:{lettera}{ultima_riga})",
        )
    lettera_totale = get_column_letter(colonna_totale)
    lettera_valore = get_column_letter(colonna_valore)
    foglio.cell(
        riga_totali,
        colonna_totale,
        f"=SUM({lettera_totale}{prima_riga}:{lettera_totale}{ultima_riga})",
    )
    foglio.cell(
        riga_totali,
        colonna_valore,
        f"=SUM({lettera_valore}{prima_riga}:{lettera_valore}{ultima_riga})",
    )

    prima_lettera_negozio = get_column_letter(primo_negozio)
    ultima_lettera_negozio = get_column_letter(ultimo_negozio)
    formato_euro = '_-* #,##0 "€"_-;\\-* #,##0 "€"_-;_-* "-"?? "€"_-;_-@_-'

    for riga, (_, record) in enumerate(dati.iterrows(), start=prima_riga):
        giacenze = [record[negozio] for negozio in negozi]
        riga_negativa = any(
            isinstance(valore, (int, float)) and valore < 0
            for valore in giacenze
        )
        valori = [
            str(record["cod.art"]),
            str(record["descrizione"]),
            *giacenze,
            None,
            record["valore"],
        ]
        for offset, valore in enumerate(valori):
            colonna = colonna_codice + offset
            cella = foglio.cell(riga, colonna, valore)
            cella.font = Font(
                name="Calibri",
                size=11,
                bold=colonna >= colonna_totale,
                color=(
                    ROSSO
                    if isinstance(valore, (int, float)) and valore < 0
                    else NERO
                ),
            )
            cella.alignment = Alignment(
                horizontal="left" if colonna <= colonna_descrizione else "center",
                vertical="center",
            )
            if riga_negativa:
                cella.fill = PatternFill("solid", fgColor=ROSSO_SFOCATO)
            cella.border = _bordi_cella(
                colonna,
                colonna_codice,
                colonna_descrizione,
                primo_negozio,
                ultimo_negozio,
                colonna_totale,
                colonna_valore,
                bordo_alto=(
                    "thick"
                    if riga == prima_riga and colonna >= colonna_totale
                    else None
                    if riga == prima_riga
                    else "thin"
                ),
                bordo_basso="thick" if riga == ultima_riga else "thin",
            )
            if colonna == colonna_valore:
                cella.number_format = formato_euro

        cella_totale = foglio.cell(
            riga,
            colonna_totale,
            f"=SUM({prima_lettera_negozio}{riga}:{ultima_lettera_negozio}{riga})",
        )
        totale_numerico = sum(
            valore for valore in giacenze if isinstance(valore, (int, float))
        )
        if totale_numerico < 0:
            cella_totale.font = Font(
                name="Calibri", size=11, bold=True, color=ROSSO
            )

    foglio.cell(riga_totali, colonna_valore).number_format = formato_euro
    intervallo_dati = (
        f"{prima_lettera_negozio}{prima_riga}:"
        f"{lettera_valore}{ultima_riga}"
    )
    foglio.conditional_formatting.add(
        intervallo_dati,
        CellIsRule(
            operator="equal",
            formula=['"-"'],
            font=Font(color=GRIGIO_TRATTINO),
        ),
    )
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    try:
        workbook.save(percorso_output)
    except Exception:
        percorso_output.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()

    return percorso_output
