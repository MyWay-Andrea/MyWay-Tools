from __future__ import annotations

import html
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def identifica_tipo_dato(nome: str | Path) -> str:
    """Ricava PEZZI o PUNTI dai marcatori presenti nel nome del file."""
    percorso = Path(nome)
    nome_file = percorso.stem.casefold()
    tipi_trovati = [
        tipo
        for tipo in ("PEZZI", "PUNTI")
        if f"_{tipo.casefold()}" in nome_file
    ]

    if not tipi_trovati:
        raise ValueError(
            f"Tipo dato non riconosciuto nel nome {percorso.name!r}: "
            "deve contenere _PEZZI oppure _PUNTI"
        )
    if len(tipi_trovati) > 1:
        raise ValueError(
            f"Tipo dato ambiguo nel nome {percorso.name!r}: "
            "sono presenti sia _PEZZI sia _PUNTI"
        )

    return tipi_trovati[0]


def normalizza_testo(valore: object) -> str:
    """Normalizza testo Excel, entita HTML, spazi e maiuscole/minuscole."""
    if valore is None:
        return ""
    return " ".join(html.unescape(str(valore)).casefold().split())


def identifica_pista(
    descrizione: object,
    parole_chiave: dict[str, tuple[str, ...]],
) -> str:
    testo = normalizza_testo(descrizione)
    piste_trovate = [
        pista
        for pista, parole in parole_chiave.items()
        if all(normalizza_testo(parola) in testo for parola in parole)
    ]
    if not piste_trovate:
        raise ValueError(f"Pista non riconosciuta dalla descrizione: {descrizione!r}")
    if len(piste_trovate) > 1:
        raise ValueError(
            f"Descrizione ambigua {descrizione!r}: riconosciute {piste_trovate}"
        )
    return piste_trovate[0]


def _numero_o_zero(valore: object, *, riferimento: str) -> int | float:
    if valore is None or (isinstance(valore, str) and not valore.strip()):
        return 0
    if isinstance(valore, bool) or not isinstance(valore, (int, float)):
        raise ValueError(f"Valore non numerico in {riferimento}: {valore!r}")
    return valore


def _numero_o_vuoto(
    valore: object, *, riferimento: str
) -> int | float | None:
    if valore is None or (isinstance(valore, str) and not valore.strip()):
        return None
    return _numero_o_zero(valore, riferimento=riferimento)


def _verifica_insieme_negozi(
    negozi_trovati: list[str],
    negozi_attesi: set[str],
    *,
    contesto: str,
) -> None:
    duplicati = sorted(
        {negozio for negozio in negozi_trovati if negozi_trovati.count(negozio) > 1}
    )
    trovati = set(negozi_trovati)
    mancanti = sorted(negozi_attesi - trovati)
    extra = sorted(trovati - negozi_attesi)
    if duplicati or mancanti or extra or len(negozi_trovati) != len(negozi_attesi):
        dettagli = []
        if mancanti:
            dettagli.append(f"mancanti: {', '.join(mancanti)}")
        if extra:
            dettagli.append(f"non previsti: {', '.join(extra)}")
        if duplicati:
            dettagli.append(f"duplicati: {', '.join(duplicati)}")
        raise ValueError(f"Negozi non coerenti in {contesto} ({'; '.join(dettagli)})")


def leggi_file_raw(
    nome_file: str,
    contenuto: bytes,
    codici_negozi: dict[str, str],
    negozi_attesi: set[str],
    parole_chiave: dict[str, tuple[str, ...]],
) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(contenuto), read_only=True, data_only=True)
    try:
        foglio = workbook.active
        pista = identifica_pista(foglio["A2"].value, parole_chiave)
        valori: dict[str, int | float | None] = {}
        negozi_trovati: list[str] = []

        for colonna in range(2, foglio.max_column):
            codice_raw = foglio.cell(1, colonna).value
            codice = str(codice_raw).strip().upper() if codice_raw is not None else ""
            if codice not in codici_negozi:
                raise ValueError(
                    f"Codice negozio sconosciuto in {nome_file}, "
                    f"cella {foglio.cell(1, colonna).coordinate}: {codice_raw!r}"
                )
            negozio = codici_negozi[codice]
            negozi_trovati.append(negozio)
            valori[negozio] = _numero_o_vuoto(
                foglio.cell(2, colonna).value,
                riferimento=f"{nome_file}!{foglio.cell(2, colonna).coordinate}",
            )

        _verifica_insieme_negozi(
            negozi_trovati, negozi_attesi, contesto=nome_file
        )
        totale_calcolato = round(
            sum(valore or 0 for valore in valori.values()),
            2,
        )
        totale_file = _numero_o_zero(
            foglio.cell(2, foglio.max_column).value,
            riferimento=f"{nome_file}!{foglio.cell(2, foglio.max_column).coordinate}",
        )
        if totale_calcolato != round(totale_file, 2):
            raise ValueError(
                f"Totale non coerente in {nome_file} per {pista}: "
                f"calcolato {totale_calcolato}, presente {totale_file}"
            )
        return {
            "pista": pista,
            "valori": valori,
            "totale": totale_calcolato,
            "file_origine": nome_file,
        }
    finally:
        workbook.close()


def leggi_tutti_file_raw(
    file_remoti: list[dict[str, Any]],
    codici_negozi: dict[str, str],
    negozi_attesi: set[str],
    parole_chiave: dict[str, tuple[str, ...]],
    piste_attese: set[str],
    tipi_dato_attesi: set[str],
) -> dict[str, dict[str, dict[str, Any]]]:
    dati_gruppi: dict[str, dict[str, dict[str, Any]]] = {
        tipo: {} for tipo in tipi_dato_attesi
    }
    for file_remoto in file_remoti:
        nome_file = str(file_remoto["name"])
        print(f"  → Lettura {nome_file}")
        tipo_dato = identifica_tipo_dato(nome_file)
        if tipo_dato not in tipi_dato_attesi:
            raise ValueError(f"Tipo dato non previsto: {tipo_dato}")
        dati = leggi_file_raw(
            nome_file,
            file_remoto["content"],
            codici_negozi,
            negozi_attesi,
            parole_chiave,
        )
        pista = dati["pista"]
        dati_piste = dati_gruppi[tipo_dato]
        if pista in dati_piste:
            raise ValueError(
                f"Pista duplicata {pista} nel gruppo {tipo_dato}: "
                f"{dati_piste[pista]['file_origine']} e {nome_file}"
            )
        dati_piste[pista] = dati
        print(f"     ✓ {tipo_dato} - {pista}: totale {dati['totale']}")

    for tipo_dato, dati_piste in dati_gruppi.items():
        trovate = set(dati_piste)
        mancanti = sorted(piste_attese - trovate)
        extra = sorted(trovate - piste_attese)
        if mancanti or extra or len(dati_piste) != len(piste_attese):
            raise ValueError(
                f"Piste non coerenti per {tipo_dato}. "
                f"Mancanti: {mancanti or 'nessuna'}; "
                f"non previste: {extra or 'nessuna'}"
            )
    return dati_gruppi
