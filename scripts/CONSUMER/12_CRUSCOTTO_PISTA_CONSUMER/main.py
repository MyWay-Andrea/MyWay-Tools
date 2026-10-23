from __future__ import annotations

import re
import shutil
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import unquote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


CARTELLA_SCRIPT = Path(__file__).resolve().parent
CARTELLA_CONSUMER = CARTELLA_SCRIPT.parent

for cartella in (CARTELLA_SCRIPT, CARTELLA_CONSUMER):
    percorso = str(cartella)
    if percorso in sys.path:
        sys.path.remove(percorso)

sys.path.insert(0, str(CARTELLA_SCRIPT))
sys.path.insert(1, str(CARTELLA_CONSUMER))

from common.negozi import carica_codici_negozi, carica_negozi
from config import (
    NOME_FILE_OUTPUT,
    NOME_FOGLIO_OUTPUT,
    NOMI_MESI,
    PAROLE_CHIAVE_PISTE,
    PISTE_ATTESE,
    SHAREPOINT_HOSTNAME,
    SHAREPOINT_LIBRARY_NAME,
    SHAREPOINT_OUTPUT_FOLDER,
    SHAREPOINT_RAW_FOLDER,
    SHAREPOINT_RAW_HOSTNAME,
    SHAREPOINT_RAW_LIBRARY_NAME,
    SHAREPOINT_RAW_SITE_PATH,
    SHAREPOINT_SITE_PATH,
    TEST_RUN,
    TEST_OUTPUT_FOLDER,
    TIPI_DATO_ATTESI,
)
from leggi_file_raw import leggi_tutti_file_raw
from invia_email import invia_report

CARTELLA_SCRIPT_GENERALE = CARTELLA_CONSUMER.parent
percorso_generale = str(CARTELLA_SCRIPT_GENERALE)
while percorso_generale in sys.path:
    sys.path.remove(percorso_generale)
sys.path.insert(0, percorso_generale)
sys.modules.pop("graph_sharepoint", None)
from graph_sharepoint import GraphSharePointClient


def _cerca_file_avanzamenti(
    client_sharepoint: GraphSharePointClient,
    drive_id: str,
    percorso_cartella: str,
) -> list[dict]:
    estensioni_excel = {".xlsx", ".xlsm", ".xls"}
    file_trovati = [
        elemento
        for elemento in client_sharepoint.elenca_file_cartella(
            drive_id,
            percorso_cartella,
        )
        if "avanzamenti" in str(elemento.get("name", "")).casefold()
        and Path(str(elemento.get("name", ""))).suffix.casefold()
        in estensioni_excel
        and not str(elemento.get("name", "")).startswith("~$")
    ]
    return sorted(
        file_trovati,
        key=lambda elemento: str(elemento.get("lastModifiedDateTime", "")),
        reverse=True,
    )


def _carica_configurazione_negozi() -> tuple[list[str], dict[str, str]]:
    negozi = [negozio.strip().upper() for negozio in carica_negozi()]
    codici = carica_codici_negozi()
    duplicati = sorted({negozio for negozio in negozi if negozi.count(negozio) > 1})
    valori_codici = list(codici.values())
    duplicati_mapping = sorted(
        {negozio for negozio in valori_codici if valori_codici.count(negozio) > 1}
    )
    mancanti_mapping = sorted(set(negozi) - set(valori_codici))
    extra_mapping = sorted(set(valori_codici) - set(negozi))

    if duplicati or duplicati_mapping or mancanti_mapping or extra_mapping:
        raise ValueError(
            "Configurazione negozi non coerente in negozi.json. "
            f"Duplicati nella lista: {duplicati or 'nessuno'}; "
            f"duplicati nei codici: {duplicati_mapping or 'nessuno'}; "
            f"senza codice: {mancanti_mapping or 'nessuno'}; "
            f"non presenti nella lista: {extra_mapping or 'nessuno'}"
        )
    return negozi, codici


def _estrai_periodo(file_raw: list[dict]) -> tuple[datetime, datetime]:
    periodi: set[tuple[str, str]] = set()
    pattern = re.compile(r"(\d{2}_\d{2}_\d{4}).*?(\d{2}_\d{2}_\d{4})")
    for file_remoto in file_raw:
        nome_file = str(file_remoto["name"])
        corrispondenza = pattern.search(Path(nome_file).stem)
        if not corrispondenza:
            raise ValueError(f"Periodo non riconosciuto nel nome: {nome_file}")
        periodi.add(corrispondenza.groups())
    if len(periodi) != 1:
        raise ValueError(f"I file RAW appartengono a periodi diversi: {sorted(periodi)}")
    inizio_testo, fine_testo = periodi.pop()
    return (
        datetime.strptime(inizio_testo, "%d_%m_%Y"),
        datetime.strptime(fine_testo, "%d_%m_%Y"),
    )


def _estrai_periodo_report(nome_file: str) -> tuple[datetime, datetime] | None:
    pattern = re.compile(
        r"^TRACCIAMENTO_CONSUMER_(\d{2}_\d{2}_\d{4})_"
        r"(\d{2}_\d{2}_\d{4})\.xlsx$",
        re.IGNORECASE,
    )
    corrispondenza = pattern.fullmatch(Path(nome_file).name)
    if not corrispondenza:
        return None
    try:
        date = tuple(
            datetime.strptime(data, "%d_%m_%Y")
            for data in corrispondenza.groups()
        )
    except ValueError:
        return None
    return date


def trova_report_precedente(
    client_sharepoint: GraphSharePointClient,
    drive_id: str,
    periodo_corrente: tuple[datetime, datetime],
) -> dict | None:
    """Trova il report con la data nel nome immediatamente precedente."""
    data_corrente = max(periodo_corrente)
    cartella_attesa = "/" + SHAREPOINT_OUTPUT_FOLDER.strip("/").casefold() + "/"
    candidati: list[tuple[datetime, str, dict]] = []

    for elemento in client_sharepoint.cerca_file(
        drive_id,
        "TRACCIAMENTO_CONSUMER",
    ):
        periodo = _estrai_periodo_report(str(elemento.get("name", "")))
        if periodo is None:
            continue

        if not elemento.get("parentReference", {}).get("path"):
            elemento = client_sharepoint.ottieni_metadati_elemento(
                drive_id,
                elemento["id"],
            )

        percorso_genitore = unquote(
            str(elemento.get("parentReference", {}).get("path", ""))
        ).replace("\\", "/").casefold()
        percorso_genitore = percorso_genitore.rstrip("/") + "/"
        if cartella_attesa not in percorso_genitore:
            continue

        data_riferimento = max(periodo)
        if data_riferimento >= data_corrente:
            continue
        candidati.append(
            (
                data_riferimento,
                str(elemento.get("lastModifiedDateTime", "")),
                elemento,
            )
        )

    if not candidati:
        return None
    return max(candidati, key=lambda candidato: (candidato[0], candidato[1]))[2]


def _numero_report_precedente(valore: object, *, riferimento: str) -> int | float:
    if isinstance(valore, bool) or not isinstance(valore, (int, float)):
        raise ValueError(
            f"Totale non numerico nel report precedente, {riferimento}: {valore!r}"
        )
    return valore


def leggi_totali_report_precedente(
    nome_file: str,
    contenuto: bytes,
) -> dict[str, dict[str, int | float]]:
    """Legge i totali PEZZI/PUNTI identificando intestazioni e piste."""
    workbook = load_workbook(BytesIO(contenuto), read_only=True, data_only=True)
    try:
        foglio = (
            workbook[NOME_FOGLIO_OUTPUT]
            if NOME_FOGLIO_OUTPUT in workbook.sheetnames
            else workbook.active
        )
        risultati: dict[str, dict[str, int | float]] = {}
        for tipo_dato in TIPI_DATO_ATTESI:
            righe_sezione = [
                riga
                for riga in range(1, foglio.max_row + 1)
                if str(foglio.cell(riga, 1).value or "").strip().upper()
                == tipo_dato
            ]
            if len(righe_sezione) != 1:
                raise ValueError(
                    f"Sezione {tipo_dato} non univoca in {nome_file}"
                )
            riga_titolo = righe_sezione[0]
            colonne_totale = [
                colonna
                for colonna in range(1, foglio.max_column + 1)
                if str(foglio.cell(riga_titolo, colonna).value or "")
                .strip()
                .upper()
                == "TOT."
            ]
            if len(colonne_totale) != 1:
                raise ValueError(
                    f"Colonna TOT. non univoca nella sezione {tipo_dato} "
                    f"di {nome_file}"
                )
            colonna_totale = colonne_totale[0]
            totali_piste: dict[str, int | float] = {}
            for riga in range(riga_titolo + 1, foglio.max_row + 1):
                pista = str(foglio.cell(riga, 1).value or "").strip().upper()
                if pista in TIPI_DATO_ATTESI:
                    break
                if pista not in PISTE_ATTESE:
                    continue
                if pista in totali_piste:
                    raise ValueError(
                        f"Pista duplicata {pista} nella sezione {tipo_dato} "
                        f"di {nome_file}"
                    )
                totali_piste[pista] = _numero_report_precedente(
                    foglio.cell(riga, colonna_totale).value,
                    riferimento=(
                        f"{nome_file}!"
                        f"{foglio.cell(riga, colonna_totale).coordinate}"
                    ),
                )
            if set(totali_piste) != set(PISTE_ATTESE):
                raise ValueError(
                    f"Piste non coerenti nella sezione {tipo_dato} di {nome_file}"
                )
            risultati[tipo_dato] = totali_piste
        return risultati
    finally:
        workbook.close()


def _crea_foglio_report(
    workbook: Workbook,
    negozi: list[str],
    codici_negozi: dict[str, str],
):
    """Crea il cruscotto direttamente, senza dipendere da un file modello."""
    foglio = workbook.active
    foglio.title = NOME_FOGLIO_OUTPUT
    foglio.sheet_view.showGridLines = False

    rosso_sfondo = "FF8B0300"
    rosso_testo = "FFC00000"
    rosso_bordo = "FFE05451"
    grigio_totale = "FF595959"
    grigio_chiaro = "FFE7E6E6"
    bianco = "FFFFFFFF"
    nero = "FF000000"
    colore_bordo_automatico = Color(auto=True)
    bordo_sottile = Side(style="thin", color=colore_bordo_automatico)
    bordo_spesso = Side(style="thick", color=colore_bordo_automatico)
    bordo_tratteggiato = Side(style="dashed", color=colore_bordo_automatico)
    bordo_rosso = Side(style="medium", color=rosso_bordo)
    nessun_bordo = Side()
    riempimento_titolo = PatternFill(
        "solid", fgColor=rosso_sfondo, bgColor=nero
    )
    riempimento_titolo_totale = PatternFill("solid", fgColor=grigio_totale)
    riempimento_totale = PatternFill("solid", fgColor=grigio_chiaro)
    font_titolo = Font(name="Calibri", size=11, bold=True, color=bianco)
    font_titolo_sintesi = Font(
        name="Calibri", size=12, bold=True, color=bianco
    )
    font_sezione = Font(
        name="Aptos Narrow",
        size=14,
        bold=True,
        color=rosso_testo,
        family=2,
        scheme="minor",
    )
    font_etichette = Font(
        name="Aptos Narrow",
        size=11,
        color=Color(theme=1),
        family=2,
        scheme="minor",
    )
    font_dati = Font(name="Calibri", size=11, color=Color(theme=1), family=2)
    font_totale = Font(
        name="Calibri", size=12, bold=True, color=nero, family=2
    )

    codice_per_negozio = {
        negozio: codice for codice, negozio in codici_negozi.items()
    }
    codici_mancanti = [
        negozio for negozio in negozi if negozio not in codice_per_negozio
    ]
    if codici_mancanti:
        raise ValueError(
            "Codice non disponibile per i negozi: " + ", ".join(codici_mancanti)
        )

    foglio.column_dimensions["A"].width = 15.55
    for colonna, negozio in enumerate(negozi, start=2):
        codice = codice_per_negozio[negozio]
        foglio.column_dimensions[get_column_letter(colonna)].width = max(
            7,
            min(len(codice) + 2, 10),
        )
    colonna_totale = len(negozi) + 2
    colonna_variazione = colonna_totale + 1
    foglio.column_dimensions[get_column_letter(colonna_totale)].width = 10
    foglio.column_dimensions[get_column_letter(colonna_variazione)].width = 12

    foglio.row_dimensions[1].height = 15.6
    foglio.merge_cells("B1:E1")
    foglio["A1"] = "Aggiornato il:"
    foglio["A1"].font = Font(name="Calibri", size=12, bold=True, family=2)
    foglio["A1"].alignment = Alignment(horizontal="right")
    foglio["B1"] = datetime.now()
    foglio["B1"].number_format = "dd/mm/yyyy hh:mm"
    foglio["B1"].font = Font(name="Calibri", size=12, bold=True, family=2)
    foglio["B1"].alignment = Alignment(horizontal="left")

    colonne = {negozio: indice + 2 for indice, negozio in enumerate(negozi)}
    intestazioni = [
        *(codice_per_negozio[negozio] for negozio in negozi),
        "TOT.",
        "Var D-1",
    ]
    for riga_titolo, tipo_dato in ((3, "PEZZI"), (11, "PUNTI")):
        foglio.row_dimensions[riga_titolo].height = 19.2
        foglio.cell(riga_titolo, 1, tipo_dato)
        for colonna, intestazione in enumerate(intestazioni, start=2):
            foglio.cell(riga_titolo, colonna, intestazione)

        for colonna in range(1, colonna_variazione + 1):
            cella = foglio.cell(riga_titolo, colonna)
            if colonna == 1:
                cella.font = font_sezione
            else:
                if colonna == colonna_totale:
                    cella.fill = riempimento_titolo_totale
                elif colonna == colonna_variazione:
                    cella.fill = riempimento_titolo_totale
                else:
                    cella.fill = riempimento_titolo
                cella.font = (
                    font_titolo_sintesi
                    if colonna in (colonna_totale, colonna_variazione)
                    else font_titolo
                )
            cella.alignment = Alignment(horizontal="center", vertical="center")
            if colonna == 1:
                cella.border = Border(
                    left=bordo_spesso,
                    right=bordo_rosso,
                    top=bordo_spesso,
                    bottom=bordo_spesso,
                )
            elif colonna == 2:
                cella.border = Border(
                    left=bordo_rosso,
                    right=bordo_rosso,
                    top=bordo_spesso,
                    bottom=nessun_bordo,
                )
            else:
                cella.border = Border(
                    left=bordo_rosso,
                    right=(
                        bordo_spesso
                        if colonna == colonna_variazione
                        else bordo_rosso
                    ),
                    top=bordo_spesso,
                    bottom=bordo_tratteggiato,
                )

        for scarto, pista in enumerate(PISTE_ATTESE, start=1):
            riga = riga_titolo + scarto
            if scarto in (1, len(PISTE_ATTESE)):
                foglio.row_dimensions[riga].height = 15
            foglio.cell(riga, 1, pista)
            for colonna in range(1, colonna_variazione + 1):
                cella = foglio.cell(riga, colonna)
                cella.font = (
                    font_etichette
                    if colonna == 1
                    else font_totale
                    if colonna in (colonna_totale, colonna_variazione)
                    else font_dati
                )
                cella.alignment = Alignment(
                    horizontal="left" if colonna == 1 else "center",
                    vertical=None if colonna == 1 else "center",
                    indent=1 if colonna == 1 else 0,
                )
                ultima_riga = scarto == len(PISTE_ATTESE)
                if colonna == 1:
                    cella.border = Border(
                        left=bordo_spesso,
                        right=bordo_spesso,
                        top=nessun_bordo,
                        bottom=bordo_spesso if ultima_riga else nessun_bordo,
                    )
                else:
                    if colonna == colonna_totale:
                        cella.fill = riempimento_totale
                    elif colonna == colonna_variazione:
                        cella.fill = riempimento_totale
                        cella.number_format = "@"
                    cella.border = Border(
                        left=(
                            bordo_spesso
                            if colonna == colonna_totale
                            else bordo_sottile
                        ),
                        right=(
                            bordo_spesso
                            if colonna == colonna_variazione
                            else bordo_rosso
                        ),
                        top=bordo_tratteggiato,
                        bottom=(
                            bordo_spesso if ultima_riga else bordo_tratteggiato
                        ),
                    )
    return foglio, colonne, colonna_totale, colonna_variazione


def _formatta_variazione(valore: int | float | None) -> str | None:
    if valore is None:
        return None
    valore = round(valore, 2)
    if valore == 0:
        return "0"
    valore_assoluto = f"{abs(valore):.2f}".rstrip("0").rstrip(".")
    valore_assoluto = valore_assoluto.replace(".", ",")
    segno = "+" if valore > 0 else "-"
    return f"{segno}{valore_assoluto}"


def genera_report(
    dati_gruppi: dict,
    negozi: list[str],
    codici_negozi: dict[str, str],
    periodo: tuple[datetime, datetime],
    cartella_output: Path,
    totali_precedenti: dict[str, dict[str, int | float]] | None = None,
) -> Path:
    inizio, fine = periodo
    cartella_output = Path(cartella_output)
    cartella_output.mkdir(parents=True, exist_ok=True)
    percorso_output = cartella_output / NOME_FILE_OUTPUT.format(
        inizio=inizio.strftime("%d_%m_%Y"),
        fine=fine.strftime("%d_%m_%Y"),
    )
    workbook = Workbook()
    try:
        foglio, colonne, colonna_totale, colonna_variazione = (
            _crea_foglio_report(workbook, negozi, codici_negozi)
        )

        righe_iniziali = {"PEZZI": 4, "PUNTI": 12}
        for tipo_dato, dati_piste in dati_gruppi.items():
            riga_iniziale = righe_iniziali[tipo_dato]
            righe_piste = {
                str(foglio.cell(riga, 1).value).strip().upper(): riga
                for riga in range(riga_iniziale, riga_iniziale + 5)
            }
            if set(righe_piste) != set(PISTE_ATTESE):
                raise ValueError(
                    f"Piste del blocco {tipo_dato} non coerenti: "
                    f"{sorted(righe_piste)}"
                )
            for pista, dati in dati_piste.items():
                riga = righe_piste[pista]
                for negozio, valore in dati["valori"].items():
                    cella = foglio.cell(riga, colonne[negozio], valore)
                    if tipo_dato == "PUNTI":
                        cella.number_format = "General"
                cella_totale = foglio.cell(riga, colonna_totale, dati["totale"])
                valore_precedente = (
                    totali_precedenti[tipo_dato][pista]
                    if totali_precedenti is not None
                    else None
                )
                variazione = (
                    round(dati["totale"] - valore_precedente, 2)
                    if valore_precedente is not None
                    else None
                )
                cella_variazione = foglio.cell(
                    riga,
                    colonna_variazione,
                    _formatta_variazione(variazione),
                )
                cella_variazione.number_format = "@"
                if variazione is not None:
                    cella_variazione.font = Font(
                        name="Calibri",
                        size=12,
                        bold=True,
                        color=(
                            "FF008000"
                            if variazione > 0
                            else "FFC00000"
                            if variazione < 0
                            else "FF000000"
                        ),
                        family=2,
                    )
                if tipo_dato == "PUNTI":
                    cella_totale.number_format = "General"
        workbook.save(percorso_output)
    except Exception:
        workbook.close()
        percorso_output.unlink(missing_ok=True)
        raise
    else:
        workbook.close()
    return percorso_output


def percorso_archivio_sharepoint(data_esecuzione: datetime) -> str:
    """Costruisce il percorso remoto base/anno/mese del report."""
    mese = f"{data_esecuzione.month:02d}_{NOMI_MESI[data_esecuzione.month]}"
    return "/".join(
        (
            SHAREPOINT_OUTPUT_FOLDER.strip("/"),
            str(data_esecuzione.year),
            mese,
        )
    )


def elimina_file_raw(
    client_sharepoint: GraphSharePointClient,
    drive_id: str,
    file_raw: list[dict],
) -> None:
    """Elimina da SharePoint esclusivamente gli elementi Graph elaborati."""
    elementi_verificati: list[tuple[str, str, str]] = []
    for file_remoto in file_raw:
        item_id = str(file_remoto.get("id", "")).strip()
        etag = str(file_remoto.get("eTag", "")).strip()
        nome_file = str(file_remoto.get("name", "")).strip()
        drive_elemento = str(
            file_remoto.get("parentReference", {}).get("driveId", "")
        ).strip()
        if not item_id or not etag or not nome_file:
            raise ValueError(f"Metadati Graph incompleti per il RAW: {nome_file!r}")
        if drive_elemento and drive_elemento != drive_id:
            raise ValueError(
                f"Eliminazione bloccata per {nome_file}: drive Graph non coerente"
            )
        elementi_verificati.append((item_id, etag, nome_file))

    for item_id, etag, nome_file in elementi_verificati:
        client_sharepoint.elimina_file(drive_id, item_id, etag)
        print(f"  ✓ Eliminato da SharePoint: {nome_file}")


def main() -> None:
    fase = "avvio"
    try:
        fase = "controllo configurazione negozi"
        print(f"\n→ {fase.capitalize()}...")
        negozi, codici_negozi = _carica_configurazione_negozi()
        print(f"✓ Configurazione negozi valida: {len(negozi)} negozi")

        fase = "connessione a SharePoint e ricerca file RAW"
        print(f"\n→ {fase.capitalize()}...")
        client_sharepoint = GraphSharePointClient()
        sito_raw = client_sharepoint.trova_sito(
            SHAREPOINT_RAW_HOSTNAME,
            SHAREPOINT_RAW_SITE_PATH,
        )
        raccolta_raw = client_sharepoint.trova_raccolta_documenti(
            sito_raw["id"],
            SHAREPOINT_RAW_LIBRARY_NAME,
        )
        drive_id_raw = raccolta_raw["id"]
        file_raw = _cerca_file_avanzamenti(
            client_sharepoint,
            drive_id_raw,
            SHAREPOINT_RAW_FOLDER,
        )
        if not file_raw:
            raise FileNotFoundError(
                "Nessun file AVANZAMENTI trovato nella cartella SharePoint "
                f"{SHAREPOINT_RAW_FOLDER}"
            )
        print(f"✓ Sito SharePoint: {sito_raw['displayName']}")
        print(f"✓ Raccolta documenti: {raccolta_raw['name']}")
        print(f"✓ File RAW remoti trovati: {len(file_raw)}")

        fase = "download file RAW in memoria"
        print(f"\n→ {fase.capitalize()}...")
        for file_remoto in file_raw:
            file_remoto["content"] = client_sharepoint.scarica_file(
                drive_id_raw,
                file_remoto["id"],
            )
            print(
                f"  ✓ Scaricato in memoria: {file_remoto['name']} "
                f"({len(file_remoto['content'])} byte)"
            )

        fase = "lettura e validazione file RAW"
        print(f"\n→ {fase.capitalize()}...")
        dati_piste = leggi_tutti_file_raw(
            file_raw,
            codici_negozi,
            set(negozi),
            PAROLE_CHIAVE_PISTE,
            set(PISTE_ATTESE),
            set(TIPI_DATO_ATTESI),
        )
        print(
            f"✓ File RAW validati: {len(file_raw)} file, "
            f"{len(dati_piste['PEZZI'])} piste PEZZI e "
            f"{len(dati_piste['PUNTI'])} piste PUNTI"
        )

        fase = "controllo periodo"
        print(f"\n→ {fase.capitalize()}...")
        periodo = _estrai_periodo(file_raw)
        print(f"✓ Periodo: {periodo[0]:%d/%m/%Y} - {periodo[1]:%d/%m/%Y}")

        fase = "ricerca report precedente su SharePoint"
        print(f"\n→ {fase.capitalize()}...")
        sito_output = client_sharepoint.trova_sito(
            SHAREPOINT_HOSTNAME,
            SHAREPOINT_SITE_PATH,
        )
        raccolta_output = client_sharepoint.trova_raccolta_documenti(
            sito_output["id"],
            SHAREPOINT_LIBRARY_NAME,
        )
        report_precedente = trova_report_precedente(
            client_sharepoint,
            raccolta_output["id"],
            periodo,
        )
        totali_precedenti = None
        if report_precedente is None:
            print("✓ Nessun report precedente: Var D-1 resterà vuota")
        else:
            contenuto_precedente = client_sharepoint.scarica_file(
                raccolta_output["id"],
                report_precedente["id"],
            )
            totali_precedenti = leggi_totali_report_precedente(
                report_precedente["name"],
                contenuto_precedente,
            )
            print(f"✓ Report precedente: {report_precedente['name']}")

        with TemporaryDirectory(prefix="cruscotto_consumer_") as cartella_temporanea:
            fase = "creazione report Excel temporaneo"
            print(f"\n→ {fase.capitalize()}...")
            output = genera_report(
                dati_piste,
                negozi,
                codici_negozi,
                periodo,
                Path(cartella_temporanea),
                totali_precedenti,
            )
            print(f"✓ Report temporaneo creato: {output.name}")

            if TEST_RUN:
                cartella_test = TEST_OUTPUT_FOLDER
                cartella_test.mkdir(parents=True, exist_ok=True)
                output_test = cartella_test / output.name
                shutil.copy2(output, output_test)
                print(f"✓ Report test salvato: {output_test}")
                print("✓ TEST RUN: upload SharePoint, email ed eliminazione RAW saltati")
            else:
                fase = "caricamento report su SharePoint"
                print(f"\n→ {fase.capitalize()}...")
                percorso_remoto = percorso_archivio_sharepoint(datetime.now())
                cartella_remota = client_sharepoint.crea_percorso_cartelle(
                    raccolta_output["id"],
                    percorso_remoto,
                )
                file_caricato = client_sharepoint.carica_file(
                    raccolta_output["id"],
                    cartella_remota["id"],
                    output,
                )
                print(f"✓ Report caricato su SharePoint: {file_caricato['webUrl']}")

                fase = "preparazione e invio email"
                print(f"\n→ {fase.capitalize()}...")
                invia_report(output, periodo)
                print("✓ Invio email completato")

                fase = "eliminazione file RAW elaborati"
                print(f"\n→ {fase.capitalize()}...")
                elimina_file_raw(client_sharepoint, drive_id_raw, file_raw)
                print(f"✓ File RAW eliminati: {len(file_raw)}")

        print("✓ File temporaneo locale eliminato")

        print("\n✓ Pipeline completata senza errori")
    except Exception as errore:
        print(f"\n✗ Errore durante la fase: {fase}")
        print(f"Dettaglio: {errore}")
        raise


if __name__ == "__main__":
    main()
