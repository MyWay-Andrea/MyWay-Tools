from __future__ import annotations

import sys
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import main


def _dati(totale_base: int) -> dict:
    return {
        tipo: {
            pista: {
                "pista": pista,
                "valori": {"NEGOZIO_TEST": totale_base + indice},
                "totale": totale_base + indice,
                "file_origine": "test.xlsx",
            }
            for indice, pista in enumerate(main.PISTE_ATTESE_PER_TIPO[tipo])
        }
        for tipo in main.TIPI_DATO_ATTESI
    }


def _dettaglio_smartphone() -> list[dict]:
    return [
        {
            "descrizione": f"Voce {indice}",
            "valori": {"NEGOZIO_TEST": indice},
            "totale": indice,
        }
        for indice in range(1, 13)
    ]


class ClientSharePointFinto:
    def cerca_file(self, drive_id: str, testo: str) -> list[dict]:
        cartella = main.SHAREPOINT_OUTPUT_FOLDER
        return [
            {
                "id": "recente",
                "name": "TRACCIAMENTO_CONSUMER_01_10_2026_08_10_2026.xlsx",
                "lastModifiedDateTime": "2026-10-08T18:00:00Z",
                "parentReference": {"path": f"/drives/{drive_id}/root:/{cartella}/2026/10_OTTOBRE"},
            },
            {
                "id": "vecchio",
                "name": "TRACCIAMENTO_CONSUMER_01_10_2026_07_10_2026.xlsx",
                "lastModifiedDateTime": "2026-10-07T18:00:00Z",
                "parentReference": {"path": f"/drives/{drive_id}/root:/{cartella}/2026/10_OTTOBRE"},
            },
            {
                "id": "corrente",
                "name": "TRACCIAMENTO_CONSUMER_01_10_2026_09_10_2026.xlsx",
                "lastModifiedDateTime": "2026-10-09T18:00:00Z",
                "parentReference": {"path": f"/drives/{drive_id}/root:/{cartella}/2026/10_OTTOBRE"},
            },
        ]


class TestVariazioneD1(unittest.TestCase):
    def test_intestazioni_con_codici_negozio_compatti(self) -> None:
        workbook = Workbook()
        try:
            foglio, _, _, _ = main._crea_foglio_report(
                workbook,
                ["CARPI", "SASSUOLO_M.POLO"],
                {"CAR": "CARPI", "SAS": "SASSUOLO_M.POLO"},
            )
            self.assertEqual(foglio["B3"].value, "CAR")
            self.assertEqual(foglio["C3"].value, "SAS")
            self.assertEqual(foglio.column_dimensions["B"].width, 7)
            self.assertEqual(foglio.column_dimensions["C"].width, 7)
            self.assertIn("B1:E1", {str(area) for area in foglio.merged_cells.ranges})
        finally:
            workbook.close()

    def test_seleziona_data_precedente_piu_recente(self) -> None:
        trovato = main.trova_report_precedente(
            ClientSharePointFinto(),
            "drive-test",
            (datetime(2026, 10, 1), datetime(2026, 10, 9)),
        )
        self.assertIsNotNone(trovato)
        self.assertEqual(trovato["id"], "recente")

    def test_genera_e_rilegge_variazioni(self) -> None:
        periodo_precedente = (datetime(2099, 10, 1), datetime(2099, 10, 8))
        periodo_corrente = (datetime(2099, 10, 1), datetime(2099, 10, 9))
        cartella = SCRIPT_DIR / "output"
        file_test = [
            cartella / "TRACCIAMENTO_CONSUMER_01_10_2099_08_10_2099.xlsx",
            cartella / "TRACCIAMENTO_CONSUMER_01_10_2099_09_10_2099.xlsx",
        ]
        try:
            precedente = main.genera_report(
                _dati(10),
                ["NEGOZIO_TEST"],
                {"TST": "NEGOZIO_TEST"},
                periodo_precedente,
                cartella,
            )
            totali_precedenti = main.leggi_totali_report_precedente(
                precedente.name,
                precedente.read_bytes(),
            )
            dati_correnti = _dati(15)
            dati_correnti["PEZZI"]["MOBILE"]["valori"]["NEGOZIO_TEST"] = 8
            dati_correnti["PEZZI"]["MOBILE"]["totale"] = 8
            dati_correnti["PUNTI"]["MOBILE"]["valori"]["NEGOZIO_TEST"] = 14.7
            dati_correnti["PUNTI"]["MOBILE"]["totale"] = 14.7
            corrente = main.genera_report(
                dati_correnti,
                ["NEGOZIO_TEST"],
                {"TST": "NEGOZIO_TEST"},
                periodo_corrente,
                cartella,
                totali_precedenti,
                _dettaglio_smartphone(),
            )
            workbook = load_workbook(corrente, data_only=False)
            try:
                foglio = workbook[main.NOME_FOGLIO_OUTPUT]
                self.assertEqual(foglio["B3"].value, "TST")
                self.assertEqual(foglio["C3"].value, "TOT.")
                self.assertEqual(foglio["D3"].value, "Var D-1")
                self.assertEqual(foglio["D4"].value, "-2")
                self.assertEqual(foglio["D4"].number_format, "@")
                self.assertEqual(foglio["D4"].font.color.rgb, "FFC00000")
                self.assertEqual(foglio["C4"].font.sz, 12)
                self.assertEqual(foglio["D4"].font.sz, 12)
                self.assertEqual(foglio["D12"].value, "+4,7")
                self.assertEqual(foglio["D12"].font.color.rgb, "FF008000")
                self.assertEqual(foglio["C4"].fill.fgColor.rgb, "FFE7E6E6")
                self.assertEqual(foglio["D4"].fill.fgColor.rgb, "FFE7E6E6")
                dettaglio = workbook[main.NOME_FOGLIO_DETTAGLIO_SMARTPHONE]
                self.assertEqual(dettaglio["B3"].value, "DETTAGLIO TELEFONI")
                self.assertEqual(dettaglio["C3"].value, "TST")
                self.assertEqual(dettaglio["D3"].value, "TOT.")
                self.assertEqual(dettaglio["C4"].value, "=+SUM(C5:C16)")
                self.assertEqual(dettaglio["B5"].value, "Voce 1")
                self.assertEqual(dettaglio["C16"].value, 12)
            finally:
                workbook.close()
        finally:
            for percorso in file_test:
                percorso.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
