from pathlib import Path
from io import BytesIO
import os
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from colorama import Fore, init
init(autoreset=True)


SCRIPT_DIR = Path(__file__).resolve().parents[3]
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from graph_sharepoint import GraphSharePointClient

SHAREPOINT_HOSTNAME = os.getenv("SHAREPOINT_HOSTNAME", "").strip()
SHAREPOINT_SCAMBIO_SITE_PATH = os.getenv("SHAREPOINT_SCAMBIO_SITE_PATH", "").strip()
SHAREPOINT_SCAMBIO_LIBRARY_NAME = os.getenv("SHAREPOINT_SCAMBIO_LIBRARY_NAME", "").strip()
SHAREPOINT_MYWAY_TOOLS_SITE_PATH = os.getenv("SHAREPOINT_MYWAY_TOOLS_SITE_PATH", "").strip()
SHAREPOINT_MYWAY_TOOLS_LIBRARY_NAME = os.getenv("SHAREPOINT_MYWAY_TOOLS_LIBRARY_NAME", "").strip()
SHAREPOINT_RAW_FOLDER = os.getenv("SHAREPOINT_RAW_FOLDER", "").strip(" /")
PATH_DATABASE = os.getenv("SHAREPOINT_SCAMBIO_PARQUET_FOLDER", "").strip(" /")
PATH_BASI_DATI = os.getenv("SHAREPOINT_CONSUMER_BASI_DATI_FOLDER", "").strip(" /")
_GRAPH_CLIENT = None
_GRAPH_DRIVE_ID = None
_GRAPH_RAW_DRIVE_ID = None


def _drive_sharepoint():
    global _GRAPH_CLIENT, _GRAPH_DRIVE_ID
    configurazione = {
        "SHAREPOINT_HOSTNAME": SHAREPOINT_HOSTNAME,
        "SHAREPOINT_SCAMBIO_SITE_PATH": SHAREPOINT_SCAMBIO_SITE_PATH,
        "SHAREPOINT_SCAMBIO_LIBRARY_NAME": SHAREPOINT_SCAMBIO_LIBRARY_NAME,
        "SHAREPOINT_SCAMBIO_PARQUET_FOLDER": PATH_DATABASE,
        "SHAREPOINT_CONSUMER_BASI_DATI_FOLDER": PATH_BASI_DATI,
    }
    mancanti = [nome for nome, valore in configurazione.items() if not valore]
    if mancanti:
        raise ValueError("Configurazione SharePoint Pedonalita incompleta: " + ", ".join(mancanti))
    if _GRAPH_CLIENT is None:
        _GRAPH_CLIENT = GraphSharePointClient()
        sito = _GRAPH_CLIENT.trova_sito(SHAREPOINT_HOSTNAME, SHAREPOINT_SCAMBIO_SITE_PATH)
        raccolta = _GRAPH_CLIENT.trova_raccolta_documenti(
            sito["id"], SHAREPOINT_SCAMBIO_LIBRARY_NAME
        )
        _GRAPH_DRIVE_ID = str(raccolta["id"])
    return _GRAPH_CLIENT, _GRAPH_DRIVE_ID


def _drive_raw_sharepoint():
    global _GRAPH_CLIENT, _GRAPH_RAW_DRIVE_ID
    configurazione = {
        "SHAREPOINT_HOSTNAME": SHAREPOINT_HOSTNAME,
        "SHAREPOINT_MYWAY_TOOLS_SITE_PATH": SHAREPOINT_MYWAY_TOOLS_SITE_PATH,
        "SHAREPOINT_MYWAY_TOOLS_LIBRARY_NAME": SHAREPOINT_MYWAY_TOOLS_LIBRARY_NAME,
        "SHAREPOINT_RAW_FOLDER": SHAREPOINT_RAW_FOLDER,
    }
    mancanti = [nome for nome, valore in configurazione.items() if not valore]
    if mancanti:
        raise ValueError("Configurazione RAW Pedonalita incompleta: " + ", ".join(mancanti))
    if _GRAPH_CLIENT is None:
        _GRAPH_CLIENT = GraphSharePointClient()
    if _GRAPH_RAW_DRIVE_ID is None:
        sito = _GRAPH_CLIENT.trova_sito(
            SHAREPOINT_HOSTNAME, SHAREPOINT_MYWAY_TOOLS_SITE_PATH
        )
        raccolta = _GRAPH_CLIENT.trova_raccolta_documenti(
            sito["id"], SHAREPOINT_MYWAY_TOOLS_LIBRARY_NAME
        )
        _GRAPH_RAW_DRIVE_ID = str(raccolta["id"])
    return _GRAPH_CLIENT, _GRAPH_RAW_DRIVE_ID


def trova_raw_pedonalita() -> dict:
    client, drive_id = _drive_raw_sharepoint()
    candidati = [
        file for file in client.elenca_file_cartella(drive_id, SHAREPOINT_RAW_FOLDER)
        if "pedonalit" in str(file.get("name", "")).casefold()
        and str(file.get("name", "")).casefold().endswith((".xlsx", ".xls"))
    ]
    if not candidati:
        raise FileNotFoundError(
            f"RAW Pedonalita non trovato su SharePoint: {SHAREPOINT_RAW_FOLDER}"
        )
    candidati.sort(key=lambda file: str(file.get("lastModifiedDateTime", "")), reverse=True)
    scelto = candidati[0]
    print(f"RAW Pedonalita SharePoint: {scelto['name']}")
    return {**scelto, "content": client.scarica_file(drive_id, scelto["id"])}


def elimina_raw_pedonalita(file_remoto: dict) -> None:
    client, drive_id = _drive_raw_sharepoint()
    client.elimina_file(
        drive_id, str(file_remoto["id"]), str(file_remoto.get("eTag", ""))
    )
    print(Fore.GREEN + f"RAW eliminato da SharePoint: {file_remoto['name']}")

#FUNZIONI COMUNI
#---------------------------------------------------------------

#FUNZIONE PER TROVARE E CONTROLLARE LA CARTELLA "RAW_PEDONALITA"
#--------------------------------------------------------------
def trova_file(path_cartelle):
    ESTENSIONI = ["*.parquet", "*.xlsx"]
    ped_DB = None
    pedonalita_lavorata  = []
    
    file_trovati = [f for ext in ESTENSIONI for f in path_cartelle.glob(ext)]
    for file in file_trovati:
        nome = file.name.lower()
        if "giorgio" in nome:
            pedonalita_lavorata.append(file)
        elif "pedonalit" in nome:
            ped_DB = file
        
    
    raw_ped = pedonalita_lavorata[-1] if pedonalita_lavorata else None
    
    if ped_DB is None and raw_ped:
        print(f" 📁 Pedonalita Preparata:    {raw_ped}\n")
    elif ped_DB and raw_ped is None:
        print(f" 📁 Pedonalita DataBase:    {ped_DB}\n")
    elif ped_DB and raw_ped:
        print(f" 📁 Pedonalita DataBase:    {ped_DB}")
        print(f" 📁 Pedonalita Preparata:    {raw_ped}\n")
    else:
        print(" ❌ Nessun file di pedonalità trovato nella cartella Raw Folder.")
        
    return ped_DB, raw_ped, 


#---------------------------------------------------------------

#FILE:
#PREPARA PEDONALITA
#-------------------------------------------------------------------------


NEGOZIO_CONDITION = {
    "28657.0W021" : "Fidenza",
    "28657.0W022" : "Piacenza Galassia",
    "28657.0W100" : "Parma Torri",
    "28657.0W101" : "Parma Eurosia",
    "28657.0W409" : "Modena",
    "28657.0W728" : "Sassuolo",
    "28657.0X132" : "Carpi",
    "28657.0Z889" : "Piacenza Centro"
}

COLONNE_ORDINATE = ["Negozio", "Data", "Ora", "Giorno","Giorno_num", "Fascia oraria", "Presenze"]

#-------------------------------------------------------------------------

#FILE:

#AGGREGA_FILE
#-------------------------------------------------------------------------
def _excel_formattato(df) -> bytes:
    buffer = BytesIO()
    df.to_excel(buffer, index=False, engine="openpyxl")
    buffer.seek(0)
    wb = load_workbook(buffer)
    ws = wb.active

    # trova colonna "Data"
    col_data_idx = None
    for i, col_name in enumerate(df.columns, start=1):
        if col_name == "Data":
            col_data_idx = i
            break

    if col_data_idx:
        col_letter = get_column_letter(col_data_idx)
        for cell in ws[col_letter][1:]:
            cell.number_format = 'DD/MM/YYYY'

    output = BytesIO()
    wb.save(output)
    return output.getvalue()

#-------------------------------------------------------------------------


def trova_database_pedonalita() -> dict | None:
    client, drive_id = _drive_sharepoint()
    candidati = [
        file for file in client.elenca_file_cartella(drive_id, PATH_DATABASE, crea=True)
        if "pedonalit" in str(file.get("name", "")).casefold()
        and str(file.get("name", "")).casefold().endswith(".parquet")
    ]
    if not candidati:
        return None
    candidati.sort(key=lambda file: str(file.get("lastModifiedDateTime", "")), reverse=True)
    scelto = candidati[0]
    return {**scelto, "content": client.scarica_file(drive_id, scelto["id"])}


def _elimina_vecchi_pedonalita(cartella: str, estensione: str) -> None:
    client, drive_id = _drive_sharepoint()
    for file in client.elenca_file_cartella(drive_id, cartella, crea=True):
        nome = str(file.get("name", ""))
        if "pedonalit" in nome.casefold() and nome.casefold().endswith(estensione):
            client.elimina_file(drive_id, str(file["id"]), str(file.get("eTag", "")))


def salva_database_pedonalita(df, ultimo_giorno: str) -> None:
    client, drive_id = _drive_sharepoint()
    parquet = BytesIO()
    df.to_parquet(parquet, index=False)
    nome_parquet = f"Pedonalità_{ultimo_giorno}.parquet"
    nome_excel = f"Pedonalità_{ultimo_giorno}.xlsx"

    _elimina_vecchi_pedonalita(PATH_DATABASE, ".parquet")
    _elimina_vecchi_pedonalita(PATH_BASI_DATI, ".xlsx")
    client.carica_bytes(
        drive_id, PATH_DATABASE, nome_parquet, parquet.getvalue(),
        content_type="application/octet-stream",
    )
    client.carica_bytes(
        drive_id, PATH_BASI_DATI, nome_excel, _excel_formattato(df),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    print(Fore.GREEN + f"SharePoint: {PATH_DATABASE}/{nome_parquet}")
    print(Fore.GREEN + f"SharePoint: {PATH_BASI_DATI}/{nome_excel}")
