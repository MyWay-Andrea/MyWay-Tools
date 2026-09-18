from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from common.negozi import PATH_TEAM_CONSUMER


NOME_FILE_OUTPUT = "TOTALE_DIGI.xlsx"
NOME_FOGLIO = "TOTALE CHIAMATE EFFETTUATE"

# Colori ripresi dal modello DIGI_TORRI.xlsx.
BLU_SCURO = "1F4E78"
CELESTE = "00B0F0"
GIALLO = "FFFF00"
NERO = "000000"
BIANCO = "FFFFFF"


def _formato_excel(nome_colonna: str) -> str:
    """Restituisce il formato Excel in base al nome della colonna."""

    nome = str(nome_colonna).strip().lower()
    if "data" in nome:
        return "dd/mm/yyyy"
    if "qt" in nome or "quantit" in nome:
        return "#,##0"
    return "@"


def _bordo_riga(indice_colonna: int, ultima_colonna: int) -> Border:
    """Restituisce il bordo del corpo tabella, con contorno esterno marcato."""

    bordo_fine = Side(style="thin", color=NERO)
    bordo_spesso = Side(style="thick", color=NERO)
    return Border(
        left=bordo_spesso if indice_colonna == 1 else bordo_fine,
        right=bordo_spesso if indice_colonna == ultima_colonna else bordo_fine,
        top=bordo_fine,
        bottom=bordo_fine,
    )


def genera_xlsx(df: pd.DataFrame, cartella_destinazione: Path = PATH_TEAM_CONSUMER) -> Path:
    """Genera TOTALE_DIGI.xlsx con lo stile del foglio DIGI dei negozi."""

    if df.empty:
        raise ValueError("Impossibile generare il file: il DataFrame e' vuoto")

    cartella_destinazione = Path(cartella_destinazione)
    if not cartella_destinazione.is_dir():
        raise FileNotFoundError(
            f"Cartella di destinazione non trovata: {cartella_destinazione}"
        )

    percorso_output = cartella_destinazione / NOME_FILE_OUTPUT
    formati_colonne = [
        _formato_excel(nome_colonna)
        for nome_colonna in df.columns
    ]
    df_excel = df.copy()
    df_excel = df_excel.astype(object).where(pd.notna(df_excel), None)

    wb = Workbook()
    ws = wb.active
    ws.title = NOME_FOGLIO
    ws.sheet_view.showGridLines = False

    for riga in dataframe_to_rows(df_excel, index=False, header=True):
        ws.append(riga)

    ultima_riga = ws.max_row
    ultima_colonna = ws.max_column
    bordo_fine = Side(style="thin", color=NERO)
    bordo_spesso = Side(style="thick", color=NERO)

    # La suddivisione cromatica segue DIGI_TORRI:
    # identificazione in blu, anagrafica in celeste, interesse in giallo.
    for indice_colonna, cella in enumerate(ws[1], start=1):
        nome = str(cella.value).strip().upper()
        if nome == "INTERESSATO A":
            colore_sfondo = GIALLO
            colore_testo = NERO
        elif nome in {"COGNOME", "NOME", "N. TELEFONO"}:
            colore_sfondo = CELESTE
            colore_testo = NERO
        else:
            colore_sfondo = BLU_SCURO
            colore_testo = BIANCO

        cella.fill = PatternFill(fill_type="solid", fgColor=colore_sfondo)
        cella.font = Font(
            name="Calibri",
            size=12,
            bold=True,
            color=colore_testo,
        )
        cella.alignment = Alignment(horizontal="center", vertical="center")
        cella.border = Border(
            left=bordo_spesso if indice_colonna == 1 else bordo_fine,
            right=(
                bordo_spesso if indice_colonna == ultima_colonna else bordo_fine
            ),
            top=bordo_spesso,
            bottom=bordo_spesso,
        )

    ws.row_dimensions[1].height = 17.25

    for indice_riga in range(2, ultima_riga + 1):
        ws.row_dimensions[indice_riga].height = 15.75
        for indice_colonna in range(1, ultima_colonna + 1):
            cella = ws.cell(indice_riga, indice_colonna)
            cella.font = Font(name="Calibri", size=11, color=NERO)
            cella.border = _bordo_riga(indice_colonna, ultima_colonna)

            cella.number_format = formati_colonne[indice_colonna - 1]

    # Chiude visivamente il contorno della tabella sull'ultima riga.
    for indice_colonna in range(1, ultima_colonna + 1):
        cella = ws.cell(ultima_riga, indice_colonna)
        cella.border = Border(
            left=bordo_spesso if indice_colonna == 1 else bordo_fine,
            right=(
                bordo_spesso if indice_colonna == ultima_colonna else bordo_fine
            ),
            top=bordo_fine,
            bottom=bordo_spesso,
        )

    larghezze = {
        "NEGOZIO": 18.0,
        "DATA": 14.85546875,
        "VENDITORE": 19.42578125,
        "COGNOME": 16.0,
        "NOME": 17.42578125,
        "N. TELEFONO": 17.42578125,
        "INTERESSATO A": 20.5703125,
    }
    for indice_colonna, nome_colonna in enumerate(df_excel.columns, start=1):
        lettera = get_column_letter(indice_colonna)
        nome_normalizzato = str(nome_colonna).strip().upper()
        ws.column_dimensions[lettera].width = larghezze.get(
            nome_normalizzato,
            15.0,
        )

    wb.save(percorso_output)
    return percorso_output
