import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from colorama import Fore, init
init(autoreset=True)

from config import (
    trova_file,
    PATH_PEDONALITA, CARTELLE_NEGOZI, MESI_EN_IT, ORDINE_MESI, MAPPA_NEGOZI, PATH_OUTPUT_REPORT
)

from estrai_trattative import main as dati_trattative

AZZURRO_HEADER  = "00B0F0"
AZZURRO_NEGOZIO = "83CCEB"
GIALLO          = "FFFF00"
NERO            = "FF000000"
BIANCO          = "FFFFFFFF"

THICK = Side(style="thick")
NO    = Side(style=None)
THIN_GIALLO = Side(style="thin", color="FFFF99")

def _fill(hex_color):
    return PatternFill(fill_type="solid", fgColor=hex_color)

def _align(h="center", v="center"):
    return Alignment(horizontal=h, vertical=v)

def _border(top=NO, bottom=NO, left=NO, right=NO):
    return Border(top=top, bottom=bottom, left=left, right=right)


def build_tabella(pivot: pd.DataFrame):
    righe, meta = [], []

    for negozio, gruppo in pivot.groupby(level="NEGOZIO", sort=False):
        subtot = gruppo.droplevel("NEGOZIO").sum()
        subtot.name = negozio
        righe.append(subtot)
        meta.append("negozio")

        for (_, venditore), row in gruppo.iterrows():
            s = pd.Series(row.values, index=pivot.columns, name=venditore)
            righe.append(s)
            meta.append("venditore")

    totale = pivot.sum()
    totale.name = "Totale complessivo"
    righe.append(totale)
    meta.append("totale")

    df = pd.DataFrame(righe)
    df.index.name = "COGNOME"
    df["Totale"] = df.sum(axis=1).astype(int)
    return df, meta


def formatta_report(pivot: pd.DataFrame, titolo: str, df_pedo: pd.DataFrame) -> None:
    df, meta = build_tabella(pivot)
    mesi    = list(pivot.columns)
    n_mesi  = len(mesi)

    wb = Workbook()
    ws = wb.active
    ws.title = "TRATTATIVE ENERGIA"

    COL_COGNOME     = 2
    COL_PRIMO_MESE  = 3
    COL_ULTIMO_MESE = COL_PRIMO_MESE + n_mesi - 1
    COL_TOTALE      = COL_ULTIMO_MESE + 1
    TUTTE_COLS      = [COL_COGNOME] + list(range(COL_PRIMO_MESE, COL_TOTALE + 1))

    # colonne pedonalità: lascia una colonna vuota dopo COL_TOTALE
    COL_PEDO_PRIMO  = COL_TOTALE + 2
    COL_PEDO_ULTIMO = COL_PEDO_PRIMO + n_mesi - 1
    TUTTE_COLS_PEDO = list(range(COL_PEDO_PRIMO, COL_PEDO_ULTIMO + 1))

    # ── larghezze colonne pedo ───────────────────────────────────────────────
    for col in TUTTE_COLS_PEDO:
        ws.column_dimensions[get_column_letter(col)].width = 12

    # ── larghezze colonne ────────────────────────────────────────────────────
    for col in TUTTE_COLS:
        ws.column_dimensions[get_column_letter(col)].width = 17.29

    # ── TITOLO (B2:ultima_mese_col, righe 2-4) ──────────────────────────────
    fine_titolo = get_column_letter(COL_ULTIMO_MESE)
    ws.merge_cells(f"B2:{fine_titolo}4")
    ws["B2"]           = "TRATTATIVE ENERGIA"
    ws["B2"].font      = Font(name="Aharoni", size=24, bold=True, color=NERO)
    ws["B2"].fill      = _fill(AZZURRO_HEADER)
    ws["B2"].alignment = _align()

    cells_titolo = list(ws[f"B2:{fine_titolo}4"])
    nr = len(cells_titolo)
    for i, row in enumerate(cells_titolo):
        nc = len(row)
        for j, cell in enumerate(row):
            cell.border = _border(
                top    = THICK if i == 0      else NO,
                bottom = THICK if i == nr - 1 else NO,
                left   = THICK if j == 0      else NO,
                right  = THICK if j == nc - 1 else NO,
            )

    # ── HEADER (riga 7) ───────────────────────────────────────────────────
    ROW_H2 = 7
  

    header_r2 = ["Etichette di riga"] + mesi + ["Totale"]

    for col_i, col in enumerate(TUTTE_COLS):
        c2 = ws.cell(ROW_H2, col)
        c2.value = header_r2[col_i]
        
        c2.font      = Font(bold=True, size=11, color=BIANCO)
        c2.fill      = _fill(AZZURRO_HEADER)
        c2.alignment = _align(h="left" if col == COL_COGNOME else "center")
        c2.border = _border(
            top   = THICK,
            left  = THICK if col == COL_COGNOME or col == COL_TOTALE    else THIN_GIALLO,
            right = THICK if col == COL_TOTALE                          else THIN_GIALLO,
        )

        ws.row_dimensions[ROW_H2].height = 28.5
        ws.column_dimensions[get_column_letter(COL_TOTALE)].width = 9

    # ── DATI ─────────────────────────────────────────────────────────────────
    ROW_START  = 8
    total_rows = len(df)

    for row_i, ((cognome, row_data), tipo) in enumerate(zip(df.iterrows(), meta)):
        excel_row = ROW_START + row_i
        is_last   = row_i == total_rows - 1

        ws.cell(excel_row, COL_COGNOME).value = f"  {cognome}"
        for i, mese in enumerate(mesi):
            ws.cell(excel_row, COL_PRIMO_MESE + i).value = int(row_data[mese])
            ws.cell(excel_row, COL_PRIMO_MESE + i).alignment = Alignment(horizontal="left", indent=3)
        ws.cell(excel_row, COL_TOTALE).value = int(row_data["Totale"])

        if tipo == "negozio":
            ws.row_dimensions[excel_row].height = 20.25
            ws.column_dimensions[get_column_letter(COL_COGNOME)].width = 26
        for col in TUTTE_COLS:
            cell = ws.cell(excel_row, col)
            is_left  = col == COL_COGNOME
            is_right = col == COL_TOTALE

            cell.border = _border(
                top    = THICK if row_i == 0 else NO,
                bottom = THICK if is_last    else NO,
                left   = THICK if is_left   or is_right else NO,
                right  = THICK if is_right              else NO,
            )

            if tipo == "negozio":
                cell.font      = Font(bold=True, size=12, color=NERO)
                cell.fill      = _fill(AZZURRO_NEGOZIO)
                cell.alignment = _align(h="center")
                cell.border = _border(
                    top    = THICK if row_i == 0 else None,#THIN_GIALLO,
                    #bottom = THIN_GIALLO,
                    left   = THICK if col == COL_COGNOME else THIN_GIALLO,
                    right  = THICK if col in (COL_TOTALE, COL_ULTIMO_MESE) else THIN_GIALLO,
                )
                


            elif tipo == "venditore":
                col_testo       = NERO
                cell.font       = Font(size=11, color=col_testo)
                cell.alignment = _align(h="center") if is_right else Alignment(horizontal="left", indent=1)

            elif tipo == "totale":
                cell.font      = Font(bold=True, size=12, color=NERO)
                cell.alignment = _align(h="left" if is_left else "center")
                # bordo thick anche in alto per chiudere la tabella
                cell.border = _border(
                    top    = THICK,
                    bottom = THICK,
                    left   = THICK if is_left   or is_right else NO,
                    right  = THICK if is_right              else NO,
                )
                if col == COL_ULTIMO_MESE:
                    cell.fill = _fill(GIALLO)
                    cell.font = Font(bold=True, size=14, color="FFFF0000")


     # ── aggregazione pedonalità per negozio e mese ───────────────────────────
    df_pedo["Data"] = pd.to_datetime(df_pedo["Data"])
    df_pedo = df_pedo[df_pedo["Data"].dt.year == 2026]
    df_pedo["MESE"] = df_pedo["Data"].dt.month_name()
    df_pedo["MESE"] = df_pedo["MESE"].map(MESI_EN_IT)
    df_pedo = df_pedo[df_pedo["MESE"].isin(mesi)]  # tiene solo i mesi delle trattative
    

    pedo_agg = df_pedo.groupby(["Negozio", "MESE"])["Presenze"].sum().unstack(fill_value=0) 
    colonne_pedo = [m for m in ORDINE_MESI if m in pedo_agg.columns]
    pedo_agg = pedo_agg[colonne_pedo]

    pedo_agg.index = pedo_agg.index.map(MAPPA_NEGOZI)
    pedo_agg = pedo_agg[pedo_agg.index.notna()]
    
    
    # totale pedonalità
    pedo_totale = pedo_agg.sum()

    # ── TITOLO PEDONALITÀ (riga 6) ───────────────────────────────────────────
    col_start_l = get_column_letter(COL_PEDO_PRIMO)
    col_end_l   = get_column_letter(COL_PEDO_ULTIMO)
    ws.merge_cells(f"{col_start_l}6:{col_end_l}6")
    c = ws.cell(6, COL_PEDO_PRIMO)
    c.value     = "PEDONALITÀ"
    c.font      = Font(bold=True, size=12, color=BIANCO)
    c.fill      = _fill(AZZURRO_HEADER)
    c.alignment = _align(h="center")
    # bordo titolo
    for j, col in enumerate(TUTTE_COLS_PEDO):
        ws.cell(6, col).border = _border(
            top    = THICK,
            left   = THICK if j == 0                        else NO,
            right  = THICK if j == len(TUTTE_COLS_PEDO) - 1 else NO,
        )

    # ── HEADER MESI PEDONALITÀ (riga 7) ─────────────────────────────────────
    for j, (col, mese) in enumerate(zip(TUTTE_COLS_PEDO, mesi)):
        c = ws.cell(7, col)
        c.value     = mese
        c.font      = Font(bold=True, size=11, color=BIANCO)
        c.fill      = _fill(AZZURRO_HEADER)
        c.alignment = _align(h="center")
        c.border    = _border(
            left   = THICK if j == 0                        else NO,
            right  = THICK if j == len(TUTTE_COLS_PEDO) - 1 else NO,
        )

    # ── DATI PEDONALITÀ ──────────────────────────────────────────────────────
    ROW_START  = 8
    for row_i, tipo in enumerate(meta):
        excel_row = ROW_START + row_i
        is_last   = row_i == len(meta) - 1
        cognome   = df.index[row_i]

        for j, (col, mese) in enumerate(zip(TUTTE_COLS_PEDO, colonne_pedo)):
            cell = ws.cell(excel_row, col)
            is_first_pedo = j == 0
            is_last_pedo  = j == len(TUTTE_COLS_PEDO) - 1

            cell.border = _border(
                top    = THICK if row_i == 0 else NO,
                bottom = THICK if is_last    else NO,
                left   = THICK if is_first_pedo else NO,
                right  = THICK if is_last_pedo  else NO,
            )

            if tipo == "negozio":
                negozio_nome = cognome.strip()
                if negozio_nome in pedo_agg.index and mese in pedo_agg.columns:
                    cell.value = int(pedo_agg.loc[negozio_nome, mese])
                cell.font      = Font(bold=False, size=12, color=NERO)
                cell.alignment = _align(h="center")
                cell.fill      = _fill(AZZURRO_NEGOZIO)

            elif tipo == "totale":
                cell.value     = int(pedo_totale[mese]) if mese in pedo_totale else 0
                cell.font      = Font(bold=True, size=12, color=NERO)
                cell.alignment = _align(h="center")
                cell.border = _border(
                    top    = THICK,
                    bottom = THICK,
                    left   = THICK if is_first_pedo else NO,
                    right  = THICK if is_last_pedo  else NO,
                )
    # ── colonne ratio ────────────────────────────────────────────────────────
    COL_RATIO_PRIMO  = COL_PEDO_ULTIMO + 2
    COL_RATIO_ULTIMO = COL_RATIO_PRIMO + n_mesi - 1
    TUTTE_COLS_RATIO = list(range(COL_RATIO_PRIMO, COL_RATIO_ULTIMO + 1))

    for col in TUTTE_COLS_RATIO:
        ws.column_dimensions[get_column_letter(col)].width = 12

    # ── aggregazione ratio ───────────────────────────────────────────────────
    # pivot trattative per negozio (solo subtotali, non venditori)
    trat_agg = {}
    for row_i, tipo in enumerate(meta):
        if tipo == "negozio":
            negozio_nome = df.index[row_i].strip()
            trat_agg[negozio_nome] = {m: int(df.iloc[row_i][m]) for m in mesi}

    def calcola_ratio(trat, pedo):
        if pedo and pedo > 0:
            return round(trat / pedo * 100, 2)
        return 0

    # totali ratio
    ratio_totale = {}
    for mese in mesi:
        t = sum(trat_agg[n][mese] for n in trat_agg if mese in trat_agg[n])
        p = int(pedo_totale[mese]) if mese in pedo_totale else 0
        ratio_totale[mese] = calcola_ratio(t, p)

    # ── TITOLO RATIO (riga 6) ────────────────────────────────────────────────
    col_start_r = get_column_letter(COL_RATIO_PRIMO)
    col_end_r   = get_column_letter(COL_RATIO_ULTIMO)
    ws.merge_cells(f"{col_start_r}6:{col_end_r}6")
    c = ws.cell(6, COL_RATIO_PRIMO)
    c.value     = "RATIO"
    c.font      = Font(bold=True, size=12, color=BIANCO)
    c.fill      = _fill(AZZURRO_HEADER)
    c.alignment = _align(h="center")
    for j, col in enumerate(TUTTE_COLS_RATIO):
        ws.cell(6, col).border = _border(
            top   = THICK,
            left  = THICK if j == 0                         else NO,
            right = THICK if j == len(TUTTE_COLS_RATIO) - 1 else NO,
        )

    # ── HEADER MESI RATIO (riga 7) ───────────────────────────────────────────
    for j, (col, mese) in enumerate(zip(TUTTE_COLS_RATIO, mesi)):
        c = ws.cell(7, col)
        c.value     = mese
        c.font      = Font(bold=True, size=11, color=BIANCO)
        c.fill      = _fill(AZZURRO_HEADER)
        c.alignment = _align(h="center")
        c.border    = _border(
            left  = THICK if j == 0                         else NO,
            right = THICK if j == len(TUTTE_COLS_RATIO) - 1 else NO,
        )

    # ── DATI RATIO ───────────────────────────────────────────────────────────
    for row_i, tipo in enumerate(meta):
        excel_row    = ROW_START + row_i
        is_last      = row_i == len(meta) - 1
        cognome      = df.index[row_i].strip()

        for j, (col, mese) in enumerate(zip(TUTTE_COLS_RATIO, mesi)):
            cell          = ws.cell(excel_row, col)
            is_first_ratio = j == 0
            is_last_ratio  = j == len(TUTTE_COLS_RATIO) - 1

            cell.border = _border(
                top    = THICK if row_i == 0 else NO,
                bottom = THICK if is_last    else NO,
                left   = THICK if is_first_ratio else NO,
                right  = THICK if is_last_ratio  else NO,
            )

            if tipo == "negozio":
                t = trat_agg.get(cognome, {}).get(mese, 0)
                p = int(pedo_agg.loc[cognome, mese]) if cognome in pedo_agg.index and mese in pedo_agg.columns else 0
                cell.value     = calcola_ratio(t, p)
                cell.font      = Font(bold=False, size=12, color=NERO)
                cell.fill      = _fill(AZZURRO_NEGOZIO)
                cell.alignment = _align(h="center")
                cell.value = calcola_ratio(t, p) / 100
                cell.number_format = "0.00%"


            elif tipo == "totale":
                ratio_val           = ratio_totale.get(mese, 0)          # già calcolato correttamente
                cell.value          = ratio_val / 100
                cell.number_format  = "0.00%"
                cell.font           = Font(bold=True, size=12, color=NERO)
                cell.alignment      = _align(h="center")
                cell.border         = _border(
                    top    = THICK,
                    bottom = THICK,
                    left   = THICK if is_first_ratio else NO,
                    right  = THICK if is_last_ratio  else NO,
                )
                if col == COL_RATIO_ULTIMO:
                    cell.fill = _fill(GIALLO)
                    cell.font = Font(bold=True, size=14, color="FFFF0000")

                    
    filename = PATH_OUTPUT_REPORT / titolo        
    wb.save(f"{filename}.xlsx")
    print(Fore.GREEN + f" ✅ Report: "+ Fore.RESET + f"{titolo}" + Fore.GREEN + " caricato con successo")
    

def main(pedo: Path, files_trattative:list[Path]):
    
    print(" === Estrazione File Tracciamento da Teams === \n")
    # 1. Estrai dati trattative energia
    df_trat = dati_trattative(files_trattative)
    
    print(" Estrazione completata.\n")

    print(" === Elaborazione Trattative === \n")

    # 2. Crea colonna mese
    df_trat["MESE"] = df_trat["DATA"].dt.month_name()
    print(Fore.GREEN + " ✓  " + Fore.RESET + "Colonna mese creata") 

    # 3. Crea pivot pe aggregare dati
    pivot_trat = df_trat.pivot_table(
        index      = ["NEGOZIO", "VENDITORE"],
        columns    = "MESE",
        values     = "DATA",
        aggfunc    = "count",
        fill_value = 0
    )
    print(Fore.GREEN + " ✓  " + Fore.RESET + "Pivot creata") 

    # 4. Converti da mesi in Eng a mesi in Ita
    pivot_trat = pivot_trat.rename(columns=MESI_EN_IT)

    # 5. Ordina colonne in modo che mesi siano corretti
    colonne_presenti = [m for m in ORDINE_MESI if m in pivot_trat.columns]
    pivot_trat = pivot_trat[colonne_presenti]

    print(Fore.GREEN + " ✓  " + Fore.RESET + "Riorganizzazione delle colonne\n") 
    
    df_pedo = pd.read_parquet(pedo)

    print(" === Creazione REPORT_TRATTATIVE_ENERGIA === \n")

    formatta_report(pivot_trat, "REPORT_TRATTATIVE_ENERGIA", df_pedo)


if __name__ == "__main__":

    file_pedo = trova_file(PATH_PEDONALITA, "Pedonalit")
    files = trova_file(CARTELLE_NEGOZI, "trattative")

    main(file_pedo, files)
