from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

from common.negozi import PATH_TEAM_CONSUMER

TGT_VENDITE = 327
CROSS_SELLING = "acquisto prodotto-servizio"


def genera_xlsx(df: pd.DataFrame, cartella_destinazione: Path = PATH_TEAM_CONSUMER,) -> Path:
    """Crea la tabella Excel aggregata e restituisce il percorso del file."""

    if df.empty:
        raise ValueError("Impossibile generare il file: il DataFrame è vuoto")

    if not cartella_destinazione.is_dir():
        raise FileNotFoundError(
            f"Cartella di destinazione non trovata: {cartella_destinazione}"
        )

    percorso_output = cartella_destinazione / "TRACCIAMENTO_GADGET.xlsx"

    # Converte NaN e NaT in celle vuote compatibili con Excel.
    df_excel = df.astype(object).where(pd.notna(df), None)

    wb = Workbook()
    ws = wb.active
    ws.title = "TRACCIAMENTO_GADGET"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    for riga in dataframe_to_rows(df_excel, index=False, header=True):
        ws.append(riga)


    ultima_riga = ws.max_row
    ultima_colonna = ws.max_column
    riferimento_tabella = (
        f"A1:{get_column_letter(ultima_colonna)}{ultima_riga}"
    )

    tabella = Table(
        displayName="TabellaTracciamentoGadget",
        ref=riferimento_tabella,
    )
    tabella.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tabella)

    riempimento_intestazione = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cella in ws[1]:
        cella.font = Font(color="FFFFFF", bold=True)
        cella.fill = riempimento_intestazione
        cella.alignment = Alignment(horizontal="center", vertical="center")

    venduti = len(df)
    cross_selling = (
        df["ACQUISTO - ATTIVAZIONE"]
        .astype("string")
        .str.strip()
        .str.casefold()
        .eq(CROSS_SELLING)
        .sum()
    )
    percentuale_cross_selling = cross_selling / venduti

    colonna_riepilogo = ultima_colonna + 2
    intestazioni_riepilogo = [
        "VENDUTI",
        "TGT VENDITE",
        "CROSS SELLING",
        "% CROSS SELLING",
    ]

    valori_riepilogo = [
        venduti,
        TGT_VENDITE,
        int(cross_selling),
        percentuale_cross_selling,
    ]

    riga_intestazioni = 5
    riga_valori = 6
    bordo_spesso = Side(style="thick", color="000000")
    bordo_fine = Side(style="thin", color="000000")
    stili_intestazioni = [
        ("FFFFFF", "008000"),
        ("FFFF00", "000000"),
        ("FF0000", "FFFFFF"),
        ("FF0000", "FFFFFF"),
    ]

    for offset, (intestazione, valore, stile) in enumerate(
        zip(intestazioni_riepilogo, valori_riepilogo, stili_intestazioni)
    ):
        colonna = colonna_riepilogo + offset
        colore_sfondo, colore_testo = stile
        prima_colonna = offset == 0
        ultima_colonna_riepilogo = offset == len(intestazioni_riepilogo) - 1

        cella_intestazione = ws.cell(
            row=riga_intestazioni,
            column=colonna,
            value=intestazione,
        )
        cella_intestazione.font = Font(color=colore_testo, bold=True)
        cella_intestazione.fill = PatternFill(
            fill_type="solid",
            fgColor=colore_sfondo,
        )
        cella_intestazione.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cella_intestazione.border = Border(
            top=bordo_spesso,
            bottom=bordo_spesso,
            left=bordo_spesso if prima_colonna else bordo_fine,
            right=bordo_spesso if ultima_colonna_riepilogo else bordo_fine,
        )

        cella_valore = ws.cell(
            row=riga_valori,
            column=colonna,
            value=valore,
        )
        cella_valore.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cella_valore.border = Border(
            top=bordo_spesso,
            bottom=bordo_spesso,
            left=bordo_spesso if prima_colonna else bordo_fine,
            right=bordo_spesso if ultima_colonna_riepilogo else bordo_fine,
        )
        ws.column_dimensions[get_column_letter(colonna)].width = max(
            len(intestazione) + 5,
            14,
        )

    ws.cell(
        row=riga_valori,
        column=colonna_riepilogo + len(intestazioni_riepilogo) - 1,
    ).number_format = "0.00%"

    for indice, nome_colonna in enumerate(df_excel.columns, start=1):
        lettera_colonna = get_column_letter(indice)

        if "data" in str(nome_colonna).strip().lower():
            for cella in ws[lettera_colonna][1:]:
                cella.number_format = "dd/mm/yyyy"

        lunghezza_massima = max(
            len(str(cella.value)) if cella.value is not None else 0
            for cella in ws[lettera_colonna]
        )
        ws.column_dimensions[lettera_colonna].width = min(
            max(lunghezza_massima + 2, 12),
            40,
        )

    wb.save(percorso_output)
    return percorso_output
