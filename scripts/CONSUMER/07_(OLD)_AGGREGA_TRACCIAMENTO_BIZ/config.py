from pathlib import Path
import pandas as pd
from datetime import date
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys
from colorama import Fore, init
init(autoreset=True)
#IMPOSTAZIONE
PATH_HOME = Path.home()

ONEDRIVE_PATH = {
    "TORRI" : PATH_HOME / "My Way S.r.l" / "TEAM - TORRI - TEAM - TORRI" / "TRACCIAMENTO",
    "CARPI": PATH_HOME / "My Way S.r.l" / "TEAM - CARPI - TEAM - CARPI" / "TRACCIAMENTO",
    "EUROSIA" : PATH_HOME / "My Way S.r.l" / "TEAM - EUROSIA - TEAM - EUROSIA" / "TRACCIAMENTO",
    "FIDENZA" : PATH_HOME / "My Way S.r.l" / "TEAM - FIDENZA - TEAM - FIDENZA" / "TRACCIAMENTO",
    "GALASSIA" : PATH_HOME / "My Way S.r.l" / "TEAM - GALASSIA - TEAM - GALASSIA" / "TRACCIAMENTO",
    "PIACENZA MEDIAWORLD"   :PATH_HOME / "My Way S.r.l" / "TEAM - PIACENZA MED - TEAM - PIACENZA MED" / "TRACCIAMENTO",
    "SASSUOLO" : PATH_HOME / "My Way S.r.l" / "TEAM - SASSUOLO - TEAM - SASSUOLO" / "TRACCIAMENTO",
}

PATH_TEAM_CONSUMER = PATH_HOME / "My Way S.r.l" / "TEAM CONSUMER - TEAM CONSUMER"

COLONNE_ORDINATE = ["NEGOZIO", "VENDITORE", "RAGIONE SOCIALE", "PARTITA IVA", "DATA INSERIMENTO", "DATA ATTIVAZIONE", "CATEGORIA", "SERVIZIO", "Qtà"]

#FUNZIONI

def trova_file(paths: dict) -> list[Path]:
    files= []
    ricerca = "*BIZ.xlsx"
    if isinstance (paths, dict):
        for name, path in paths.items():
            temp_list_file = list(path.glob(ricerca)) 
            for file in temp_list_file:
                files.append(file)
                if temp_list_file:
                    print(f"{name} ✓")

    return files

   
    
def _read_file(file) -> pd.DataFrame | None:
    if not file.is_file():
        return None
    
    df = pd.read_excel(file, dtype = str)
    df = df.reindex(columns = COLONNE_ORDINATE).copy()
    df["NEGOZIO"] = file.name.split("_")[0]
    return df

def create_aggregate_df(files_list:list) -> pd.DataFrame | None:
    '''
    legge la lista di file estratti dalle cartelle di teams e li trasforma in df, 
    li inserisce in un una lista, e ritorna la lista aggregata sotto forma 
    di un unico df
    '''
    lista_df = []
    for file in files_list:
        df = _read_file(file)
        if df is not None:
            lista_df.append(df)
        
    return pd.concat(lista_df,ignore_index = True) if lista_df else None


def normalize_df(df:pd.DataFrame) -> pd.DataFrame | None:
    '''
    prende in ingresso il df aggreagto e normalizza le colonne data e p.iva 
    '''
    if df is None:
        print("df is None")
        return None
    
    #formatta P.iva
    df["PARTITA IVA"] = df["PARTITA IVA"].astype(str).str.zfill(11)

    #formatta date 
    df["DATA INSERIMENTO"] = pd.to_datetime(df["DATA INSERIMENTO"], format = "%d/%m/%Y", errors = "coerce")
    df["DATA ATTIVAZIONE"] = pd.to_datetime(df["DATA ATTIVAZIONE"], format ="%d/%m/%Y", errors = "coerce")

    #formatta Qtà (tutte le celle senza esplicitazione della Qtà verrà inserito 1)
    df["Qtà"] =pd.to_numeric(df["Qtà"], errors = "coerce").fillna(1)
    return df


#cancellare file vecchio
def delete_old_file(path, param = None):
    import time
    if param is None:
        print("Non è stato inserito il parametro.\nEsco..👋")
        sys.exit(1)
    if path.is_dir():
        files_list = path.glob(param)
        for file in files_list:
            file.unlink()
            time.sleep(2)
            if file.exists():
                print(Fore.RED + "\n ❌ Errore -> il file non è stato cancellato")
            else:
                print("\nfile " + Fore.RED + "cancellato" + Fore.RESET +  " correttamente")
        
#per creazione file report
'''
Genera il foglio excel con 3 fogli:
1. Il primo foglio è quello "DB" dove saranno stampati tutti i record 
    def df aggregato.
2. Il secondo foglio "Venditori" sarà un tabella pivot, con nelle colonne 
    i valori dei venditori, sulle righe le categorie seguide dai loro servizi
    di riferimento, e nei valori il conteggio per venditore di quanti servizi 
    ha venduto, e nella riga della categoria c'è il subtotale 
3. il terzo foglio "Negozi" è l stessa tabella dei venditori solo che al posto
    di aggregare per venditori si fa direttamente in modo macroscopico per negozio
'''
def _save_first_sheet(df: pd.DataFrame, path: Path) -> Path:
    oggi = date.today().strftime("%d-%m-%y")  
    filename = path / f"TRACCIAMENTO_PISTA_BIZ_{oggi}.xlsx"
    df.to_excel(filename, index=False, sheet_name="DB")
    print(" - Foglio DB creato ✓")
    return filename


def _scrivi_foglio_pivot(ws, df: pd.DataFrame, col_gruppo: str):
    """Scrive un foglio pivot: righe = CATEGORIA > SERVIZIO, colonne = col_gruppo."""
    # Sfondi
    BLU        = PatternFill('solid', start_color='4472C4')
    BLU_CHIARO = PatternFill('solid', start_color='D9E1F2')
    BLU_MED    = PatternFill('solid', start_color='8EA9DB')
    GRIGIO     = PatternFill('solid', start_color='F2F2F2')
    BIANCO     = PatternFill('solid', start_color='FFFFFF')

    # Font
    font_h        = Font(bold=True, color='FFFFFF')
    font_bold     = Font(bold=True, color='000000')
    font_normale  = Font(color='000000')
    font_zero     = Font(color='DEDCDC')

    # Bordo griglia grigio
    grigio_border = Side(style='thin', color='DEDCDC')
    bordo         = Border(
        left=grigio_border, right=grigio_border,
        top=grigio_border,  bottom=grigio_border
    )
    center = Alignment(horizontal='center', vertical='center')

    gruppi    = sorted(df[col_gruppo].dropna().unique())
    ORDINE_CATEGORIE = ["Voce_BIZ", "Dati_BIZ", "Fissa_BIZ", "Solution_BIZ", "Altro"]
    categorie_presenti = df['CATEGORIA'].dropna().unique()
    '''
    categorie = [c for c in ORDINE_CATEGORIE if c in categorie_presenti] + \
                [c for c in categorie_presenti if c not in ORDINE_CATEGORIE]
    
    doppio controllo perchè cerca prima e mette in ordine in base alla lista "ORDINE_CATEGORIE"
    e poi fa un altro controllo ep vedere se nei dati c'è qualche categoria nuova che non è presente 
    nella lista, es. qualche promo estiva per sempio
    ["Voce_BIZ", "Dati_BIZ", "Fissa_BIZ", "Solution_BIZ", "Altro"] + ["Promo estiva"]
    '''

    categorie = [c for c in ORDINE_CATEGORIE if c in categorie_presenti] + \
                [c for c in categorie_presenti if c not in ORDINE_CATEGORIE]
    
    tot_col   = len(gruppi) + 2

    # --- INTESTAZIONE ---
    cell = ws.cell(row=1, column=1, value='CATEGORIA / SERVIZIO')
    cell.fill, cell.font, cell.border = BLU, font_h, bordo
    cell.alignment = Alignment(horizontal='left', vertical='center')

    for c_idx, g in enumerate(gruppi, start=2):
        cell = ws.cell(row=1, column=c_idx, value=g)
        cell.fill, cell.font, cell.alignment, cell.border = BLU, font_h, center, bordo

    cell = ws.cell(row=1, column=tot_col, value='Totale')
    cell.fill, cell.font, cell.alignment, cell.border = BLU, font_h, center, bordo

    # --- RIGA SUBTOTALE GRUPPI ---
     
    cell = ws.cell(row=2, column=1, value='Totale contratti')
    cell.fill, cell.font, cell.border = BLU_MED, font_bold, bordo
    cell.alignment = Alignment(horizontal='left', vertical='center')

    for c_idx, g in enumerate(gruppi, start=2):
        count = int(df.loc[df[col_gruppo] == g, 'Qtà'].sum())
        cell = ws.cell(row=2, column=c_idx, value=count)
        cell.fill      = BLU_MED
        cell.font      = font_zero if count == 0 else font_bold
        cell.alignment = center
        cell.border    = bordo

    gran_tot = int(df['Qtà'].sum())
    cell = ws.cell(row=2, column=tot_col, value=gran_tot)
    cell.fill, cell.font, cell.alignment, cell.border = BLU_MED, font_bold, center, bordo

    # --- BORDO SPESSO SOTTO RIGA 2 ---
    for col in range(1, tot_col + 1):
        ws.cell(row=2, column=col).border = Border(
            left=grigio_border, right=grigio_border,
            top=grigio_border,  bottom=Side(style='medium', color='000000')
        )

    # --- RIGHE DATI ---
    row = 3
    for cat in categorie:
        df_cat = df[df['CATEGORIA'] == cat]
        servizi = sorted(df_cat['SERVIZIO'].dropna().unique())

        # Riga CATEGORIA (subtotale)
        cell = ws.cell(row=row, column=1, value=cat)
        cell.fill, cell.font, cell.border = BLU_CHIARO, font_bold, bordo

        for c_idx, g in enumerate(gruppi, start=2):
            count = int(df_cat.loc[df_cat[col_gruppo] == g, 'Qtà'].sum())
            cell = ws.cell(row=row, column=c_idx, value=count if count else 0)
            cell.fill      = BLU_CHIARO
            cell.font      = font_zero if count == 0 else font_bold
            cell.alignment = center
            cell.border    = bordo

        tot = int(df_cat['Qtà'].sum())
        cell = ws.cell(row=row, column=tot_col, value=tot)
        cell.fill, cell.font, cell.alignment, cell.border = BLU_CHIARO, font_bold, center, bordo
        row += 1

        # Righe SERVIZIO
        for i, serv in enumerate(servizi):
            df_serv = df_cat[df_cat['SERVIZIO'] == serv]
            fill = GRIGIO if i % 2 == 0 else BIANCO

            cell = ws.cell(row=row, column=1, value=f'    {serv}')
            cell.fill, cell.font, cell.border = fill, font_normale, bordo

            tot_riga = 0
            for c_idx, g in enumerate(gruppi, start=2):
                count = int(df_serv.loc[df_serv[col_gruppo] == g, 'Qtà'].sum())
                cell = ws.cell(row=row, column=c_idx, value=count if count else 0)
                cell.fill      = fill
                cell.font      = font_zero if count == 0 else font_normale
                cell.alignment = center
                cell.border    = bordo
                tot_riga += count

            cell = ws.cell(row=row, column=tot_col, value=tot_riga)
            cell.fill      = fill
            cell.font      = font_zero if tot_riga == 0 else font_normale
            cell.alignment = center
            cell.border    = bordo
            row += 1

    # --- LARGHEZZE & ALTEZZE---
    #larghezza
    ws.column_dimensions['A'].width = 30
    for i in range(2, tot_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16

    #altezza
    ws.row_dimensions[1].height = 30 
    ws.row_dimensions[2].height = 30  


def _create_pivot_sheets(filename: Path, df: pd.DataFrame):
    wb = load_workbook(filename)
    ws_venditori = wb.create_sheet("Venditori")
    _scrivi_foglio_pivot(ws_venditori, df, col_gruppo='VENDITORE')
    ws_negozi = wb.create_sheet("Negozi")
    _scrivi_foglio_pivot(ws_negozi, df, col_gruppo='NEGOZIO')
    wb.save(filename)
    print(" - Fogli pivot creati ✓")


def write_file_xlsx(df: pd.DataFrame, path: Path):
    print("\nInizio generazione report...")
    try:
        output_file = _save_first_sheet(df, path)
        _create_pivot_sheets(output_file, df)
        print(Fore.GREEN + "\n ✅ Report salvato in: " + Fore.RESET + f"{output_file} ✓")
    except Exception as e:
        print(Fore.RED + f"\n ❌ Errore durante la creazione del file:\n{e}")
    